"""Satış ekibi için açıklayıcı Excel raporu: Bursa_{N}_Satisci_Bolgeleme.xlsx

Sayfalar
    OZET            başlık bandı, KPI kartları, bölge tablosu (canlı formül), grafikler
    PERSONEL_ATAMA  bölge başına satış sorumlusu seçimi (açılır liste, otomatik atama yok)
    BOLGE_DETAY     coğrafi göstergeler: ofise mesafe, kompaktlık, rota yükü, ilçe/mahalle
    BINA_ATAMA      19.706 binanın tamamı ve bölgesi (düzenlenebilir: bölge no)
    ILCE_MATRIS     ilçe × bölge (RES HP, bina)
    MAHALLE_MATRIS  mahalle × bölge; birden çok bölgeye bölünen mahalleler işaretli
    METODOLOJI      tek ekran, sade anlatım
    VERI_KALITESI   data/master/veri_kalitesi.json özeti
    PARAMETRELER    algoritma sürümü, seed, hedef, süre, tarih
    LISTE (gizli)   açılır liste ve bölge adları

OZET, PERSONEL_ATAMA ve matrislerdeki sayılar BINA_ATAMA'dan canlı hesaplanır (SUMIFS/COUNTIFS,
okunur olsun diye adlandırılmış aralıklarla: ATAMA_BOLGE, ATAMA_RES_HP ...). Bir binanın bölge
numarası değiştirilirse toplamlar ve sapmalar kendiliğinden güncellenir.

openpyxl formüllerin sonucunu hesaplamaz; kaydettikten sonra formül hücrelerine pandas ile
hesaplanan önbellek değerlerini yazıyoruz. Böylece hesap yapmayan görüntüleyiciler (telefon,
e-posta önizlemesi, Korumalı Görünüm) de doğru sayıları gösterir; Excel açılışta yine yeniden hesaplar.
"""
from __future__ import annotations

import datetime as dt
import json
import math
import re
import sys
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

import numpy as np
import openpyxl
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import AxDataSource, NumData, NumFmt, NumVal, StrData, StrRef, StrVal
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun, RichTextProperties
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from . import config
from .enrich import tr_norm
from .metrics import renkler as varsayilan_renkler
from .partition import ALGORITMA_SURUMU

# ----------------------------------------------------------------------------- görünüm
LACIVERT = "1B2A41"
VURGU = "FFC400"
BANT_ALT = "C9D3E3"
ZEMIN = "F4F6F9"
ZEBRA = "F7F9FC"
CIZGI = "D9DEE7"
GRI = "6B7280"
KOYU = "1F2937"
GIRDI = "FFF8DB"            # kullanıcının dolduracağı hücreler
YESIL, YESIL_Z = "15803D", "DCFCE7"
AMBER, AMBER_Z = "B45309", "FEF3C7"
KIRMIZI, KIRMIZI_Z = "B91C1C", "FEE2E2"
HEDEF_CIZGI = "C0392B"

FMT_SAYI = "#,##0"
FMT_YUZDE = "%0.0"                     # Türkçe yazım: %48,3
FMT_SAPMA = '+%0.00;"−"%0.00;%0.00'   # +%0,51 / −%0,35 / %0,00 (işaret %'den önce)
FMT_KOD = '"B"0'                       # hücre değeri 3, görünen "B3"
FMT_KM = "0.0"
FMT_MATRIS = '#,##0;-#,##0;"-"'

AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
OLCU_ETIKET = {"res_hp": "RES HP", "firsat": "Fırsat", "toplam_hp": "Toplam HP", "bina": "Bina"}
YESIL_SINIR, AMBER_SINIR = 0.02, 0.05  # sapma renkleri: ±%2 yeşil, %2–5 amber, üstü kırmızı

# BINA_ATAMA kolonları: (anahtar, başlık, genişlik, sayı biçimi)
BINA_KOLONLARI = [
    ("bolge", "Bölge", 8, FMT_KOD),
    ("bolge_adi", "Bölge Adı", 30, None),
    ("sorumlu", "Satış Sorumlusu", 24, None),
    ("bina_serial", "Bina Serial", 17, None),
    ("tellcordia_id", "Tellcordia ID", 15, None),
    ("location_id", "Location ID", 12, None),
    ("ad", "Bina Adı", 34, None),
    ("site_adi", "Site Adı", 24, None),
    ("blok_adi", "Blok", 12, None),
    ("kapi_no", "Kapı No", 9, None),
    ("mahalle", "Mahalle", 18, None),
    ("adres", "Cadde / Sokak", 28, None),
    ("ilce", "İlçe", 14, None),
    ("il", "İl", 8, None),
    ("obek", "Öbek", 20, None),
    ("lat", "Enlem", 11, "0.000000"),
    ("lon", "Boylam", 11, "0.000000"),
    ("kat", "Kat", 6, "0"),
    ("toplam_hp", "Toplam HP", 10, FMT_SAYI),
    ("res_hp", "RES HP", 10, FMT_SAYI),
    ("aktif_res", "Aktif Abone (RES)", 11, FMT_SAYI),
    ("firsat", "Fırsat", 10, FMT_SAYI),
    ("penetrasyon", "Penetrasyon", 12, FMT_YUZDE),
    ("not", "Not", 16, None),
]
BK = {k: get_column_letter(i + 1) for i, (k, *_r) in enumerate(BINA_KOLONLARI)}

# BINA_ATAMA üzerindeki adlandırılmış aralıklar (formüller okunur olsun diye)
ATAMA_ADLARI = {"bolge": "ATAMA_BOLGE", "bina_serial": "ATAMA_SERIAL", "ilce": "ATAMA_ILCE", "mahalle": "ATAMA_MAHALLE",
                "toplam_hp": "ATAMA_TOPLAM_HP", "res_hp": "ATAMA_RES_HP", "aktif_res": "ATAMA_AKTIF", "firsat": "ATAMA_FIRSAT"}

PARAM_ETIKET = {
    "algoritma": "Algoritma sürümü", "olcu": "Denge ölçüsü", "haric_il": "Hariç tutulan il", "mod": "Çözüm modu",
    "tohum_sayisi": "Denenen başlangıç sayısı", "secilen_tohum": "Seçilen başlangıç", "seed": "Rastgelelik tohumu (seed)",
    "birim_sayisi": "Birim (site grubu) sayısı", "dis_iterasyon": "Merkez güncelleme turu", "grup_sayisi": "Üst grup sayısı",
    "grup_boyutlari": "Grup başına bölge", "ada_duzeltme": "Ada onarımı (birim)", "ince_ayar": "Genel ince ayar",
    "denenen_grup_sayilari": "Denenen üst grup sayıları",
}


# ----------------------------------------------------------------------------- küçük yardımcılar
def _tr(x: float, ondalik: int = 0) -> str:
    """Türkçe sayı yazımı: 19.706 · 37.097,25"""
    s = f"{x:,.{ondalik}f}"
    return s.replace(",", "·").replace(".", ",").replace("·", ".")


EKSI = "−"                       # tipografik eksi: "−%0,35" (uygulamayla aynı yazım)


def _tr_yuzde(x: float, ondalik: int = 0) -> str:
    """Türkçe yüzde: işaret %'den önce gelir. Sıfıra yuvarlananda işaret yazılmaz."""
    v = x * 100
    if abs(v) < 0.5 * 10.0 ** (-ondalik):
        return "%" + _tr(0.0, ondalik)
    return (EKSI if v < 0 else "") + "%" + _tr(abs(v), ondalik)


def _tr_isaretli(x: float, ondalik: int = 2) -> str:
    """Her zaman işaretli yüzde: +%0,51 / −%0,35 / %0,00."""
    v = x * 100
    if abs(v) < 0.5 * 10.0 ** (-ondalik):
        return "%" + _tr(0.0, ondalik)
    return ("+" if v > 0 else EKSI) + "%" + _tr(abs(v), ondalik)


def _temiz_ad(seri: pd.Series) -> pd.Series:
    """Ham veride metin olarak gelen 'Null' / 'None' / 'nan' adlarını boşa çevirir.

    Kaynak CSV'de site_adi/ad/blok_adi alanlarında 559 hücre kelimesi kelimesine "Null" yazıyor;
    filtrelenince "Null" diye bir site adı görünüyordu. Uygulama tarafında selectors.temizAd
    aynı işi yapıyor — Excel de aynı listeyi kullanır.
    """
    m = seri.astype("string")
    return m.where(~m.str.strip().str.lower().isin(["null", "none", "nan", "<na>"]), "").fillna("")


def _yuzde_yuvarla(x: float, ondalik: int = 2) -> float:
    """Oranı, yüzdesi `ondalik` basamağa yuvarlanmış hâline eşitler.

    OZET ve PARAMETRELER sayfalarındaki `ROUND(...,4)` formülleriyle AYNI değeri verir:
    0,000613 → %0,06. (Eskiden yukarı yuvarlanıyordu ve kapak %0,07 derken dört satır
    aşağıdaki tablo %0,06 diyordu.)
    """
    k = 10.0 ** (ondalik + 2)
    return round(x, ondalik + 2) if k else x


def _tr_tarih(d: dt.datetime) -> str:
    return f"{d.day} {AYLAR[d.month - 1]} {d.year}"


def _dolgu(renk: str) -> PatternFill:
    return PatternFill("solid", start_color=renk, end_color=renk)


def _hex(renk: str) -> str:
    return renk.lstrip("#").upper()


def _yazi_rengi(renk: str) -> str:
    r, g, b = (int(renk[i:i + 2], 16) for i in (0, 2, 4))
    return LACIVERT if 0.299 * r + 0.587 * g + 0.114 * b > 140 else "FFFFFF"


def _acik_ton(renk: str, beyaz: float = 0.72) -> str:
    r, g, b = (int(renk[i:i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{round(c + (255 - c) * beyaz):02X}" for c in (r, g, b))


def _bos(x) -> bool:
    return x is None or (isinstance(x, float) and math.isnan(x)) or (isinstance(x, str) and not x.strip())


_TR_ALFABE = "0123456789aâbcçdefgğhıiîjklmnoöprsştuûüvwxyz"


def _tr_anahtar(s) -> str:
    """Türkçe alfabetik sıralama anahtarı (Ç, Ğ, İ, Ö, Ş, Ü doğru yerde)."""
    s = "" if _bos(s) else str(s).replace("I", "ı").replace("İ", "i").lower()
    return "".join(chr(0x100 + _TR_ALFABE.index(c)) if c in _TR_ALFABE else c for c in s)


def _adres(cadde, sokak) -> str | None:
    c = "" if _bos(cadde) else str(cadde).strip()
    s = "" if _bos(sokak) else str(sokak).strip()
    if not c or not s:
        return (s or c) or None
    return s if tr_norm(c) == tr_norm(s) else f"{c}, {s}"


INCE = Side(style="thin", color=CIZGI)
BEYAZ_KALIN = Side(style="thick", color="FFFFFF")


def _yazi(sz: int, renk: str, kalin: bool = False) -> RichText:
    cp = CharacterProperties(sz=sz, b=kalin, solidFill=renk)
    return RichText(bodyPr=RichTextProperties(), p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _grafik_basligi(metin: str) -> Title:
    cp = CharacterProperties(sz=1100, b=True, solidFill=LACIVERT)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=metin)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def _seri_onbellegi(seri, degerler, kategoriler=None, bicim: str = "General") -> None:
    """Grafik serisine önbellek değerleri: hesap yapmayan görüntüleyicilerde de grafik dolu görünür."""
    seri.val.numRef.numCache = NumData(formatCode=bicim, ptCount=len(degerler),
                                       pt=[NumVal(idx=i, v=float(v)) for i, v in enumerate(degerler)])
    if kategoriler is not None:
        f = seri.cat.numRef.f if seri.cat.numRef is not None else seri.cat.strRef.f
        seri.cat = AxDataSource(strRef=StrRef(f=f, strCache=StrData(ptCount=len(kategoriler),
                                                                    pt=[StrVal(idx=i, v=str(k)) for i, k in enumerate(kategoriler)])))


def _cizgi_yerlestir(ws, grafik, kol0: int, satir0: int, kol1: int, satir1: int, off0: int = 0, off1: int = 0) -> None:
    """Grafiği iki hücre köşesi arasına yerleştirir (kolon/satır 0 tabanlı, ofset piksel)."""
    grafik.anchor = TwoCellAnchor(_from=AnchorMarker(col=kol0, colOff=pixels_to_EMU(off0), row=satir0, rowOff=0),
                                  to=AnchorMarker(col=kol1, colOff=pixels_to_EMU(off1), row=satir1, rowOff=0))
    ws.add_chart(grafik)


# ----------------------------------------------------------------------------- rapor
class _Rapor:
    def __init__(self, df: pd.DataFrame, sonuc, bolgeler: list[dict], cikti: Path):
        self.df = df
        self.s = sonuc
        self.cikti = cikti
        self.olcu = sonuc.olcu
        self.olcu_etiket = OLCU_ETIKET.get(self.olcu, self.olcu)
        self.n = int(getattr(sonuc, "n", 0) or max(b["bolge"] for b in bolgeler))
        self.tarih = dt.datetime.now()
        self.bolgeler = self._bolgeleri_tamamla(bolgeler)
        self.renk = [b["renk"] for b in self.bolgeler]
        self.wb = Workbook()
        self.onbellek: dict[str, dict[str, object]] = defaultdict(dict)
        self.veri = self._bina_tablosu()
        self.son_satir = len(self.veri) + 1          # BINA_ATAMA'daki son veri satırı
        self.ozet_hucre: dict[tuple[int, str], str] = {}   # (bölge, ölçü) -> OZET hücresi
        self.kpi: dict[str, str] = {}
        self.algoritma = str(sonuc.parametreler.get("algoritma") or ALGORITMA_SURUMU)
        # Fırsat tanımının kaynak veriyle çeliştiği binalar (fırsat negatife düşemez)
        abone, hp = df["aktif_res"], df["res_hp"]
        self.abone_fazla = int((abone > hp).sum())
        self.abone_fark = int((abone - hp).clip(lower=0).sum())
        self.hp_sifir = int(((hp == 0) & (abone > 0)).sum())
        self.mahalle_sayisi = int(df.groupby(["ilce", "mahalle"]).ngroups)
        self.firsat_notu = (f"{_tr(self.abone_fazla)} binada aktif abone sayısı RES HP'yi aşıyor; bu binalarda fırsat "
                            f"0 sayılır (toplam fark {_tr(self.abone_fark)} hane). Fırsat = Σ maks(RES HP − aktif abone, 0).")

    # --- veri hazırlığı -------------------------------------------------------
    def _bolgeleri_tamamla(self, bolgeler: list[dict]) -> list[dict]:
        yedek = varsayilan_renkler(self.n)
        sirali = {b["bolge"]: b for b in bolgeler}
        out = []
        for k in range(1, self.n + 1):
            b = dict(sirali.get(k, {"bolge": k, "bos": True}))
            b.setdefault("kod", f"B{k}")
            b.setdefault("ad", f"B{k} (boş bölge)")
            b.setdefault("kisa_ad", b["ad"])
            b["renk"] = _hex(b.get("renk") or yedek[k - 1])
            out.append(b)
        return out

    def _bina_tablosu(self) -> pd.DataFrame:
        d = self.df
        v = pd.DataFrame({
            "bolge": d["bina_serial"].map(self.s.atama).fillna(0).astype(int),
            "bina_serial": d["bina_serial"], "tellcordia_id": d["tellcordia_id"], "location_id": d["location_id"],
            "ad": _temiz_ad(d["ad"]), "site_adi": _temiz_ad(d["site_adi"]), "blok_adi": _temiz_ad(d["blok_adi"]),
            "kapi_no": d["kapi_no"],
            "mahalle": d["mahalle"], "adres": [_adres(c, s) for c, s in zip(d["cadde"], d["sokak"])],
            "ilce": d["ilce"], "il": d["il"], "obek": d["obek"], "lat": d["lat"], "lon": d["lon"], "kat": d["kat"],
            "toplam_hp": d["toplam_hp"], "res_hp": d["res_hp"], "aktif_res": d["aktif_res"], "firsat": d["firsat"],
            "penetrasyon": d["penetrasyon"],
        })
        # Penetrasyon en çok %100: aboneyi HP'yi aşan binalarda oran 1, HP'si 0 olanlarda boş
        v["penetrasyon"] = np.where(v["res_hp"] > 0, np.minimum(v["aktif_res"] / v["res_hp"].where(v["res_hp"] > 0, 1), 1.0), np.nan)
        v["not"] = np.where((v["res_hp"] == 0) & (v["aktif_res"] > 0), "HP 0, abone var",
                            np.where(v["aktif_res"] > v["res_hp"], "abone > HP", None))
        ad = {b["bolge"]: b["ad"] for b in self.bolgeler}
        v.insert(1, "bolge_adi", v["bolge"].map(ad).fillna("?"))
        v.insert(2, "sorumlu", "")
        # Adsız binalar mahallenin sonuna düşsün (ilk satırlar boş görünmesin)
        v["_k0"] = v["ad"].map(lambda x: 1 if _bos(x) else 0)
        v["_k1"], v["_k2"], v["_k3"] = v["ilce"].map(_tr_anahtar), v["mahalle"].map(_tr_anahtar), v["ad"].map(_tr_anahtar)
        v = v.sort_values(["bolge", "_k1", "_k2", "_k0", "_k3", "bina_serial"]).drop(columns=["_k0", "_k1", "_k2", "_k3"])
        return v.reset_index(drop=True)

    def _bolge_toplamlari(self) -> pd.DataFrame:
        g = self.veri.groupby("bolge")
        t = pd.DataFrame({"bina": g.size(), "res_hp": g["res_hp"].sum(), "toplam_hp": g["toplam_hp"].sum(),
                          "aktif_res": g["aktif_res"].sum(), "firsat": g["firsat"].sum()})
        return t.reindex(range(1, self.n + 1), fill_value=0)

    def _denge(self) -> tuple[float, float, int]:
        """(en düşük sapma, en yüksek sapma, en büyük mutlak fark) — BINA_ATAMA'daki veriden türetilir."""
        t = self._bolge_toplamlari()
        toplam = float(len(self.veri)) if self.olcu == "bina" else float(self.veri[self.olcu].sum())
        hedef = toplam / self.n
        deger = t["bina"].astype(float) if self.olcu == "bina" else t[self.olcu].astype(float)
        sap = (deger / hedef - 1).to_numpy()
        return float(sap.min()), float(sap.max()), int(math.ceil(float(np.abs(deger - hedef).max()) - 1e-9))

    def _denge_cumlesi(self) -> str:
        alt, ust, fark = self._denge()
        return (f"bütün bölgeler hedefin {_tr_isaretli(_yuzde_yuvarla(alt))} … {_tr_isaretli(_yuzde_yuvarla(ust))} "
                f"aralığında (hedeften en büyük sapma: {_tr(fark)} hane)")

    # --- hücre yardımcıları ---------------------------------------------------
    def formul(self, ws, ref: str, formul: str, deger, **stil):
        c = ws[ref]
        c.value = formul
        self.onbellek[ws.title][ref] = deger
        self._stil(c, **stil)
        return c

    def yaz(self, ws, ref: str, deger, **stil):
        c = ws[ref]
        c.value = deger
        self._stil(c, **stil)
        return c

    @staticmethod
    def _stil(c, font=None, fill=None, fmt=None, align=None, border=None):
        if font is not None:
            c.font = font
        if fill is not None:
            c.fill = fill
        if fmt is not None:
            c.number_format = fmt
        if align is not None:
            c.alignment = align
        if border is not None:
            c.border = border

    def bant(self, ws, son_kol: int, baslik: str, alt: str, ilk_kol: int = 2) -> None:
        """Sayfa üstü lacivert başlık bandı (satır 2–4) ve sarı vurgu çizgisi (satır 5)."""
        ws.row_dimensions[1].height = 8
        ws.row_dimensions[2].height = 18
        ws.row_dimensions[3].height = 32
        ws.row_dimensions[4].height = 20
        ws.row_dimensions[5].height = 4
        for k in range(ilk_kol, son_kol + 1):
            for r in (2, 3, 4):
                ws.cell(r, k).fill = _dolgu(LACIVERT)
            ws.cell(5, k).fill = _dolgu(VURGU)
        h = get_column_letter(ilk_kol)
        self.yaz(ws, f"{h}2", config.ORGANIZASYON, font=Font(size=9, bold=True, color=VURGU), align=Alignment(indent=1, vertical="bottom"))
        self.yaz(ws, f"{h}3", baslik, font=Font(size=20, bold=True, color="FFFFFF"), align=Alignment(indent=1, vertical="center"))
        self.yaz(ws, f"{h}4", alt, font=Font(size=11, color=BANT_ALT), align=Alignment(indent=1, vertical="top"))

    def baslik_satiri(self, ws, satir: int, kolonlar: list[str], ilk_kol: int = 2, yukseklik: float = 32) -> None:
        for i, metin in enumerate(kolonlar):
            self.yaz(ws, f"{get_column_letter(ilk_kol + i)}{satir}", metin,
                     font=Font(bold=True, color="FFFFFF", size=10), fill=_dolgu(LACIVERT),
                     align=Alignment(horizontal="center", vertical="center", wrap_text=True))
        ws.row_dimensions[satir].height = yukseklik

    def rozet(self, ws, ref: str, k: int, **ek):
        """Bölge rengiyle boyanmış, "B3" gösteren bölge hücresi (değer = bölge no)."""
        r = self.renk[k - 1]
        return self.yaz(ws, ref, k, font=Font(bold=True, color=_yazi_rengi(r)), fill=_dolgu(r), fmt=FMT_KOD,
                        align=Alignment(horizontal="center", vertical="center"), **ek)

    @staticmethod
    def genislik(ws, genislikler: dict[str, float]) -> None:
        for k, w in genislikler.items():
            ws.column_dimensions[k].width = w

    # --- sayfalar ---------------------------------------------------------------
    def olustur(self) -> Workbook:
        wb = self.wb
        ozet = wb.active
        ozet.title = "OZET"
        for ad in ["PERSONEL_ATAMA", "BOLGE_DETAY", "BINA_ATAMA", "ILCE_MATRIS", "MAHALLE_MATRIS",
                   "METODOLOJI", "VERI_KALITESI", "PARAMETRELER", "LISTE"]:
            wb.create_sheet(ad)
        self._adlar()
        self._liste()
        self._bina_atama()
        self._ozet()
        self._personel()
        self._bolge_detay()
        self._ilce_matris()
        self._mahalle_matris()
        self._metodoloji()
        self._veri_kalitesi()
        self._parametreler()

        renk = {"OZET": VURGU, "PERSONEL_ATAMA": VURGU, "BOLGE_DETAY": VURGU, "BINA_ATAMA": LACIVERT,
                "ILCE_MATRIS": LACIVERT, "MAHALLE_MATRIS": LACIVERT}
        for ws in wb.worksheets:
            ws.sheet_properties.tabColor = renk.get(ws.title, "9CA3AF")
            if ws.title not in ("BINA_ATAMA",):
                ws.sheet_view.showGridLines = False
        wb["LISTE"].sheet_state = "hidden"
        wb.active = 0
        p = wb.properties
        p.title = f"Bursa {self.n} Satışçı Bölgeleme"
        p.subject = f"{self.n} satış bölgesi · {self.olcu_etiket} dengeli"
        p.creator = config.ORGANIZASYON_KISA
        p.keywords = "fiber, satış, bölgeleme, Bursa"
        return wb

    def _adlar(self) -> None:
        n = self.n
        adlar = {ad: f"{quote_sheetname('BINA_ATAMA')}!${BK[k]}:${BK[k]}" for k, ad in ATAMA_ADLARI.items()}
        personel = len(self._personel_listesi())
        adlar.update({
            "PERSONEL": f"LISTE!$A$2:$A${personel + 1}",
            "BOLGE_ADLARI": f"LISTE!$E$2:$E${n + 1}",
            "SORUMLU_BOLGE": f"PERSONEL_ATAMA!$B$10:$B${9 + n}",
            "SORUMLULAR": f"PERSONEL_ATAMA!$H$10:$H${9 + n}",
        })
        for ad, ref in adlar.items():
            self.wb.defined_names.add(DefinedName(ad, attr_text=ref))

    @staticmethod
    def _personel_listesi() -> list[str]:
        e = config.SATIS_EKIBI
        return list(e["sorumlular"]) + [e["takim_lideri"]]

    def _olcu_formulu(self, olcu: str, kriter: str) -> str:
        if olcu == "bina":
            return f"COUNTIFS(ATAMA_BOLGE,{kriter})"
        return f"SUMIFS({ATAMA_ADLARI[olcu]},ATAMA_BOLGE,{kriter})"

    # LISTE (gizli): açılır liste, bölge adları, grafik etiketleri, hedef çizgisi
    def _liste(self) -> None:
        """A: personel (açılır liste) · C–I: bölge sözlüğü (E = BOLGE_ADLARI, H = grafik etiketi) · J: hedef çizgisi."""
        ws = self.wb["LISTE"]
        ws["A1"] = "Personel"
        for i, ad in enumerate(self._personel_listesi(), start=2):
            ws.cell(i, 1, ad)
        for j, b in {3: "No", 4: "Kod", 5: "Bölge adı", 6: "Kısa ad", 8: "Grafik etiketi", 9: "Renk", 10: "Hedef"}.items():
            ws.cell(1, j, b)
        for k, b in enumerate(self.bolgeler, start=1):
            r = k + 1
            ws.cell(r, 3, k)
            ws.cell(r, 4, b["kod"])
            ws.cell(r, 5, b["ad"])
            ws.cell(r, 6, b["kisa_ad"])
            ws.cell(r, 8, f"{b['kod']} · {b['kisa_ad']}" if self.n <= 12 else b["kod"])
            ws.cell(r, 9, "#" + b["renk"])
        self.genislik(ws, {"A": 28, "E": 40, "F": 22, "H": 28})

    # OZET ----------------------------------------------------------------------
    def _ozet(self) -> None:
        ws = self.wb["OZET"]
        n, t = self.n, self._bolge_toplamlari()
        kolonlar = [("bina", "Bina"), ("res_hp", "RES HP"), ("aktif_res", "Aktif Abone"), ("firsat", "Fırsat")]
        if self.olcu not in dict(kolonlar):
            kolonlar.append((self.olcu, self.olcu_etiket))
        kolonlar += [("penetrasyon", "Penetrasyon"), ("sapma", "Sapma\n(hedefe göre)")]
        harf = {k: get_column_letter(4 + i) for i, (k, _) in enumerate(kolonlar)}
        son_kol = 3 + len(kolonlar)
        self.genislik(ws, {"A": 2.5, "B": 8, "C": 42, **{get_column_letter(4 + i): 16 for i in range(len(kolonlar))}})

        toplam = {k: float(self.veri[k].sum()) for k in ("res_hp", "toplam_hp", "aktif_res", "firsat")}
        toplam["bina"] = float(len(self.veri))
        hedef = toplam[self.olcu] / n
        sapma = (t[self.olcu] / hedef - 1).tolist()
        self.bant(ws, son_kol, "Bursa Fiber Satış Bölgeleme",
                  f"{n} Satışçı  ·  {self.olcu_etiket} dengeli  ·  {_tr_tarih(self.tarih)}")

        # KPI kartları (satır 7–9) ------------------------------------------------
        il = self.veri["il"].value_counts()
        pen = toplam["aktif_res"] / max(toplam["res_hp"], 1)
        bina_alt = " · ".join(f"{i} {_tr(il[i])}" for i in il.index) if len(il) <= 3 else f"{self.veri['ilce'].nunique()} ilçe"
        tablo_ilk, tablo_son = 13, 12 + n
        sap = harf["sapma"]
        kartlar = [
            ("BİNA", "=COUNTA(ATAMA_SERIAL)-1", toplam["bina"], FMT_SAYI, bina_alt),
            ("RES HP", "=SUM(ATAMA_RES_HP)", toplam["res_hp"], FMT_SAYI, "konut home pass"),
            ("AKTİF ABONE", "=SUM(ATAMA_AKTIF)", toplam["aktif_res"], FMT_SAYI, f"penetrasyon {_tr_yuzde(pen, 1)}"),
            # Uzun açıklama 10. satırdaki dipnotta; karta sığan kısa etiket kullanılır (kısaltma okunmaz hâle geliyor)
            ("FIRSAT", "=SUM(ATAMA_FIRSAT)", toplam["firsat"], FMT_SAYI, "boş konut HP ⁽*⁾"),
            ("HEDEF / SATIŞÇI",
             f"=D8/{n}" if self.olcu == "bina" else f"=SUM({ATAMA_ADLARI[self.olcu]})/{n}", hedef, FMT_SAYI,
             f"{self.olcu_etiket} · {n} satışçı"),
            ("MAKS. SAPMA", f"=MAX(MAX({sap}{tablo_ilk}:{sap}{tablo_son}),-MIN({sap}{tablo_ilk}:{sap}{tablo_son}))",
             # önbellek, formülün göreceği değerle birebir aynı olmalı: I13:I20 ROUND(...,4)'lü
             max(abs(round(x, 4)) for x in sapma), "%0.00", "±%2 içi = yeşil"),
        ]
        ws.row_dimensions[6].height = 14
        ws.row_dimensions[7].height = 20
        ws.row_dimensions[8].height = 36
        ws.row_dimensions[9].height = 18
        for i, (etiket, f, deger, fmt, alt) in enumerate(kartlar):
            h = get_column_letter(4 + i)
            kenar = dict(left=BEYAZ_KALIN, right=BEYAZ_KALIN)
            self.yaz(ws, f"{h}7", etiket, font=Font(size=9, bold=True, color=GRI), fill=_dolgu(ZEMIN),
                     align=Alignment(horizontal="center", vertical="bottom"),
                     border=Border(top=Side(style="thick", color=VURGU if i < 4 else LACIVERT), **kenar))
            self.formul(ws, f"{h}8", f, deger, font=Font(size=20, bold=True, color=LACIVERT), fill=_dolgu(ZEMIN), fmt=fmt,
                        align=Alignment(horizontal="center", vertical="center"), border=Border(**kenar))
            self.yaz(ws, f"{h}9", alt, font=Font(size=8, color=GRI), fill=_dolgu(ZEMIN),
                     align=Alignment(horizontal="center", vertical="top", shrink_to_fit=True), border=Border(**kenar))
            self.kpi[etiket] = f"{h}8"
        # Sol blok: plan tanımı
        for r in (7, 8, 9):
            ws.merge_cells(f"B{r}:C{r}")
        kenar_sol = Border(left=Side(style="thick", color=VURGU))
        self.yaz(ws, "B7", "PLAN", font=Font(size=9, bold=True, color=GRI), align=Alignment(indent=1, vertical="bottom"), border=kenar_sol)
        self.yaz(ws, "B8", f"{n} bölge · {n} satışçı", font=Font(size=18, bold=True, color=LACIVERT),
                 align=Alignment(indent=1, vertical="center"), border=kenar_sol)
        d_alt, d_ust, d_fark = self._denge()
        self.yaz(ws, "B9", f"Sapma {_tr_isaretli(_yuzde_yuvarla(d_alt))} … {_tr_isaretli(_yuzde_yuvarla(d_ust))} "
                           f"· hedeften en büyük sapma {_tr(d_fark)} hane · site blokları bölünmez · bölgeler iç içe geçmez",
                 font=Font(size=8, color=GRI), align=Alignment(indent=1, vertical="top"), border=kenar_sol)
        # Fırsat dipnotu (KPI bandının altı)
        ws.row_dimensions[10].height = 13
        self.yaz(ws, "B10", "(*) " + self.firsat_notu, font=Font(size=8, italic=True, color=GRI),
                 align=Alignment(indent=1, vertical="center"))

        # Bölge tablosu -------------------------------------------------------------
        self.yaz(ws, "B11", "BÖLGELER", font=Font(size=11, bold=True, color=LACIVERT))
        self.yaz(ws, f"{get_column_letter(son_kol)}11", "BINA_ATAMA sayfasından canlı hesaplanır",
                 font=Font(size=8, italic=True, color=GRI), align=Alignment(horizontal="right"))
        self.baslik_satiri(ws, 12, ["Bölge", "Bölge Adı"] + [b for _, b in kolonlar])
        for k in range(1, n + 1):
            r = 12 + k
            b = self.bolgeler[k - 1]
            ws.row_dimensions[r].height = 20
            alt = Border(bottom=INCE)
            self.rozet(ws, f"B{r}", k, border=alt)
            self.yaz(ws, f"C{r}", b["ad"], font=Font(color=KOYU), align=Alignment(indent=1, vertical="center"), border=alt)
            for key, _ in kolonlar:
                h = harf[key]
                if key == "penetrasyon":
                    deger = t.at[k, "aktif_res"] / t.at[k, "res_hp"] if t.at[k, "res_hp"] else 0
                    f = f"=IF({harf['res_hp']}{r}>0,{harf['aktif_res']}{r}/{harf['res_hp']}{r},0)"
                    fmt = FMT_YUZDE
                elif key == "sapma":
                    deger = round(sapma[k - 1], 4)      # ±%0,00 gösterimi için biçimle aynı basamak
                    hd = self._abs(self.kpi["HEDEF / SATIŞÇI"])
                    f = f"=IF({hd}>0,ROUND({harf[self.olcu]}{r}/{hd}-1,4),0)"
                    fmt = FMT_SAPMA
                else:
                    deger = float(t.at[k, key])
                    f = "=" + self._olcu_formulu(key, f"$B{r}")
                    fmt = FMT_SAYI
                    self.ozet_hucre[(k, key)] = f"{h}{r}"
                self.formul(ws, f"{h}{r}", f, deger, fmt=fmt, border=alt, font=Font(color=KOYU, bold=key == "sapma"),
                            align=Alignment(horizontal="center" if key == "sapma" else "right", vertical="center", indent=0 if key == "sapma" else 1))
        # Toplam satırı
        rt = 13 + n
        ust = Border(top=Side(style="medium", color=LACIVERT))
        self.yaz(ws, f"B{rt}", None, fill=_dolgu(ZEMIN), border=ust)
        self.yaz(ws, f"C{rt}", "TOPLAM", font=Font(bold=True, color=LACIVERT), fill=_dolgu(ZEMIN), border=ust, align=Alignment(indent=1))
        for key, _ in kolonlar:
            h = harf[key]
            if key == "penetrasyon":
                f, deger, fmt = f"=IF({harf['res_hp']}{rt}>0,{harf['aktif_res']}{rt}/{harf['res_hp']}{rt},0)", pen, FMT_YUZDE
            elif key == "sapma":
                f, deger, fmt = None, None, None
            else:
                f, deger, fmt = f"=SUM({h}{tablo_ilk}:{h}{tablo_son})", float(t[key].sum()), FMT_SAYI
            if f:
                self.formul(ws, f"{h}{rt}", f, deger, font=Font(bold=True, color=LACIVERT), fill=_dolgu(ZEMIN), fmt=fmt, border=ust,
                            align=Alignment(horizontal="right", indent=1))
            else:
                self.yaz(ws, f"{h}{rt}", None, fill=_dolgu(ZEMIN), border=ust)
        ws.row_dimensions[rt].height = 22

        # Sapma renkleri: tabloda hücre hücre, KPI kartında dolgu tüm karta (yazı rengi yalnız sayıya)
        kademeler = [(f"<={YESIL_SINIR}", YESIL_Z, YESIL), (f"<={AMBER_SINIR}", AMBER_Z, AMBER),
                     (f">{AMBER_SINIR}", KIRMIZI_Z, KIRMIZI)]
        tablo_aralik = f"{sap}{tablo_ilk}:{sap}{tablo_son}"
        for j, (kosul, zemin, yazi) in enumerate(kademeler):
            ws.conditional_formatting.add(tablo_aralik, FormulaRule(
                formula=[f"ABS({sap}{tablo_ilk}){kosul}"], fill=_dolgu(zemin), font=Font(color=yazi, bold=True),
                stopIfTrue=j < 2))
        m = re.fullmatch(r"([A-Z]+)(\d+)", self.kpi["MAKS. SAPMA"])
        kart_kol, kart_satir = m.group(1), m.group(2)
        kart_kosul = f"ABS(${kart_kol}${kart_satir})"
        for j, (kosul, zemin, yazi) in enumerate(kademeler):
            ws.conditional_formatting.add(f"{kart_kol}7:{kart_kol}9", FormulaRule(
                formula=[f"{kart_kosul}{kosul}"], fill=_dolgu(zemin), stopIfTrue=j < 2))
            ws.conditional_formatting.add(f"{kart_kol}{kart_satir}", FormulaRule(
                formula=[f"{kart_kosul}{kosul}"], font=Font(color=yazi, bold=True), stopIfTrue=j < 2))

        # Lejant + tutarlılık uyarısı
        rl = rt + 2
        self.yaz(ws, f"C{rl}", "Sapma rengi:", font=Font(size=9, color=GRI), align=Alignment(horizontal="right", indent=1))
        for h, metin, z, y in ((harf["bina"], "±%2 içinde", YESIL_Z, YESIL), (harf["res_hp"], "%2 – %5", AMBER_Z, AMBER),
                               (harf["aktif_res"], "%5 üzeri", KIRMIZI_Z, KIRMIZI)):
            self.yaz(ws, f"{h}{rl}", metin, font=Font(size=9, bold=True, color=y), fill=_dolgu(z), align=Alignment(horizontal="center"))
        kontrol = f"{harf['bina']}{rt}"
        self.formul(ws, f"C{rl + 1}",
                    f'=IF({kontrol}<>{self.kpi["BİNA"]},"UYARI: "&({self.kpi["BİNA"]}-{kontrol})&" binanın bölge numarası geçersiz (1–{n} olmalı)","")',
                    "", font=Font(size=9, bold=True, color=KIRMIZI))
        self.yaz(ws, f"{harf['firsat']}{rl}", "Bir binanın bölgesini BINA_ATAMA'da değiştirin; bu tablo kendiliğinden güncellenir.",
                 font=Font(size=8, italic=True, color=GRI))

        # Grafikler ----------------------------------------------------------------
        liste = self.wb["LISTE"]
        for k in range(1, n + 1):
            self.formul(liste, f"J{k + 1}", f"=OZET!{self._abs(self.kpi['HEDEF / SATIŞÇI'])}", hedef)
        etiket_uzun = [liste.cell(k + 1, 8).value for k in range(1, n + 1)]
        kodlar = [b["kod"] for b in self.bolgeler]
        grafik_olculer = [self.olcu] + [k for k in ("firsat", "bina", "res_hp") if k != self.olcu][:2]
        g0 = rl + 3
        ana = self._cubuk_grafik(
            ws, grafik_olculer[0], f"{self.olcu_etiket} · bölge başına   (kırmızı çizgi: hedef {_tr(hedef)})",
            Reference(liste, min_col=8, min_row=2, max_row=n + 1), etiket_uzun, t[grafik_olculer[0]].tolist(), harf, tablo_ilk, tablo_son,
            hedef_ref=Reference(liste, min_col=10, min_row=2, max_row=n + 1), hedef=hedef)
        _cizgi_yerlestir(ws, ana, 1, g0 - 1, son_kol, g0 + 15)
        basliklar = {"firsat": "Fırsat (boş konut HP) · bölge başına", "bina": "Bina sayısı · bölge başına",
                     "res_hp": "RES HP · bölge başına", "toplam_hp": "Toplam HP · bölge başına"}
        g1 = g0 + 17
        # Yarım genişlikteki iki grafik: sütun piksel genişliklerinden orta nokta hesabı
        px = [self._kolon_px(ws, get_column_letter(c)) for c in range(2, son_kol + 1)]
        orta = sum(px) / 2
        kol, birikim = 2, 0.0
        while birikim + px[kol - 2] < orta:
            birikim += px[kol - 2]
            kol += 1
        ofset = int(orta - birikim)
        for j, olcu in enumerate(grafik_olculer[1:3]):
            g = self._cubuk_grafik(ws, olcu, basliklar[olcu], Reference(liste, min_col=4, min_row=2, max_row=n + 1), kodlar,
                                   t[olcu].tolist(), harf, tablo_ilk, tablo_son)
            if j == 0:
                _cizgi_yerlestir(ws, g, 1, g1 - 1, kol - 1, g1 + 15, 0, max(ofset - 6, 0))
            else:
                _cizgi_yerlestir(ws, g, kol - 1, g1 - 1, son_kol, g1 + 15, ofset + 6, 0)

        ws.print_area = f"A1:{get_column_letter(son_kol + 1)}{g1 + 16}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1      # tek sayfaya sığsın
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.ozet_harf = harf
        self.ozet_toplam_satiri = rt

    @staticmethod
    def _abs(ref: str) -> str:
        m = re.fullmatch(r"([A-Z]+)(\d+)", ref)
        return f"${m.group(1)}${m.group(2)}"

    @staticmethod
    def _kolon_px(ws, harf: str) -> float:
        w = ws.column_dimensions[harf].width or 8.43
        return math.trunc(w * 7 + 5)

    def _cubuk_grafik(self, ws, olcu, baslik, kat_ref, kategoriler, degerler, harf, ilk, son, hedef_ref=None, hedef=None):
        ch = BarChart()
        ch.type, ch.grouping, ch.gapWidth = "col", "clustered", 45
        ch.add_data(Reference(ws, min_col=openpyxl.utils.column_index_from_string(harf[olcu]), min_row=ilk, max_row=son), titles_from_data=False)
        ch.set_categories(kat_ref)
        s = ch.series[0]
        s.tx = SeriesLabel(v=OLCU_ETIKET.get(olcu, olcu))
        s.spPr = GraphicalProperties(solidFill=LACIVERT)
        for i, r in enumerate(self.renk):
            s.dPt.append(DataPoint(idx=i, spPr=GraphicalProperties(solidFill=r, ln=LineProperties(noFill=True))))
        if self.n <= 16:
            s.dLbls = DataLabelList()
            s.dLbls.showVal = True
            for a in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
                setattr(s.dLbls, a, False)
            s.dLbls.position = "outEnd"
            s.dLbls.numFmt = FMT_SAYI
            s.dLbls.txPr = _yazi(900 if self.n <= 10 else 700, KOYU, True)
        _seri_onbellegi(s, degerler, kategoriler, FMT_SAYI)
        ch.title = _grafik_basligi(baslik)
        ch.legend = None
        ch.roundedCorners = False
        ch.visible_cells_only = False
        ch.graphical_properties = GraphicalProperties(ln=LineProperties(solidFill=CIZGI))
        y, x = ch.y_axis, ch.x_axis
        y.delete = x.delete = False
        y.numFmt = NumFmt(formatCode=FMT_SAYI, sourceLinked=False)
        y.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="E5E7EB")))
        y.scaling.min = 0
        y.majorTickMark = x.majorTickMark = "none"
        y.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
        x.spPr = GraphicalProperties(ln=LineProperties(solidFill="9CA3AF"))
        y.txPr = _yazi(800, GRI)
        x.txPr = _yazi(900 if self.n <= 10 else 700, KOYU)
        if hedef_ref is not None:
            lc = LineChart()
            lc.add_data(hedef_ref, titles_from_data=False)
            ls = lc.series[0]
            ls.tx = SeriesLabel(v="Hedef")
            ls.spPr = GraphicalProperties(ln=LineProperties(solidFill=HEDEF_CIZGI, w=28575, prstDash="dash"))
            ls.marker = Marker(symbol="none")
            ls.smooth = False
            ls.val.numRef.numCache = NumData(formatCode=FMT_SAYI, ptCount=self.n, pt=[NumVal(idx=i, v=float(hedef)) for i in range(self.n)])
            lc.x_axis, lc.y_axis = ch.x_axis, ch.y_axis
            ch += lc
        return ch

    # PERSONEL_ATAMA ------------------------------------------------------------
    def _personel(self) -> None:
        ws = self.wb["PERSONEL_ATAMA"]
        n, t = self.n, self._bolge_toplamlari()
        e = config.SATIS_EKIBI
        self.genislik(ws, {"A": 2.5, "B": 8, "C": 40, "D": 12, "E": 12, "F": 10, "G": 13, "H": 34, "I": 32})
        self.bant(ws, 9, "Personel Ataması",
                  f"Satış Müdürü: {e['mudur']}   ·   Takım Lideri: {e['takim_lideri']}   ·   {len(e['sorumlular'])} satış sorumlusu")
        self.yaz(ws, "B7", "Her bölge için sarı hücreden satış sorumlusunu seçin. Otomatik atama yapılmamıştır; seçim satış müdürüne aittir.",
                 font=Font(size=10, color=KOYU))
        self.yaz(ws, "B8", "Seçilen isim BINA_ATAMA sayfasında o bölgedeki tüm binalara kendiliğinden yazılır. Aynı kişi iki bölgeye seçilirse satır turuncu olur.",
                 font=Font(size=9, italic=True, color=GRI))
        self.baslik_satiri(ws, 9, ["Bölge", "Bölge Adı", "RES HP", "Fırsat", "Bina", "Ofise Mesafe\n(km)", "Satış Sorumlusu  ▼", "Not"])
        alt = Border(bottom=INCE)
        girdi_kenar = Border(bottom=INCE, left=INCE, right=INCE)
        for k in range(1, n + 1):
            r = 9 + k
            b = self.bolgeler[k - 1]
            ws.row_dimensions[r].height = 24
            self.rozet(ws, f"B{r}", k, border=alt)
            self.yaz(ws, f"C{r}", b["ad"], align=Alignment(indent=1, vertical="center"), border=alt)
            for h, key in (("D", "res_hp"), ("E", "firsat"), ("F", "bina")):
                self.formul(ws, f"{h}{r}", f"=OZET!{self.ozet_hucre[(k, key)]}", float(t.at[k, key]), fmt=FMT_SAYI, border=alt,
                            align=Alignment(horizontal="right", vertical="center", indent=1))
            self.yaz(ws, f"G{r}", round(float(b.get("ofis_km", 0) or 0), 1), fmt=FMT_KM, border=alt,
                     align=Alignment(horizontal="center", vertical="center"))
            self.yaz(ws, f"H{r}", None, fill=_dolgu(GIRDI), border=girdi_kenar, font=Font(bold=True, color=LACIVERT),
                     align=Alignment(vertical="center", indent=1))
            self.yaz(ws, f"I{r}", None, border=alt, align=Alignment(vertical="center", wrap_text=True))
        ilk, son = 10, 9 + n
        dv = DataValidation(type="list", formula1="PERSONEL", allow_blank=True, showDropDown=False,
                            showErrorMessage=True, errorStyle="warning", errorTitle="Listede olmayan isim",
                            error="Bu isim satış ekibi listesinde yok. Yine de kullanmak istiyor musunuz?")
        dv.add(f"H{ilk}:H{son}")
        ws.add_data_validation(dv)
        ws.conditional_formatting.add(f"B{ilk}:I{son}", FormulaRule(
            formula=[f'AND($H{ilk}<>"",COUNTIF($H${ilk}:$H${son},$H{ilk})>1)'], fill=_dolgu(AMBER_Z), font=Font(bold=True, color=AMBER)))
        r = son + 2
        self.formul(ws, f"H{r}", f'=COUNTA(H{ilk}:H{son})&" / {n} bölge atandı"', f"0 / {n} bölge atandı",
                    font=Font(bold=True, color=LACIVERT), align=Alignment(indent=1))
        self.yaz(ws, f"C{r}", "Ofise mesafe: bölge merkezinin Nilüfer ofisine kuş uçuşu uzaklığı.", font=Font(size=8, italic=True, color=GRI))
        ws.freeze_panes = "A10"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    # BOLGE_DETAY ---------------------------------------------------------------
    def _bolge_detay(self) -> None:
        ws = self.wb["BOLGE_DETAY"]
        n = self.n
        self.genislik(ws, {"A": 2.5, "B": 8, "C": 30, "D": 11, "E": 11, "F": 11, "G": 11, "H": 10, "I": 9, "J": 9,
                           "K": 24, "L": 32, "M": 13})
        self.bant(ws, 13, "Bölge Detayı", "Coğrafya ve saha yükü göstergeleri  ·  plan anındaki atamaya göre")
        self.baslik_satiri(ws, 7, ["Bölge", "Bölge Adı", "Ofise\nMesafe (km)", "Ort. Merkeze\nUzaklık (km)", "En Uzak\nBina (km)",
                                   "Rota Yükü\n(km)", "Yerleşim\nalanı (km²)", "Mahalle", "Site /\nBirim", "İlçe Payları (RES HP)",
                                   "En Büyük 5 Mahalle · fırsat", "Bölge\nMerkezi"], yukseklik=40)
        alt = Border(bottom=INCE)
        ust = Alignment(vertical="top", wrap_text=True, indent=1)
        for k in range(1, n + 1):
            r, b = 7 + k, self.bolgeler[k - 1]
            if b.get("bos"):
                self.rozet(ws, f"B{r}", k, border=alt)
                self.yaz(ws, f"C{r}", b["ad"], border=alt)
                continue
            ilceler = b.get("ilceler", [])
            coklu = len(ilceler) > 1
            ilce_satir = [f"{i['ilce']}  {_tr_yuzde(i['pay']) if i['pay'] >= 0.005 else '<%1'}" for i in ilceler]
            mah = sorted(b.get("ust_mahalleler", []), key=lambda m: -m["firsat"])[:5]
            mah_satir = [f"{j}. {m['mahalle']}{' (' + m['ilce'] + ')' if coklu else ''}  ·  {_tr(m['firsat'])}" for j, m in enumerate(mah, 1)]
            ws.row_dimensions[r].height = max(len(ilce_satir), len(mah_satir), 2) * 13 + 8
            self.rozet(ws, f"B{r}", k, border=alt)
            self.yaz(ws, f"C{r}", b["ad"], border=alt, align=ust, font=Font(bold=True, color=KOYU))
            for h, key, fmt in (("D", "ofis_km", FMT_KM), ("E", "ort_merkez_km", "0.00"), ("F", "maks_merkez_km", FMT_KM),
                                ("G", "mst_km", FMT_KM), ("H", "alan_km2", FMT_KM)):
                v = b.get(key)
                self.yaz(ws, f"{h}{r}", None if v is None else round(float(v), 2), fmt=fmt, border=alt,
                         align=Alignment(horizontal="right", vertical="center", indent=1))
            self.yaz(ws, f"I{r}", b.get("mahalle_sayisi"), border=alt, align=Alignment(horizontal="center", vertical="center"))
            self.yaz(ws, f"J{r}", b.get("site"), fmt=FMT_SAYI, border=alt, align=Alignment(horizontal="center", vertical="center"))
            self.yaz(ws, f"K{r}", "\n".join(ilce_satir), border=alt, align=ust, font=Font(size=9, color=KOYU))
            self.yaz(ws, f"L{r}", "\n".join(mah_satir), border=alt, align=ust, font=Font(size=9, color=KOYU))
            lon, lat = b["merkez"]
            c = self.yaz(ws, f"M{r}", "Haritada aç", border=alt, font=Font(size=9, color="1D4ED8", underline="single"),
                         align=Alignment(horizontal="center", vertical="top"))
            c.hyperlink = f"https://www.google.com/maps/search/?api=1&query={lat:.5f},{lon:.5f}"
        son = 7 + n
        for h in "DEFGH":
            ws.conditional_formatting.add(f"{h}8:{h}{son}", ColorScaleRule(
                start_type="min", start_color="FFFFFF", end_type="max", end_color="C7D7F5"))
        notlar = [
            "Ofise mesafe: bölge merkezinin Nilüfer ofisine kuş uçuşu uzaklığı.",
            "Ort. merkeze uzaklık: binaların bölge merkezine RES HP ağırlıklı ortalama uzaklığı — küçükse bölge derli topludur.",
            "En uzak bina: bölge merkezine en uzak binanın kuş uçuşu mesafesi.",
            "Rota yükü: bölgedeki tüm site ve binaları birbirine bağlayan en kısa ağın uzunluğu — saha gezisi yükünün göstergesi.",
            "Yerleşim alanı: bölgedeki binaların kapladığı alan (bina tamponu, komşu bölgeyle paylaşılan sınırdan kırpılır) — idari yüzölçümü değildir.",
            "Mahalle: (ilçe, mahalle) çifti olarak sayılır; aynı mahalle adı farklı ilçelerde ayrı sayılır.",
            "Bu sayfa plan anındaki atamayı gösterir; BINA_ATAMA değişikliklerinde güncellenmez.",
        ]
        self.yaz(ws, f"B{son + 2}", "Açıklamalar", font=Font(bold=True, size=10, color=LACIVERT))
        for i, m in enumerate(notlar):
            self.yaz(ws, f"B{son + 3 + i}", m, font=Font(size=9, color=GRI))
        ws.freeze_panes = "D8"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    # BINA_ATAMA ------------------------------------------------------------------
    def _bina_atama(self) -> None:
        ws = self.wb["BINA_ATAMA"]
        v = self.veri
        son = self.son_satir
        for i, (key, baslik, w, _f) in enumerate(BINA_KOLONLARI, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.append([b for _, b, *_r in BINA_KOLONLARI])
        kayit = self.onbellek["BINA_ATAMA"]
        kolon = [k for k, *_r in BINA_KOLONLARI]
        veriler = v[kolon].astype(object).where(v[kolon].notna(), None)
        b_ad, s_ad = BK["bolge_adi"], BK["sorumlu"]
        for r, satir in enumerate(veriler.itertuples(index=False, name=None), start=2):
            satir = list(satir)
            for j, x in enumerate(satir):
                if isinstance(x, str) and x.startswith("="):
                    satir[j] = "'" + x      # metin formül sanılmasın
                elif isinstance(x, (np.integer,)):
                    satir[j] = int(x)
                elif isinstance(x, (np.floating,)):
                    satir[j] = None if np.isnan(x) else float(x)
            kayit[f"{b_ad}{r}"] = satir[1]
            kayit[f"{s_ad}{r}"] = ""
            satir[1] = f'=IFERROR(INDEX(BOLGE_ADLARI,$A{r}),"?")'
            satir[2] = f'=IFERROR(INDEX(SORUMLULAR,MATCH($A{r},SORUMLU_BOLGE,0))&"","")'
            ws.append(satir)
        # sayı biçimleri ve hizalama
        for i, (key, _b, _w, fmt) in enumerate(BINA_KOLONLARI, start=1):
            if fmt is None and key != "bolge":
                continue
            for (c,) in ws.iter_rows(min_row=2, max_row=son, min_col=i, max_col=i):
                c.number_format = fmt
                if key == "bolge":
                    c.alignment = Alignment(horizontal="center")
        baslik_font = Font(bold=True, color="FFFFFF", size=10)
        for c in ws[1]:
            c.font = baslik_font
            c.fill = _dolgu(LACIVERT)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        son_harf = get_column_letter(len(BINA_KOLONLARI))
        ws.freeze_panes = "E2"
        ws.auto_filter.ref = f"A1:{son_harf}{son}"
        ws.print_title_rows = "1:1"
        ws.page_setup.orientation = "landscape"
        # Bölge renkleri: 19 bin tek tek dolgu yerine bölge başına bir koşullu biçim kuralı
        for k in range(1, self.n + 1):
            r = self.renk[k - 1]
            ws.conditional_formatting.add(f"A2:A{son}", FormulaRule(formula=[f"$A2={k}"], fill=_dolgu(r), font=Font(bold=True, color=_yazi_rengi(r)), stopIfTrue=True))
            ws.conditional_formatting.add(f"B2:B{son}", FormulaRule(formula=[f"$A2={k}"], fill=_dolgu(_acik_ton(r)), stopIfTrue=True))
        ws.conditional_formatting.add(f"C2:{son_harf}{son}", FormulaRule(formula=["MOD(ROW(),2)=0"], fill=_dolgu(ZEBRA)))
        dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(self.n), showErrorMessage=True,
                            errorTitle="Geçersiz bölge", error=f"Bölge numarası 1 ile {self.n} arasında bir tam sayı olmalıdır.")
        dv.add(f"A2:A{son}")
        ws.add_data_validation(dv)
        self.yok_sayilan_hatalar = {"BINA_ATAMA": f"A1:{son_harf}{son}"}

    # Matrisler -------------------------------------------------------------------
    def _matris_basligi(self, ws, satir: int, ilk_kol: int, etiket_satiri: bool = True) -> None:
        for k in range(1, self.n + 1):
            h = get_column_letter(ilk_kol + k - 1)
            self.rozet(ws, f"{h}{satir}", k, border=Border(bottom=Side(style="thin", color="FFFFFF")))
            if etiket_satiri:
                self.yaz(ws, f"{h}{satir + 1}", self.bolgeler[k - 1]["kisa_ad"], font=Font(size=8, color=GRI),
                         align=Alignment(horizontal="center", vertical="top", wrap_text=True))

    def _ilce_matris(self) -> None:
        ws = self.wb["ILCE_MATRIS"]
        n, v = self.n, self.veri
        ilk_kol = 4
        son_kol = ilk_kol + n            # toplam kolonu
        self.genislik(ws, {"A": 2.5, "B": 9, "C": 18, **{get_column_letter(ilk_kol + i): 12 for i in range(n)},
                           get_column_letter(son_kol): 13})
        self.bant(ws, son_kol, "İlçe × Bölge", "Her ilçenin RES HP'si ve binaları hangi bölgelere düşüyor  ·  BINA_ATAMA'dan canlı")
        ilce_sira = (v.groupby(["il", "ilce"])["res_hp"].sum().reset_index()
                     .assign(_il=lambda x: x["il"].ne("Bursa")).sort_values(["_il", "res_hp"], ascending=[True, False]))
        satirlar = list(zip(ilce_sira["il"], ilce_sira["ilce"]))
        r = 7
        for baslik, key, fmt in ((f"RES HP", "res_hp", FMT_MATRIS), ("Bina sayısı", "bina", FMT_MATRIS)):
            pv = (v.assign(bina=1).pivot_table(index="ilce", columns="bolge", values=key, aggfunc="sum", fill_value=0)
                  .reindex(columns=range(1, n + 1), fill_value=0))
            self.yaz(ws, f"B{r}", baslik.upper(), font=Font(bold=True, size=11, color=LACIVERT))
            hr = r + 1
            self.yaz(ws, f"B{hr}", "İl", font=Font(bold=True, color="FFFFFF"), fill=_dolgu(LACIVERT), align=Alignment(horizontal="center", vertical="center"))
            self.yaz(ws, f"C{hr}", "İlçe", font=Font(bold=True, color="FFFFFF"), fill=_dolgu(LACIVERT), align=Alignment(indent=1, vertical="center"))
            self._matris_basligi(ws, hr, ilk_kol)
            self.yaz(ws, f"{get_column_letter(son_kol)}{hr}", "Toplam", font=Font(bold=True, color="FFFFFF"), fill=_dolgu(LACIVERT),
                     align=Alignment(horizontal="center", vertical="center"))
            ws.row_dimensions[hr].height = 22
            ws.row_dimensions[hr + 1].height = 24
            ilk_veri = hr + 2
            for i, (il, ilce) in enumerate(satirlar):
                rr = ilk_veri + i
                self.yaz(ws, f"B{rr}", il, font=Font(color=GRI, size=9), border=Border(bottom=INCE), align=Alignment(horizontal="center"))
                self.yaz(ws, f"C{rr}", ilce, font=Font(bold=True, color=KOYU), border=Border(bottom=INCE), align=Alignment(indent=1))
                for k in range(1, n + 1):
                    h = get_column_letter(ilk_kol + k - 1)
                    kriter = f"ATAMA_ILCE,$C{rr},ATAMA_BOLGE,{h}${hr}"
                    f = f"=COUNTIFS({kriter})" if key == "bina" else f"=SUMIFS(ATAMA_RES_HP,{kriter})"
                    self.formul(ws, f"{h}{rr}", f, float(pv.at[ilce, k]) if ilce in pv.index else 0.0, fmt=fmt, border=Border(bottom=INCE))
                h1, h2, ht = get_column_letter(ilk_kol), get_column_letter(son_kol - 1), get_column_letter(son_kol)
                self.formul(ws, f"{ht}{rr}", f"=SUM({h1}{rr}:{h2}{rr})", float(pv.loc[ilce].sum()) if ilce in pv.index else 0.0,
                            fmt=FMT_SAYI, font=Font(bold=True, color=LACIVERT), fill=_dolgu(ZEMIN), border=Border(bottom=INCE))
            son_veri = ilk_veri + len(satirlar) - 1
            rt = son_veri + 1
            ust = Border(top=Side(style="medium", color=LACIVERT))
            self.yaz(ws, f"B{rt}", None, fill=_dolgu(ZEMIN), border=ust)
            self.yaz(ws, f"C{rt}", "TOPLAM", font=Font(bold=True, color=LACIVERT), fill=_dolgu(ZEMIN), border=ust, align=Alignment(indent=1))
            for kk in range(ilk_kol, son_kol + 1):
                h = get_column_letter(kk)
                deger = float(pv.values.sum()) if kk == son_kol else float(pv[kk - ilk_kol + 1].sum())
                self.formul(ws, f"{h}{rt}", f"=SUM({h}{ilk_veri}:{h}{son_veri})", deger, fmt=FMT_SAYI,
                            font=Font(bold=True, color=LACIVERT), fill=_dolgu(ZEMIN), border=ust)
            aralik = f"{get_column_letter(ilk_kol)}{ilk_veri}:{get_column_letter(son_kol - 1)}{son_veri}"
            ws.conditional_formatting.add(aralik, ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                                 end_type="max", end_color="6F9BEA"))
            r = rt + 3
        ws.freeze_panes = "D7"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    def _mahalle_matris(self) -> None:
        ws = self.wb["MAHALLE_MATRIS"]
        n, v = self.n, self.veri
        ilk_kol = 9                        # I: B1
        son_kol = ilk_kol + n - 1
        self.genislik(ws, {"A": 14, "B": 24, "C": 12, "D": 8, "E": 9, "F": 9, "G": 10, "H": 12,
                           **{get_column_letter(ilk_kol + i): 10 for i in range(n)}})
        pv = (v.pivot_table(index=["ilce", "mahalle"], columns="bolge", values="res_hp", aggfunc="sum", fill_value=0)
              .reindex(columns=range(1, n + 1), fill_value=0))
        bina = v.groupby(["ilce", "mahalle"]).size()
        ilce_hp = v.groupby("ilce")["res_hp"].sum()
        sira = (pd.DataFrame({"res_hp": pv.sum(axis=1)}).reset_index()
                .assign(_i=lambda x: -x["ilce"].map(ilce_hp), _m=lambda x: -x["res_hp"])
                .sort_values(["_i", "ilce", "_m", "mahalle"]))
        # bölünme: RES HP'si iki veya daha fazla bölgeye dağılan mahalle
        bolge_sayisi = (pv > 0).sum(axis=1)
        bolunmus = int((bolge_sayisi >= 2).sum())

        self.yaz(ws, "A1", "Mahalle × Bölge  ·  RES HP", font=Font(bold=True, size=14, color=LACIVERT))
        hr, ilk = 4, 5
        son = ilk + len(sira) - 1
        self.formul(ws, "A2", f'={len(sira)}&" mahalle  ·  "&COUNTIF(H{ilk}:H{son},"Bölünmüş")&" mahalle birden çok bölgeye bölünmüş (turuncu satırlar)"',
                    f"{len(sira)} mahalle  ·  {bolunmus} mahalle birden çok bölgeye bölünmüş (turuncu satırlar)",
                    font=Font(size=10, color=GRI))
        basliklar = ["İlçe", "Mahalle", "Toplam\nRES HP", "Bina", "Bölge\nSayısı", "Ana\nBölge", "Ana Bölge\nPayı", "Durum"]
        for i, b in enumerate(basliklar, start=1):
            self.yaz(ws, f"{get_column_letter(i)}{hr}", b, font=Font(bold=True, color="FFFFFF", size=10), fill=_dolgu(LACIVERT),
                     align=Alignment(horizontal="center", vertical="center", wrap_text=True))
        self._matris_basligi(ws, hr, ilk_kol, etiket_satiri=False)
        for k in range(1, n + 1):   # kısa adlar başlığın üstünde (filtre satırı tek satır kalsın)
            self.yaz(ws, f"{get_column_letter(ilk_kol + k - 1)}{hr - 1}", self.bolgeler[k - 1]["kisa_ad"],
                     font=Font(size=7, color=GRI), align=Alignment(horizontal="center", vertical="bottom", wrap_text=True))
        ws.row_dimensions[hr - 1].height = 30
        ws.row_dimensions[hr].height = 32
        h1, h2 = get_column_letter(ilk_kol), get_column_letter(son_kol)
        # RES HP'si 0 olan mahalleler için bölge dağılımı bina sayısından okunur (satırlar boş kalmasın).
        # Bina sayıları gizli bir yardımcı blokta tutulur: dizi biçimli COUNTIFS her Excel sürümünde çalışmıyor.
        bina_pv = (v.assign(_b=1).pivot_table(index=["ilce", "mahalle"], columns="bolge", values="_b", aggfunc="sum", fill_value=0)
                   .reindex(columns=range(1, n + 1), fill_value=0))
        yrd_ilk = son_kol + 2
        y1, y2 = get_column_letter(yrd_ilk), get_column_letter(yrd_ilk + n - 1)
        self.yaz(ws, f"{y1}{hr}", "yardımcı: bina sayısı (RES HP = 0 satırları için)", font=Font(size=8, color=GRI))
        for k in range(n):
            ws.column_dimensions[get_column_letter(yrd_ilk + k)].hidden = True
            ws.column_dimensions[get_column_letter(yrd_ilk + k)].width = 8
        for i, (ilce, mahalle) in enumerate(zip(sira["ilce"], sira["mahalle"])):
            r = ilk + i
            deg = pv.loc[(ilce, mahalle)]
            bdeg = bina_pv.loc[(ilce, mahalle)]
            top = float(deg.sum())
            kaynak = deg if top > 0 else bdeg
            say = int((kaynak > 0).sum())
            ana = int(kaynak.idxmax()) if float(kaynak.sum()) > 0 else None
            self.yaz(ws, f"A{r}", ilce, font=Font(color=GRI))
            self.yaz(ws, f"B{r}", mahalle, font=Font(bold=True, color=KOYU))
            satir = f"{h1}{r}:{h2}{r}"
            basliklar_ref = f"${h1}${hr}:${h2}${hr}"
            bina_sayisi = int(bina.loc[(ilce, mahalle)])
            self.formul(ws, f"C{r}", f"=SUM({satir})", top, fmt=FMT_SAYI)
            self.yaz(ws, f"D{r}", bina_sayisi, fmt=FMT_SAYI)
            if top > 0:
                kaynak_satir = satir
            else:                         # yardımcı blok yalnız HP'si olmayan satırlarda doldurulur
                kaynak_satir = f"{y1}{r}:{y2}{r}"
                for k in range(1, n + 1):
                    h = get_column_letter(ilk_kol + k - 1)
                    self.formul(ws, f"{get_column_letter(yrd_ilk + k - 1)}{r}",
                                f"=COUNTIFS(ATAMA_ILCE,$A{r},ATAMA_MAHALLE,$B{r},ATAMA_BOLGE,{h}${hr})",
                                float(bdeg[k]), fmt=FMT_SAYI)
            self.formul(ws, f"E{r}", f'=COUNTIF({kaynak_satir},">0")', say, align=Alignment(horizontal="center"))
            self.formul(ws, f"F{r}",
                        f'=IFERROR(INDEX({basliklar_ref},MATCH(MAX({kaynak_satir}),{kaynak_satir},0)),"")',
                        ana if ana is not None else "", fmt=FMT_KOD, align=Alignment(horizontal="center"))
            self.formul(ws, f"G{r}",
                        f'=IF(C{r}>0,MAX({satir})/C{r},IF(D{r}>0,MAX({kaynak_satir})/D{r},""))',
                        float(deg.max() / top) if top > 0 else (float(bdeg.max()) / max(bina_sayisi, 1) if say else ""),
                        fmt="%0", align=Alignment(horizontal="center"))
            self.formul(ws, f"H{r}", f'=IF(E{r}>=2,"Bölünmüş",IF(C{r}=0,"HP yok",""))',
                        "Bölünmüş" if say >= 2 else ("HP yok" if top == 0 else ""), align=Alignment(horizontal="center"))
            for k in range(1, n + 1):
                h = get_column_letter(ilk_kol + k - 1)
                self.formul(ws, f"{h}{r}", f"=SUMIFS(ATAMA_RES_HP,ATAMA_ILCE,$A{r},ATAMA_MAHALLE,$B{r},ATAMA_BOLGE,{h}${hr})",
                            float(deg[k]), fmt=FMT_MATRIS)
        ws.conditional_formatting.add(f"A{ilk}:H{son}", FormulaRule(formula=[f"$E{ilk}>=2"], fill=_dolgu(AMBER_Z)))
        ws.conditional_formatting.add(f"H{ilk}:H{son}", FormulaRule(formula=[f"$E{ilk}>=2"], font=Font(bold=True, color=AMBER)))
        ws.conditional_formatting.add(f"{h1}{ilk}:{h2}{son}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                                          end_type="max", end_color="6F9BEA"))
        ws.freeze_panes = f"C{ilk}"
        ws.auto_filter.ref = f"A{hr}:{h2}{son}"
        ws.print_title_rows = f"{hr}:{hr}"
        ws.sheet_view.showGridLines = True

    # METODOLOJI -------------------------------------------------------------------
    def _metodoloji(self) -> None:
        ws = self.wb["METODOLOJI"]
        k = self._kalite()
        n = self.n
        es = k.get("eslesme", {})
        tekil = int(k.get("tekil_bina", len(self.veri)))
        eslesen = sum(v for key, v in es.items() if key != "YOK")
        hedef = self.s.hedef
        birim = self.s.parametreler.get("birim_sayisi", "")
        self.genislik(ws, {"A": 2.5, "B": 4, "C": 118})
        self.bant(ws, 3, "Bölgeler nasıl oluşturuldu?", "Tek sayfada yöntem  ·  ayrıntı gerekmez")
        bolumler = [
            ("Veri kaynakları", [
                f"data.xlsx — Superonline fiber bina listesi: {_tr(tekil)} tekil bina, HP ve aktif abone sayıları.",
                f"OneMap bina katmanı — her binanın konumu, taban poligonu, mahalle/sokak ve kat bilgisi.",
                f"Eşleşme {_tr_yuzde(eslesen / max(tekil, 1))}: " + " + ".join(f"{_tr(v)} {key}" for key, v in es.items() if key != "YOK")
                + " ile; konumu eksik bina yok.",
            ]),
            ("Tanımlar", [
                "RES HP: binadaki konut (residential) home pass sayısı — satılabilir konut bağlantısı.",
                "Fırsat = RES HP − aktif konut abonesi  (boş konut HP; negatif çıkarsa 0 kabul edilir).",
                f"{_tr(self.abone_fazla)} binada abone sayısı RES HP'yi aşıyor; oralarda fırsat 0 sayılır "
                f"(toplam fark {_tr(self.abone_fark)} hane).",
                "Penetrasyon = aktif konut abonesi ÷ RES HP  (en çok %100).",
                "Mahalle sayısı (ilçe, mahalle) çifti olarak sayılır: "
                f"{_tr(self.mahalle_sayisi)} mahalle, {self.veri['ilce'].nunique()} ilçe.",
            ]),
            ("Kurallar", [
                f"Her satışçıya eşit {self.olcu_etiket}: hedef {_tr(hedef)}; {self._denge_cumlesi()}.",
                "Site blokları bölünmez: aynı sitenin tüm blokları tek parça olarak aynı satışçıya gider.",
                "Bölgeler iç içe geçmez: bölge sınırları Voronoi mozaiğinden kırpılır, komşu bölgelerle ortak sınırı paylaşır.",
            ]),
            (f"Algoritma — dengeli güç diyagramı ({self.algoritma})", [
                f"1. Binalar site gruplarına toplanır ({_tr(birim) if birim else '—'} birim).",
                f"2. {n} başlangıç merkezi seçilir; birkaç farklı başlangıç denenir.",
                "3. Her birim en yakın merkeze bağlanır; her merkeze bir “çekim gücü” verilerek bölge yükleri hedefe eşitlenir.",
                "4. Merkezler bölgenin ağırlık merkezine kaydırılır ve 3. adım tekrarlanır (merkezler 5 m'den az oynayana kadar).",
                "5. Cila: sınır birimleri serbest bırakılıp tam bir taşıma problemi (doğrusal programlama) çözülür; "
                "kalan fark açgözlü sınır takasıyla kapatılır.",
                "6. Ada onarımı: bütün yakın komşuları başka bölgede kalan uzak birim o bölgeye verilir.",
                "7. Denemelerden en derli toplu ve dengeli olan seçilir. Aynı veriyle her zaman aynı sonuç çıkar.",
            ]),
            ("Yeniden bölme", [
                "Farklı satışçı sayısı için:   python bolge.py --n 10",
                "Ham veri güncellenince önce:   python -m dsale.enrich   (sonra yukarıdaki komut)",
                "Bu dosyada elle düzeltme: BINA_ATAMA'da binanın “Bölge” numarasını değiştirin; OZET kendiliğinden güncellenir.",
            ]),
        ]
        r = 7
        for baslik, maddeler in bolumler:
            self.yaz(ws, f"B{r}", baslik, font=Font(bold=True, size=12, color=LACIVERT),
                     border=Border(bottom=Side(style="thin", color=VURGU)))
            ws.cell(r, 3).border = Border(bottom=Side(style="thin", color=VURGU))
            ws.row_dimensions[r].height = 22
            r += 1
            for m in maddeler:
                komut = "python " in m
                self.yaz(ws, f"B{r}", "" if m[:2].rstrip(".").isdigit() else "•", font=Font(color=VURGU, bold=True),
                         align=Alignment(horizontal="center"))
                self.yaz(ws, f"C{r}", m, font=Font(size=10, color=KOYU, name="Consolas" if komut else "Calibri"),
                         align=Alignment(vertical="center"))
                ws.row_dimensions[r].height = 17
                r += 1
            r += 1
        ws.print_area = f"A1:C{r}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    # VERI_KALITESI ------------------------------------------------------------------
    @staticmethod
    def _kalite() -> dict:
        try:
            return json.load(open(config.QUALITY_JSON, encoding="utf-8"))
        except (OSError, ValueError):
            return {}

    def _veri_kalitesi(self) -> None:
        ws = self.wb["VERI_KALITESI"]
        k = self._kalite()
        self.genislik(ws, {"A": 2.5, "B": 36, "C": 16, "D": 13, "E": 100})
        self.bant(ws, 5, "Veri Kalitesi", "Kaynak veride bulunan sorunlar ve nasıl ele alındıkları")
        if not k:
            self.yaz(ws, "B7", f"{config.QUALITY_JSON.name} bulunamadı. Önce:  python -m dsale.enrich", font=Font(bold=True, color=KIRMIZI))
            return
        es = k.get("eslesme", {})
        tekil = int(k.get("tekil_bina", 0))
        tekrar = k.get("tekrar_eden_satir", [])
        tekrar_sayi = len(tekrar) - len({t["bina_serial"] for t in tekrar})
        bozuk = k.get("bozuk_tellcordia_id", [])
        uyusmaz = k.get("ilce_uyusmazligi", [])
        kat = k.get("kat_kaynagi", {})
        eslesen = sum(v for key, v in es.items() if key != "YOK")
        il = k.get("il_dagilimi", {})
        TAMAM, DUZ, BILGI = ("Tamam", YESIL_Z, YESIL), ("Düzeltildi", AMBER_Z, AMBER), ("Bilgi", "E5E7EB", "374151")
        tekrar_grup: dict[str, dict] = {}
        for t in tekrar:
            g = tekrar_grup.setdefault(t["bina_serial"], {"ad": " ".join(str(t.get("ad") or "").split()), "hp": []})
            g["hp"].append(_tr(t.get("res_hp") or 0))
        satirlar = [
            ("Kaynak satır (data.xlsx)", k.get("data_satir"), BILGI, f"{_tr(tekil)} tekil bina."),
            ("Tekrar eden satır", tekrar_sayi, DUZ if tekrar_sayi else TAMAM,
             "; ".join(f"{s} ({g['ad']}) {len(g['hp'])} kez var (RES HP {' / '.join(g['hp'])}) — HP'si yüksek kayıt tutuldu."
                       for s, g in tekrar_grup.items()) or "Yok."),
            ("Bozuk Tellcordia ID (bilimsel gösterim)", len(bozuk), DUZ if bozuk else TAMAM,
             "Excel'de sayıya dönüşmüş kimlik (ör. " + (bozuk[0]["tellcordia_id"] if bozuk else "") + "); bu binalar Bina Serial ile eşleştirildi."
             if bozuk else "Yok."),
            ("OneMap eşleşmesi", eslesen / max(tekil, 1), TAMAM if eslesen == tekil else DUZ,
             " + ".join(f"{_tr(v)} {key}" for key, v in es.items()) + " ile eşleşti."),
            ("Konumu eksik bina", k.get("konum_eksik", 0), TAMAM if not k.get("konum_eksik") else DUZ, "Tüm binalar OneMap bina poligonunun merkezine konumlandı."),
            ("İlçe uyuşmazlığı (CRM ↔ harita)", len(uyusmaz), DUZ if uyusmaz else TAMAM,
             "; ".join(f"{u['bina_serial']}: CRM {u['ilce_crm']}, harita {u['ilce']}" for u in uyusmaz) + " → harita esas alındı." if uyusmaz else "Yok."),
            ("Aktif abone > RES HP", self.abone_fazla, DUZ if self.abone_fazla else TAMAM,
             self.firsat_notu + f" Bunların {_tr(self.hp_sifir)} tanesinde RES HP 0 olduğu hâlde aktif abone var; "
                                "o binalarda penetrasyon boş bırakıldı (BINA_ATAMA → Not sütunu)."),
            ("Kat sayısı tahmini", kat.get("Tahmin", 0), BILGI,
             f"OneMap'te kat bilgisi olan {_tr(kat.get('OneMap', 0))} bina; eksikler konut sayısı ve taban alanından tahmin edildi "
             f"(yalnızca 3B görünüm için)."),
            ("Site grubu (bölünmeyen birim)", k.get("site_grup_sayisi"), BILGI, "Aynı sitenin blokları tek birim sayıldı; 350 m'den uzak bloklar ayrı grup."),
            ("Mahalle sayısı (ilçe × mahalle)", self.mahalle_sayisi, BILGI,
             f"Mahalle kimliği (ilçe, mahalle) çiftidir: {self.mahalle_sayisi} çift, "
             f"{self.veri['mahalle'].nunique()} farklı mahalle adı (aynı ad birden çok ilçede geçebilir). "
             "OneMap mahalle adları sadeleştirildi (Mh./Mah. ekleri kaldırıldı)."),
            ("İl dağılımı", sum(il.values()), BILGI, " · ".join(f"{a} {_tr(v)}" for a, v in il.items()) + " bina."),
        ]
        self.baslik_satiri(ws, 7, ["Kontrol", "Sonuç", "Durum", "Açıklama"], yukseklik=24)
        for i, (kontrol, sonuc, (durum, zemin, yazi), aciklama) in enumerate(satirlar):
            r = 8 + i
            alt = Border(bottom=INCE)
            ws.row_dimensions[r].height = 14 * math.ceil(len(aciklama) / 115) + 8
            self.yaz(ws, f"B{r}", kontrol, font=Font(bold=True, color=KOYU), border=alt, align=Alignment(indent=1, vertical="center"))
            fmt = "%0.0" if isinstance(sonuc, float) and sonuc <= 1 else FMT_SAYI
            self.yaz(ws, f"C{r}", sonuc, fmt=fmt, border=alt, font=Font(bold=True, color=LACIVERT),
                     align=Alignment(horizontal="right", vertical="center", indent=1))
            self.yaz(ws, f"D{r}", durum, fill=_dolgu(zemin), font=Font(bold=True, size=9, color=yazi), border=alt,
                     align=Alignment(horizontal="center", vertical="center"))
            self.yaz(ws, f"E{r}", aciklama, font=Font(size=9, color=KOYU), border=alt, align=Alignment(indent=1, vertical="center", wrap_text=True))
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        r = 8 + len(satirlar) + 1
        self.yaz(ws, f"B{r}", "Kaynak", font=Font(bold=True, size=9, color=GRI))
        self.yaz(ws, f"C{r}", f"OneMap çekim: {str(k.get('onemap_cekim_tarihi', ''))[:10]}  ·  {k.get('onemap_kaynak', '')}", font=Font(size=9, color=GRI))
        if bozuk or tekrar:
            r += 2
            self.yaz(ws, f"B{r}", "AYRINTI", font=Font(bold=True, size=10, color=LACIVERT))
            r += 1
            self.baslik_satiri(ws, r, ["Bina Serial", "Tür", "Değer", "Bina adı"], yukseklik=20)
            for x in bozuk:
                r += 1
                for h, val in zip("BCDE", (x["bina_serial"], "Bozuk ID", x["tellcordia_id"], x.get("ad"))):
                    self.yaz(ws, f"{h}{r}", val, font=Font(size=9), border=Border(bottom=INCE))
            for x in tekrar:
                r += 1
                for h, val in zip("BCDE", (x["bina_serial"], "Tekrar", f"RES HP {x.get('res_hp')}", x.get("ad"))):
                    self.yaz(ws, f"{h}{r}", val, font=Font(size=9), border=Border(bottom=INCE))

    # PARAMETRELER ------------------------------------------------------------------
    def _parametreler(self) -> None:
        ws = self.wb["PARAMETRELER"]
        s, n = self.s, self.n
        self.genislik(ws, {"A": 2.5, "B": 34, "C": 70})
        self.bant(ws, 3, "Parametreler", "Bu planı birebir yeniden üretmek için gereken her şey")
        sapmalar = self._denge()[:2]
        m = config.MASTER_CSV
        haric = s.parametreler.get("haric_il") or []
        komut = f"python bolge.py --n {n} --olcu {self.olcu}" + (" --haric-il " + " ".join(haric) if haric else "")
        satirlar: list[tuple[str, object, str | None]] = [
            ("Plan", self.cikti.name, None),
            ("Satışçı (bölge) sayısı", n, "0"),
            ("Denge ölçüsü", f"{self.olcu_etiket} ({self.olcu}) — {config.OLCULER.get(self.olcu, '')}", None),
            ("Hedef (bölge başına)", float(s.hedef), "#,##0.00"),
            ("En düşük sapma", round(float(min(sapmalar)), 4), FMT_SAPMA),
            ("En yüksek sapma", round(float(max(sapmalar)), 4), FMT_SAPMA),
            ("Hedeften en büyük sapma (hane)", self._denge()[2], FMT_SAYI),
            ("Algoritma", f"{self.algoritma} — kapasite kısıtlı güç diyagramı (dengeli k-means) + LP cilası", None),
        ]
        binlik = {"birim_sayisi"}            # seed binlik ayraçsız kalmalı (kopyalanabilir olsun)
        for key, val in s.parametreler.items():
            if key in ("algoritma", "olcu"):
                continue
            if isinstance(val, (list, tuple)):
                val = ", ".join(map(str, val)) or "—"
            sayi = isinstance(val, (int, np.integer)) and not isinstance(val, bool)
            satirlar.append((PARAM_ETIKET.get(key, key), val,
                             (FMT_SAYI if key in binlik else "0") if sayi else None))
        satirlar += [
            ("Çalışma süresi (sn)", float(s.sure_sn), "0.00"),
            ("Oluşturulma", self.tarih.strftime("%d.%m.%Y %H:%M"), None),
            ("Veri", f"{m.relative_to(config.ROOT).as_posix()}  ·  {_tr(len(self.df))} bina  ·  "
                     f"{dt.datetime.fromtimestamp(m.stat().st_mtime).strftime('%d.%m.%Y %H:%M')}" if m.exists() else str(m), None),
            ("Ofis konumu", f"{config.OFIS['lat']:.6f}, {config.OFIS['lon']:.6f}  ({config.OFIS.get('plus_code', '')})", None),
            ("Komut", komut, None),
            ("Yazılım", f"Python {sys.version.split()[0]} · pandas {pd.__version__} · openpyxl {openpyxl.__version__}", None),
        ]
        self.baslik_satiri(ws, 7, ["Parametre", "Değer"], yukseklik=22)
        for i, (ad, val, fmt) in enumerate(satirlar):
            r = 8 + i
            alt = Border(bottom=INCE)
            ws.row_dimensions[r].height = 18
            self.yaz(ws, f"B{r}", ad, font=Font(bold=True, color=KOYU), border=alt, align=Alignment(indent=1, vertical="center"))
            self.yaz(ws, f"C{r}", val, fmt=fmt, border=alt, font=Font(color=LACIVERT, name="Consolas" if ad == "Komut" else "Calibri"),
                     align=Alignment(horizontal="left", vertical="center", indent=1))
        ws.print_area = f"A1:C{7 + len(satirlar)}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True


# ----------------------------------------------------------------------------- önbellek (formül sonuçları)
_FORMUL_HUCRESI = re.compile(r'<c r="([A-Z]{1,3}[0-9]+)"([^>]*)><f>([^<]*)</f><v\s*/>')
_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
       "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
       "p": "http://schemas.openxmlformats.org/package/2006/relationships"}
_IGNORED_SONRASI = ("<smartTags", "<drawing", "<legacyDrawing", "<picture", "<oleObjects", "<controls",
                    "<webPublishItems", "<tableParts", "<extLst", "</worksheet>")


def _sayfa_yollari(dosyalar: dict[str, bytes]) -> dict[str, str]:
    wbx = ET.fromstring(dosyalar["xl/workbook.xml"])
    rels = ET.fromstring(dosyalar["xl/_rels/workbook.xml.rels"])
    hedef = {r.get("Id"): r.get("Target") for r in rels.findall("p:Relationship", _NS)}
    yollar = {}
    for s in wbx.find("m:sheets", _NS):
        t = hedef[s.get(f"{{{_NS['r']}}}id")]
        yollar[s.get("name")] = t.lstrip("/") if t.startswith("/") else "xl/" + t
    return yollar


def _onbellek_yaz(yol: Path, onbellek: dict[str, dict[str, object]], yok_sayilan: dict[str, str]) -> int:
    """Kaydedilmiş xlsx'te formül hücrelerine sonuç değerini, veri sayfasına 'metin olarak sayı' muafiyetini yazar."""
    with zipfile.ZipFile(yol) as z:
        bilgiler = z.infolist()
        dosyalar = {i.filename: z.read(i.filename) for i in bilgiler}
    yollar = _sayfa_yollari(dosyalar)
    yazilan = 0
    for sayfa in set(onbellek) | set(yok_sayilan):
        parca = yollar[sayfa]
        xml = dosyalar[parca].decode("utf-8")
        degerler = onbellek.get(sayfa, {})

        def degistir(m: re.Match) -> str:
            nonlocal yazilan
            ref, oz, f = m.groups()
            if ref not in degerler:
                return m.group(0)
            v = degerler[ref]
            if v is None or (isinstance(v, float) and not math.isfinite(v)):
                return m.group(0)
            oz = re.sub(r'\st="[^"]*"', "", oz)
            yazilan += 1
            if isinstance(v, str):
                return f'<c r="{ref}"{oz} t="str"><f>{f}</f><v>{escape(v)}</v>'
            if isinstance(v, (bool, np.bool_)):
                return f'<c r="{ref}"{oz} t="b"><f>{f}</f><v>{int(v)}</v>'
            v = float(v)
            return f'<c r="{ref}"{oz}><f>{f}</f><v>{int(v) if v.is_integer() else repr(v)}</v>'

        if degerler:
            xml = _FORMUL_HUCRESI.sub(degistir, xml)
        if sayfa in yok_sayilan:
            ek = f'<ignoredErrors><ignoredError sqref="{yok_sayilan[sayfa]}" numberStoredAsText="1"/></ignoredErrors>'
            konum = min(i for i in (xml.find(t) for t in _IGNORED_SONRASI) if i >= 0)
            xml = xml[:konum] + ek + xml[konum:]
        dosyalar[parca] = xml.encode("utf-8")
    gecici = yol.with_name(yol.stem + ".tmp.xlsx")
    with zipfile.ZipFile(gecici, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as z:
        for i in bilgiler:
            z.writestr(i.filename, dosyalar[i.filename])
    gecici.replace(yol)
    return yazilan


# ----------------------------------------------------------------------------- giriş noktası
def build(df: pd.DataFrame, sonuc, bolgeler: list[dict], cikti: Path) -> Path:
    """Excel raporunu üretir ve dosya yolunu döndürür (bolge.py çağırır)."""
    cikti = Path(cikti)
    cikti.mkdir(parents=True, exist_ok=True)
    r = _Rapor(df, sonuc, bolgeler, cikti)
    wb = r.olustur()
    yol = cikti / f"Bursa_{r.n}_Satisci_Bolgeleme.xlsx"
    wb.save(yol)
    beklenen = sum(len(v) for v in r.onbellek.values())
    yazilan = _onbellek_yaz(yol, r.onbellek, r.yok_sayilan_hatalar)
    if yazilan != beklenen:
        print(f"  ! Excel önbellek: {yazilan}/{beklenen} formül hücresine değer yazılabildi")
    return yol
