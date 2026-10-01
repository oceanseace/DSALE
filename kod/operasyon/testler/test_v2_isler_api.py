"""İş uçları, role göre kırpma ve KVKK (F3, F7, F21; spec §2.5, §5). Gerçek yetki (saha.yetki) + v2 yönlendiricileri."""
from __future__ import annotations

import io
import json

import pytest

import yollar
from operasyon.testler import yardim as y
from operasyon.v2 import akis, saklama

ORNEK = yollar.KOD / "operasyon" / "v2" / "ornek"


def _gorukle(v2db, conn, n=4):
    y.aktar(v2db, [y.satir(i, mahalle="Görükle") for i in range(1, n + 1)])
    return sorted(r[0] for r in conn.execute("SELECT is_no FROM is_emri"))


def test_obek_isleri_alanlar(istemci, jeton, v2db, conn, kisi, kim):
    """F3: öbeğe tıklayınca öbekteki işler — müşteri, kısa adres, task, durum, kalan süre, randevu, atanan."""
    nolar = _gorukle(v2db, conn)
    akis.ata(conn, nolar[0], kisi["teknik1"], kim("operasyon"),
             randevu={"bas": "2026-10-01 10:00:00", "bit": "2026-10-01 12:00:00", "teyitli": True})
    h = jeton(kisi["operasyon"])
    d = istemci.get("/api/isler", headers=h).json()
    oid = conn.execute("SELECT obek_id FROM is_emri WHERE is_no=?", (nolar[0],)).fetchone()[0]
    isler = [x for x in d["isler"] if (x["obek"] or {}).get("id") == oid]
    assert len(isler) == 4
    obekler = istemci.get("/api/obekler", headers=h).json()["obekler"]
    assert next(o for o in obekler if o["id"] == oid)["acik"] == 4          # satır sayısı = sayaç
    x = next(i for i in isler if i["is_no"] == nolar[0])
    assert x["musteri_adi"] == "Deneme Müşteri" and x["musteri_no"] and "Görükle" in x["kisa_adres"]
    assert x["task_adi"] == "Bağlantı Problemi" and x["durum"] == "atandi" and isinstance(x["kalan_dk"], int)
    assert x["randevu"] == {"bas": "2026-10-01 10:00:00", "bit": "2026-10-01 12:00:00", "teyitli": True, "kaynak": "biz"}
    assert x["atanan"]["id"] == kisi["teknik1"] and x["boss_ekip"] is None
    assert d["sayac"]["acik"] == 4 and d["hikaye"].startswith("Bugün 4 açık iş var")


def test_teknik_yalniz_kendi_isi_digerine_404(istemci, jeton, v2db, conn, kisi, kim):
    """F7: teknik yalnız kendisine atanmış işi görür; başkasınınki 404 (varlığı sızmaz)."""
    nolar = _gorukle(v2db, conn)
    op = kim("operasyon")
    akis.ata(conn, nolar[0], kisi["teknik1"], op)
    akis.ata(conn, nolar[1], kisi["teknik2"], op)
    h = jeton(kisi["teknik1"])
    d = istemci.get("/api/isler", headers=h).json()
    assert [x["is_no"] for x in d["isler"]] == [nolar[0]]
    assert istemci.get(f"/api/isler/{nolar[1]}", headers=h).status_code == 404
    r = istemci.post(f"/api/isler/{nolar[1]}/durum", headers=h, json={"yeni": "yolda", "istemci_id": "x1"})
    assert r.status_code == 404 and r.json()["kod"] == "is_yok"
    m = istemci.get("/api/islerim", headers=h).json()
    assert [x["is_no"] for x in m["isler"]] == [nolar[0]]
    assert istemci.get("/api/isler/excel", headers=h).status_code == 403
    assert istemci.get("/api/obekler", headers=h).status_code == 403


def test_musteri_alanlari_rol(istemci, jeton, v2db, conn, kisi, kim):
    """F21: tam bilgi operasyon/yöneticide; teknikte kendi AÇIK işinde kısa ad; kapanınca görmez; satış hiç."""
    nolar = _gorukle(v2db, conn, 2)
    akis.ata(conn, nolar[0], kisi["teknik1"], kim("operasyon"))
    for kid in (kisi["operasyon"], kisi["yonetici"]):
        a = istemci.get(f"/api/isler/{nolar[0]}", headers=jeton(kid)).json()
        assert a["musteri_adi"] == "Deneme Müşteri" and a["adres"] and a["musteri_no"]
    a = istemci.get(f"/api/isler/{nolar[0]}", headers=jeton(kisi["teknik1"])).json()
    assert a["musteri_adi"] == "Deneme M." and a["adres"] and "boss_ekip" not in a and "aramalar" not in a
    akis.gecis(conn, nolar[0], "cozuldu", kim("teknik1"), istemci_id="c1", evde_miydi=True)
    m = istemci.get("/api/islerim", headers=jeton(kisi["teknik1"])).json()
    biten = next(x for x in m["bitenler"] if x["is_no"] == nolar[0])
    assert not {"musteri_adi", "musteri_no", "kisa_adres"} & set(biten)       # anahtar HİÇ yok
    assert istemci.get("/api/isler", headers=jeton(kisi["satis"])).status_code == 403


def test_erisim_kaydi_gunde_bir(istemci, jeton, v2db, conn, kisi):
    nolar = _gorukle(v2db, conn, 2)
    h = jeton(kisi["operasyon"])
    for _ in range(3):
        istemci.get(f"/api/isler/{nolar[0]}", headers=h)
    istemci.get("/api/isler", headers=h)
    assert istemci.post(f"/api/isler/{nolar[0]}/kopya", headers=h, json={"alan": "musteri_no"}).status_code == 204
    say = dict(conn.execute("SELECT eylem, COUNT(*) FROM erisim_kaydi WHERE kullanici_id=? GROUP BY eylem",
                            (kisi["operasyon"],)).fetchall())
    assert say["is_ayrinti"] == 1 and say["is_liste"] == 1 and say["kopya_musteri_no"] == 1
    adet = conn.execute("SELECT adet FROM erisim_kaydi WHERE eylem='is_liste'").fetchone()[0]
    assert adet == 2


def test_saklama_30_gun_hmac(v2db, conn, saat):
    satirlar = [y.satir(1, musteri_no="933333333"), y.satir(2)]
    y.aktar(v2db, satirlar)
    y.aktar(v2db, satirlar[1:])                                              # 1. iş kapandı
    no = "400000001"
    ozet = conn.execute("SELECT musteri_ozet FROM is_emri WHERE is_no=?", (no,)).fetchone()[0]
    assert ozet and ozet != "933333333"
    saat.ilerlet(days=29)
    assert saklama.uygula(conn)["musteri_temizlenen"] == 0
    saat.ilerlet(days=2)
    assert saklama.uygula(conn)["musteri_temizlenen"] == 1
    r = conn.execute("SELECT musteri_adi, musteri_no, adres, musteri_ozet FROM is_emri WHERE is_no=?", (no,)).fetchone()
    assert tuple(r) == (None, None, None, ozet)                               # yalnız HMAC özeti kalır
    assert conn.execute("SELECT musteri_no FROM is_emri WHERE is_no='400000002'").fetchone()[0]   # açık iş dokunulmaz


def test_excel_rol_ve_kayit(istemci, jeton, v2db, conn, kisi):
    from openpyxl import load_workbook
    _gorukle(v2db, conn, 3)
    r = istemci.get("/api/isler/excel", headers=jeton(kisi["operasyon"]))
    assert r.status_code == 200 and r.headers["cache-control"] == "no-store"
    ws = load_workbook(io.BytesIO(r.content)).active
    basliklar = [c.value for c in ws[1]]
    assert "Müşteri Adı" in basliklar and ws.max_row == 4
    assert conn.execute("SELECT adet FROM erisim_kaydi WHERE eylem='excel'").fetchone()[0] == 3
    assert istemci.get("/api/isler/excel", headers=jeton(kisi["satis"])).status_code == 403


def test_degisim_imleci(istemci, jeton, v2db, conn, kisi, kim):
    nolar = _gorukle(v2db, conn, 3)
    h = jeton(kisi["operasyon"])
    imlec = istemci.get("/api/isler", headers=h).json()["imlec"]
    assert istemci.get(f"/api/isler/degisim?imlec={imlec}", headers=h).json()["isler"] == []
    akis.ata(conn, nolar[1], kisi["teknik1"], kim("operasyon"))
    d = istemci.get(f"/api/isler/degisim?imlec={imlec}", headers=h).json()
    assert [x["is_no"] for x in d["isler"]] == [nolar[1]] and d["imlec"] > imlec
    # teknikte başkasına geçen iş "görünmez" listesine düşer
    ht = jeton(kisi["teknik1"])
    im2 = istemci.get("/api/isler", headers=ht).json()["imlec"]
    akis.ata(conn, nolar[1], kisi["teknik2"], kim("operasyon"))
    d = istemci.get(f"/api/isler/degisim?imlec={im2}", headers=ht).json()
    assert d["gorunmez"] == [nolar[1]] and d["isler"] == []


def test_eski_uclar_410(monkeypatch, v2db):
    """Eski İş emirleri ekranı kaldırılınca (ya da SAHA_ESKI_IS_EMRI=0) /api/is-emri/* 410 'yenilendi'."""
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from operasyon.testler.conftest import uygulama
    from operasyon.v2 import api as v2api
    from saha import yetki
    monkeypatch.setenv("SAHA_ESKI_IS_EMRI", "0")
    app = uygulama()
    for r in v2api.eski_yonlendirici(yetki.Bagimliliklar()):
        app.include_router(r)
    with TestClient(app) as c:
        for yontem, yol in (("GET", "/api/is-emri"), ("POST", "/api/is-emri/yukle"), ("DELETE", "/api/is-emri/x")):
            r = c.request(yontem, yol)
            assert r.status_code == 410 and r.json()["kod"] == "yenilendi"
    assert isinstance(FastAPI, type)


def _anahtarlar(ornek, gercek, yol, eksik: list, izinli_fazla=frozenset()):
    """Örnekteki her anahtar gerçek yanıtta da olmalı (iç içe sözlüklerde ve listelerin ilk öğesinde)."""
    if isinstance(ornek, dict) and isinstance(gercek, dict):
        for k in ornek:
            if k not in gercek:
                eksik.append(f"{yol}.{k}")
            else:
                _anahtarlar(ornek[k], gercek[k], f"{yol}.{k}", eksik, izinli_fazla)
        for k in gercek:
            if k not in ornek and k not in izinli_fazla:
                eksik.append(f"{yol}.{k} (örnekte yok)")
    elif isinstance(ornek, list) and isinstance(gercek, list) and ornek and gercek \
            and isinstance(ornek[0], dict) and isinstance(gercek[0], dict):
        _anahtarlar(ornek[0], gercek[0], f"{yol}[0]", eksik, izinli_fazla)


def test_ornek_yanitlar_sema_uyumu(istemci, jeton, v2db, conn, kisi, kim):
    """operasyon/v2/ornek/*.json (sahte.ts'ten) gerçek yanıtlarla aynı şekilde: arayüz sahte kipten gerçeğe geçince kırılmaz."""
    nolar = _gorukle(v2db, conn, 4)
    op = kim("operasyon")
    akis.ata(conn, nolar[0], kisi["teknik1"], op)
    akis.gecis(conn, nolar[1], "askida", op, neden="abone", uyanma="2026-10-02 10:00:00")
    h = jeton(kisi["operasyon"])
    dizin = json.loads((ORNEK / "_dizin.json").read_text(encoding="utf-8"))
    uclar = {
        "GET /api/isler": "/api/isler", "GET /api/isler/{is_no}": f"/api/isler/{nolar[1]}",
        "GET /api/aranacaklar": "/api/aranacaklar", "GET /api/obekler": "/api/obekler",
        "GET /api/isler/ayarlar": "/api/isler/ayarlar", "GET /api/isler/kurallar": "/api/isler/kurallar",
        "GET /api/isler/teknikler": "/api/isler/teknikler", "GET /api/boss/giden": "/api/boss/giden",
        "GET /api/takip": "/api/takip", "GET /api/kesinti": "/api/kesinti", "GET /api/kesinti/aday": "/api/kesinti/aday",
        "GET /api/ilceler": "/api/ilceler", "GET /api/isler/boss-ekip": "/api/isler/boss-ekip",
    }
    # Role/duruma bağlı, örnekte olmayabilen ek alanlar (sözleşmede isteğe bağlı: tipler.ts '?')
    fazla = frozenset({"calisiyor", "klasorler_varsayilan", "yontem", "fark", "btk_durduran_aski", "teknik_kapasite",
                       "aski_nedeni", "uyanma", "talep_ulasamama_sms", "musteri_tel"})
    eksik: list[str] = []
    for uc, yol in uclar.items():
        r = istemci.get(yol, headers=h)
        assert r.status_code == 200, (uc, r.status_code, r.text[:200])
        ornek = json.loads((ORNEK / dizin[uc]).read_text(encoding="utf-8"))
        _anahtarlar(ornek, r.json(), uc, eksik, fazla)
    assert eksik == []


def test_boss_giden_ek4(istemci, jeton, v2db, conn, kisi, kim):
    """EK-4: eklentiye hazır giden kutusu — BOSS'a birebir yazılacak değerler; işlendi → kutudan çıkar; rapor doğrular."""
    satirlar = [y.satir(i, mahalle="Görükle") for i in range(1, 3)]
    y.aktar(v2db, satirlar)
    no = "400000001"
    akis.ata(conn, no, kisi["teknik1"], kim("operasyon"),
             randevu={"bas": "2026-10-01 10:00:00", "bit": "2026-10-01 12:00:00", "teyitli": True})
    h = jeton(kisi["operasyon"])
    g = istemci.get("/api/boss/giden", headers=h).json()
    assert [x["is_no"] for x in g] == [no]
    assert g[0]["alanlar"] == {"ekip": "ALI DENEME EKIBI", "randevu_baslangic": "2026-10-01 10:00:00",
                               "randevu_bitis": "2026-10-01 12:00:00"}
    assert g[0]["neden"] == "ekip,randevu" and g[0]["boss_task_no"] == no and g[0]["olusma"]
    assert istemci.get("/api/boss/giden", headers=jeton(kisi["teknik1"])).status_code == 403
    r = istemci.post(f"/api/boss/giden/{no}/islendi", headers=h)
    assert r.status_code == 200 and r.json()["boss"]["islendi"]
    assert istemci.get("/api/boss/giden", headers=h).json() == []
    # sonraki rapor BOSS'ta aynı ekip ve randevuyu gösterir → bayrak kalkar, "BOSS'ta doğrulandı"
    satirlar[0].update({"Ekip": "ALI DENEME EKIBI", "Randevu Başlangıç Tarihi": "2026-10-01 10:00:00",
                        "Randevu Bitiş Tarihi": "2026-10-01 12:00:00"})
    y.aktar(v2db, satirlar)
    r = conn.execute("SELECT boss_bekleyen FROM is_emri WHERE is_no=?", (no,)).fetchone()
    assert r[0] is None
    assert "BOSS'ta doğrulandı" in conn.execute("SELECT notu FROM is_emri_olay WHERE is_no=? AND tur='aktarim_degisti' "
                                                 "ORDER BY id DESC", (no,)).fetchone()[0]


def test_kurallar_uclari(istemci, jeton, v2db, conn, kisi):
    y.aktar(v2db, [y.satir(1, task="Modem Değişikliği", baslangic="2026-09-30 09:00:00")])
    assert conn.execute("SELECT son24 FROM is_emri").fetchone()[0] == "2026-10-07 09:00:00"     # EK-12.1: 7 gün
    k = istemci.get("/api/isler/kurallar", headers=jeton(kisi["operasyon"])).json()
    assert k["hedef_saatleri"]["teyit_bekliyor"] and k["arama_merdiveni"]["deger"]["ara_saat"] == 3
    yeni = {"hedef_saatleri": {"varsayilan": 24, "kurallar": [{"desen": "MODEM DEGISIKLIGI", "saat": 72}]}}
    assert istemci.put("/api/isler/kurallar", headers=jeton(kisi["operasyon"]), json=yeni).status_code == 403
    bozuk = istemci.put("/api/isler/kurallar", headers=jeton(kisi["yonetici"]), json={"serit_kurallari": [{"desen": ""}]})
    assert bozuk.status_code == 422 and bozuk.json()["kod"] == "alan_eksik"
    r = istemci.put("/api/isler/kurallar", headers=jeton(kisi["yonetici"]), json=yeni)
    assert r.status_code == 200 and r.json()["hedef_saatleri"]["degisti"]
    assert conn.execute("SELECT son24 FROM is_emri").fetchone()[0] == "2026-10-03 09:00:00"     # açık işe uygulandı
    r = istemci.put("/api/isler/kurallar", headers=jeton(kisi["yonetici"]), json={"hedef_saatleri": None})
    assert not r.json()["hedef_saatleri"]["degisti"]
    assert conn.execute("SELECT COUNT(*) FROM yonetim_kaydi WHERE hedef='is_kurallari'").fetchone()[0] == 2


@pytest.mark.parametrize("yol", ["/api/isler/kurallar", "/api/kesinti", "/api/kesinti/aday"])
def test_yeni_uclar_satisciya_kapali(istemci, jeton, v2db, kisi, yol):
    assert istemci.get(yol, headers=jeton(kisi["satis"])).status_code == 403
