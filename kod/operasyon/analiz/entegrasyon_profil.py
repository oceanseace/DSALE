"""Entegrasyon profili — kaynak dışa aktarımlarının anahtar, biçim ve eşleşme ölçümleri.

    .venv/Scripts/python.exe operasyon/analiz/entegrasyon_profil.py

Okur (salt okunur): BOSS TeknikTaskDetayRaporu.xlsx, Fox ReportResultAçık/Askı.xls,
PS26/data.xlsx (ORIGN, TICKET, GUZERGAH, ALTYAPI), veri/raw/data.xlsx (önceki tur raporu),
veri/raw/onemap_bina_bursa.json, veri/master/bina_master.csv.

Yazar: operasyon/analiz/cikti/entegrasyon_profil.json

GİZLİLİK: çıktıya yalnız sayım, oran, kategori, ilçe ve süre yazılır. Müşteri adı, adres,
telefon, müşteri no, cihaz seri no, teknisyen/satıcı adı ÇIKTIYA GİRMEZ. Ağ çağrısı yoktur.
"""
from __future__ import annotations

import json
import re
import sys
import warnings
from pathlib import Path

import numpy as np
import pandas as pd

warnings.filterwarnings("ignore")

if __package__ in (None, ""):  # doğrudan çalıştırıldı: kod/ içe aktarma yoluna (yollar)
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import yollar  # noqa: E402
CIKTI = Path(__file__).resolve().parent / "cikti"
MASAUSTU = Path(r"C:/Users/EXT03426951/Desktop")
BOSS_XLSX = MASAUSTU / "TeknikTaskDetayRaporu.xlsx"
FOX_ACIK = MASAUSTU / "ReportResultAçık.xls"
FOX_ASKI = MASAUSTU / "ReportResultAskı.xls"
TUR_YENI = yollar.PS26 / "data.xlsx"
TUR_ESKI = yollar.VERI / "raw" / "data.xlsx"
ONEMAP = yollar.VERI / "raw" / "onemap_bina_bursa.json"
MASTER = yollar.VERI / "master" / "bina_master.csv"
PS26_DATA = yollar.PS26 / "data"

# Kişisel veri taşıyan kolonlar (yalnız sınıflandırma; değerleri okunup yazılmaz)
KISISEL = {
    "BOSS": ["Müşteri No", "Müşteri Adı", "Adres", "Cihaz Seri No", "Cihaz Seri No 2", "Superbox GSM No",
             "Satış Temsilcisi Ad Soyad", "Satış Temsilcisi Kullanıcı Kodu", "Ekip", "Teknisyen", "2. Teknisyen",
             "En Son İşlem Yapan Kullanıcı", "Son Açıklama"],
    "FOX": ["Müşteri No", "Akışı Başlatan", "Üstlenen Kodu - Adı"],
}

KURULUM_RE = re.compile(r"kurulum", re.IGNORECASE)
TEL_RE = re.compile(r"(?<!\d)(?:\+?90|0)?\s?5\d{2}\s?\d{3}\s?\d{2}\s?\d{2}(?!\d)")


def yuvarla(x, n=3):
    return None if x is None or (isinstance(x, float) and np.isnan(x)) else round(float(x), n)


def dagilim(s: pd.Series, n: int = 40) -> dict:
    return {str(k): int(v) for k, v in s.value_counts(dropna=False).head(n).items()}


def yuzdelik(s: pd.Series) -> dict:
    s = pd.to_numeric(s, errors="coerce").dropna()
    if s.empty:
        return {}
    q = s.quantile([0.5, 0.9, 0.95, 0.99])
    return {"adet": int(len(s)), "medyan": yuvarla(q[0.5], 2), "p90": yuvarla(q[0.9], 2),
            "p95": yuvarla(q[0.95], 2), "p99": yuvarla(q[0.99], 2), "maks": yuvarla(s.max(), 2)}


def kurulum_mu(ad: str) -> bool:
    """Kullanıcı kuralı: adında 'kurulum' geçen = yeni müşteri ekibi; 'Kurulum Taskı Ürememiş' hariç."""
    ad = str(ad or "")
    return bool(KURULUM_RE.search(ad.replace("İ", "i").lower())) and "ürememiş" not in ad.lower()


# ----------------------------------------------------------------------------- okuma
def oku():
    boss = pd.read_excel(BOSS_XLSX, dtype={"Lokasyon": str})
    fa = pd.read_excel(FOX_ACIK)
    fa["_dosya"] = "acik"
    fs = pd.read_excel(FOX_ASKI)
    fs["_dosya"] = "aski"
    fox = pd.concat([fa, fs], ignore_index=True)
    master = pd.read_csv(MASTER, dtype=str, encoding="utf-8-sig")
    return boss, fox, master


def _kod(s: pd.Series) -> pd.Series:
    s = s.astype(object).where(s.notna(), "")
    return s.astype(str).str.strip().str.replace(r"\.0$", "", regex=True).replace({"nan": np.nan, "None": np.nan, "": np.nan})


def lokasyon_norm(s: pd.Series) -> pd.Series:
    """Bina lokasyon anahtarı: 'O…' olduğu gibi; yalnız rakamsa 8 haneye sıfır doldurulur."""
    s = _kod(s)
    rakam = s.str.fullmatch(r"\d+").fillna(False).astype(bool)
    return s.where(~rakam, s.str.zfill(8))


# ----------------------------------------------------------------------------- BOSS
def boss_profili(boss: pd.DataFrame, master: pd.DataFrame) -> dict:
    p: dict = {"satir": int(len(boss)), "kolon": int(boss.shape[1]), "kolonlar": list(boss.columns),
               "task_no_tekil": bool(boss["Task No"].is_unique),
               "kisisel_kolonlar": [c for c in KISISEL["BOSS"] if c in boss.columns]}
    p["bos_oran"] = {c: yuvarla(boss[c].isna().mean()) for c in boss.columns}
    p["tamamen_bos_kolonlar"] = [c for c in boss.columns if boss[c].isna().all()]
    p["sabit_kolonlar"] = [c for c in boss.columns if boss[c].nunique(dropna=True) == 1 and not boss[c].isna().all()]
    p["zaman_araligi"] = {"ilk_baslangic": str(boss["Task Başlangıç Tarihi"].min()),
                          "son_baslangic": str(boss["Task Başlangıç Tarihi"].max())}
    for c in ["Task Durumu", "Randevu Durumu", "Askıya Alınma Nedeni", "Merkeze Gönder Statüsü", "SL", "Segment",
              "Ürün Bilgisi", "Web/Mobil", "İl"]:
        p.setdefault("dagilim", {})[c] = dagilim(boss[c])
    p["dagilim"]["Task Adı"] = dagilim(boss["Task Adı"], 60)
    p["ekip_sayisi_tekil"] = int(boss["Ekip"].nunique())
    p["teknisyen_sayisi_tekil"] = int(boss["Teknisyen"].nunique())

    # Kurulum / mevcut müşteri ayrımı
    kur = boss["Task Adı"].map(kurulum_mu)
    p["kurulum_ayrimi"] = {"kurulum_yeni_musteri": int(kur.sum()), "mevcut_musteri": int((~kur).sum())}

    # Lokasyon → bina eşleşmesi
    lok = _kod(boss["Lokasyon"])
    desen = lok.dropna().str.replace(r"\d", "9", regex=True)
    loc_set = set(_kod(master["location_id"]).dropna())
    serial_set = set(master["bina_serial"].dropna())
    tell_set = set(_kod(master["tellcordia_id"]).dropna())
    eslesti = lok.isin(loc_set)
    p["lokasyon"] = {
        "dolu": int(lok.notna().sum()), "bos": int(lok.isna().sum()), "tekil": int(lok.nunique()),
        "desen": {k: int(v) for k, v in desen.value_counts().items()},
        "master_location_id_eslesen_satir": int(eslesti.sum()),
        "master_location_id_eslesme_orani_dolularda": yuvarla(eslesti[lok.notna()].mean()),
        "master_location_id_eslesme_orani_tum": yuvarla(eslesti.mean()),
        "bina_serial_ile_eslesen": int(lok.isin(serial_set).sum()),
        "tellcordia_ile_eslesen": int(lok.isin(tell_set).sum()),
        "eslesen_tekil_bina": int(lok[eslesti].nunique()),
        "desene_gore_eslesme": {
            k: {"satir": int(len(g)), "eslesen": int(g.isin(loc_set).sum())}
            for k, g in lok.dropna().groupby(desen)
        },
        "urune_gore_bos_lokasyon": {str(k): int(v) for k, v in boss.loc[lok.isna(), "Ürün Bilgisi"].fillna("(boş)").value_counts().items()},
        "urune_gore_eslesme_orani": {str(k): yuvarla(g.mean()) for k, g in eslesti.groupby(boss["Ürün Bilgisi"].fillna("(boş)"))},
        "kategoriye_gore_eslesme_orani": {
            "kurulum": yuvarla(eslesti[kur].mean()), "mevcut_musteri": yuvarla(eslesti[~kur].mean())},
        "il_disi_eslesmeyen_dolu": int((lok.notna() & ~eslesti & ~boss["İl"].isin(["Bursa", "Yalova"])).sum()),
    }
    # Normalizasyon: tur raporu Excel'i sayısal Location Id'lerin baştaki sıfırlarını atıyor
    # ('113704' ↔ OneMap ENTEGRASYON_ID / BOSS Lokasyon '00113704'). Kural: yalnız rakamsa 8 haneye sıfırla doldur.
    loc_norm = set(lokasyon_norm(_kod(master["location_id"])).dropna())
    lok_n = lokasyon_norm(lok)
    eslesti_n = lok_n.isin(loc_norm)
    musteri_bina = boss.loc[eslesti_n].groupby("Müşteri No")["Lokasyon"].first()
    kurtarilan = (~eslesti_n) & boss["Müşteri No"].isin(musteri_bina.index)
    p["lokasyon"]["normalizasyon"] = {
        "kural": "Lokasyon/Location Id yalnız rakamsa zfill(8); 'O' ile başlıyorsa olduğu gibi",
        "master_sifir_kaybi_olan_bina": int((_kod(master["location_id"]).str.fullmatch(r"\d+") &
                                              (_kod(master["location_id"]).str.len() < 8)).sum()),
        "eslesen_satir_normalize": int(eslesti_n.sum()),
        "eslesme_orani_dolularda_normalize": yuvarla(eslesti_n[lok.notna()].mean()),
        "eslesme_orani_tum_normalize": yuvarla(eslesti_n.mean()),
        "ayni_musterinin_baska_isinden_kurtarilan": int(kurtarilan.sum()),
        "toplam_binaya_baglanan_orani": yuvarla((eslesti_n | kurtarilan).mean()),
        "kategoriye_gore_normalize": {"kurulum": yuvarla(eslesti_n[kur].mean()), "mevcut_musteri": yuvarla(eslesti_n[~kur].mean())},
        "fiber_urunde_normalize": yuvarla(eslesti_n[boss["Ürün Bilgisi"].eq("Fiber")].mean()),
    }
    eslesti = eslesti_n
    lok = lok_n

    # aynı binada birden çok açık iş (toplu arıza / ortak altyapı sinyali)
    bina_basina = lok[eslesti].value_counts()
    p["lokasyon"]["ayni_binada_is"] = {"1": int((bina_basina == 1).sum()), "2": int((bina_basina == 2).sum()),
                                       "3-4": int(((bina_basina >= 3) & (bina_basina <= 4)).sum()),
                                       "5+": int((bina_basina >= 5).sum()), "en_cok": int(bina_basina.max()) if len(bina_basina) else 0}

    # Serbest metinde telefon numarası izi (yalnız sayım)
    for c in ["Son Açıklama", "Adres"]:
        p.setdefault("serbest_metinde_telefon_izi", {})[c] = int(boss[c].fillna("").astype(str).str.contains(TEL_RE).sum())

    # Zaman damgası doluluğu: gerçek SLA ölçümü için hangi olaylar görünür?
    zaman_kolonlari = ["Task Başlangıç Tarihi", "Task Bitiş Tarihi", "Randevu Başlangıç Tarihi", "Randevu Bitiş Tarihi",
                       "Teknik Ekip Konum Paylaşma Tarihi", "Teknik Ekip İşe Başlama Tarihi", "Teknik Ekip Bitirme Tarihi"]
    p["zaman_damgasi_dolu"] = {c: int(boss[c].notna().sum()) for c in zaman_kolonlari}
    p["SL_suresi_sa"] = yuzdelik(boss["SL Süresi(Sa)"])
    p["teknik_ekip_SL_dk"] = {"sifir": int((boss["Teknik Ekip SL Süresi(Dk)"] == 0).sum()),
                              "negatif": int((boss["Teknik Ekip SL Süresi(Dk)"] < 0).sum()),
                              "pozitif": int((boss["Teknik Ekip SL Süresi(Dk)"] > 0).sum())}
    return p


# ----------------------------------------------------------------------------- FOX
def _kalan(s: pd.Series) -> pd.Series:
    return pd.to_numeric(s.astype(str).str.replace(".", "", regex=False).str.replace(",", ".", regex=False), errors="coerce")


def fox_profili(fox: pd.DataFrame, boss: pd.DataFrame) -> dict:
    p: dict = {"satir": {k: int(v) for k, v in fox["_dosya"].value_counts().items()},
               "kolonlar": [c for c in fox.columns if c != "_dosya"],
               "akis_no_tekil": bool(fox["Akış No"].is_unique),
               "kisisel_kolonlar": KISISEL["FOX"],
               "dosya_bicimi": "BIFF .xls (Composite Document, 'NPOI' ile üretilmiş) — xlrd ile okunur; HTML değil"}
    for c in ["Akış Tipi", "Akış Statüsü", "Adım Statüsü", "Hedef SL", "Siparişi Oluşturan Kanal Bilgisi", "Müsteri İl"]:
        p.setdefault("dagilim", {})[c] = {f"{d}|{k}": int(v) for (d, k), v in fox.groupby("_dosya")[c].value_counts(dropna=False).items()}
    p["dagilim"]["Task Adı"] = dagilim(fox["Task Adı"], 60)
    kalan = _kalan(fox["Kalan Süre"])
    p["kalan_sure_saat"] = {
        "bicim": "virgüllü ondalık metin, saat; negatif = hedef SL aşıldı; Hedef SL=0 ise 0,00",
        "negatif_asilmis": int((kalan < 0).sum()), "sifir": int((kalan == 0).sum()), "pozitif": int((kalan > 0).sum()),
        "hedef_sl_0_iken_sifir": int(((fox["Hedef SL"] == 0) & (kalan == 0)).sum()),
        "negatiflerde_gecikme_saat": yuzdelik(-kalan[kalan < 0]),
    }
    p["hedef_sl_task_adina_gore"] = {
        str(k): {str(h): int(n) for h, n in g.value_counts().items()} for k, g in fox.groupby("Task Adı")["Hedef SL"]
    }
    p["ilce_bos_oran"] = yuvarla(fox["Müsteri İlce"].isna().mean())
    p["urun_bos_oran"] = yuvarla(fox["Ürün"].isna().mean())

    # Uzlaştırma
    fk, bk = set(fox["Akış No"]), set(boss["Task No"])
    fo = fox[~fox["Akış No"].isin(bk)]
    bo = boss[~boss["Task No"].isin(fk)]
    j = boss.merge(fox, left_on="Task No", right_on="Akış No", suffixes=("_b", "_f"))
    r = {
        "fox_toplam": len(fk), "boss_toplam": len(bk), "ortak": len(fk & bk),
        "yalniz_fox": len(fk - bk), "yalniz_boss": len(bk - fk),
        "yalniz_fox_dosya": dagilim(fo["_dosya"]),
        "yalniz_fox_task": dagilim(fo["Task Adı"]),
        "yalniz_fox_akis_statusu": dagilim(fo["Akış Statüsü"]),
        "yalniz_fox_kurulum_mu": {"kurulum": int(fo["Task Adı"].map(kurulum_mu).sum()),
                                  "mevcut": int((~fo["Task Adı"].map(kurulum_mu)).sum())},
        "yalniz_fox_yas_gun": yuzdelik((pd.Timestamp("2026-09-29 17:00") - fo["Başlangıç Tarihi"]).dt.total_seconds() / 86400),
        "yalniz_boss_task": dagilim(bo["Task Adı"]),
        "yalniz_boss_durum": dagilim(bo["Task Durumu"]),
        "musteri_no_tutarli_oran": yuvarla((j["Müşteri No_b"] == j["Müşteri No_f"]).mean()),
        "durum_capraz_boss_x_foxdosya": {f"{a}|{b}": int(v) for (a, b), v in j.groupby(["Task Durumu", "_dosya"]).size().items()},
        "durum_capraz_boss_x_foxakis": {f"{a}|{b}": int(v) for (a, b), v in j.groupby(["Task Durumu", "Akış Statüsü"]).size().items()},
        "fox_askida_boss_acik": int(((j["_dosya"] == "aski") & (j["Task Durumu"] != "Askıya alındı")).sum()),
        "boss_askida_fox_acik": int(((j["_dosya"] == "acik") & (j["Task Durumu"] == "Askıya alındı")).sum()),
        "fox_to_boss_gecikme_dk": yuzdelik((j["Task Başlangıç Tarihi"] - j["Başlangıç Tarihi"]).dt.total_seconds() / 60),
        "fox_baslangic_to_atanma_dk": yuzdelik((j["Atanma Zamanı"] - j["Başlangıç Tarihi"]).dt.total_seconds() / 60),
    }
    # Ad eşleme tablosu Fox → BOSS (tek kaynak sözlük)
    ad = j.groupby(["Task Adı_f", "Task Adı_b"]).size().reset_index(name="adet")
    r["ad_eslemesi_fox_boss"] = [
        {"fox": a, "boss": b, "adet": int(n), "kurulum_mu": kurulum_mu(b)} for a, b, n in ad.itertuples(index=False)
    ]
    p["uzlastirma"] = r
    return p


# ----------------------------------------------------------------------------- tur raporu farkı
def tur_farki() -> dict:
    kol = ["Bina Serial Number", "Tellcordia ID", "Location Id", "Adı", "Site Adı", "Blok Adı", "Kapı No", "Ilçe", "IL",
           "Obek Adı", "Kurumsal Sözleşme Tipi", "Aktiflik", "Sales Ready Tarihi", "Toplam HP", "Soho Op", "RES HP",
           "Aktif Abone Residential Segment"]

    def yukle(yol: Path) -> pd.DataFrame:
        d = pd.read_excel(yol, sheet_name="ORIGN", dtype={"Tellcordia ID": str, "Location Id": str, "Müşteri No": str})
        d.columns = [str(c).strip() for c in d.columns]
        return d

    eski, yeni = yukle(TUR_ESKI), yukle(TUR_YENI)
    p: dict = {"eski_dosya": "veri/raw/data.xlsx (2026-09-19)", "yeni_dosya": "PS26/data.xlsx ORIGN (2026-09-29)",
               "eski_satir": int(len(eski)), "yeni_satir": int(len(yeni)),
               "kolonlar_ayni": list(eski.columns) == list(yeni.columns)}
    p["eski_tekrar_eden_serial"] = int(eski["Bina Serial Number"].duplicated().sum())
    p["yeni_tekrar_eden_serial"] = int(yeni["Bina Serial Number"].duplicated().sum())
    e = eski.drop_duplicates("Bina Serial Number").set_index("Bina Serial Number")
    y = yeni.drop_duplicates("Bina Serial Number").set_index("Bina Serial Number")
    yeni_id, silinen = y.index.difference(e.index), e.index.difference(y.index)
    ortak = y.index.intersection(e.index)
    p["yeni_bina"] = int(len(yeni_id))
    p["kalkan_bina"] = int(len(silinen))
    degisen: dict = {}
    for c in kol[1:]:
        if c not in e.columns or c not in y.columns:
            continue
        # NaN != NaN tuzağı: boşları "" yap, sonra karşılaştır (pandas str dtype NaN'ı korur)
        a = e.loc[ortak, c].astype(object).where(e.loc[ortak, c].notna(), "").astype(str).str.strip()
        b = y.loc[ortak, c].astype(object).where(y.loc[ortak, c].notna(), "").astype(str).str.strip()
        n = int((a != b).sum())
        if n:
            degisen[c] = n
    p["ortak_binada_degisen_alan"] = degisen
    hp_a = pd.to_numeric(e.loc[ortak, "RES HP"], errors="coerce").fillna(0)
    hp_b = pd.to_numeric(y.loc[ortak, "RES HP"], errors="coerce").fillna(0)
    p["res_hp"] = {"eski_toplam": int(pd.to_numeric(e["RES HP"], errors="coerce").fillna(0).sum()),
                   "yeni_toplam": int(pd.to_numeric(y["RES HP"], errors="coerce").fillna(0).sum()),
                   "ortakta_degisen_bina": int((hp_a != hp_b).sum()),
                   "ortakta_net_degisim": int((hp_b - hp_a).sum())}
    # yeni binaların OneMap'te koordinatı var mı (Tellcordia ID = OneMap ID)?
    j = json.load(open(ONEMAP, encoding="utf-8"))
    om_id = {str(int(f["a"]["ID"])) for f in j["features"] if f["a"].get("ID") is not None}
    om_loc = {f["a"]["LOCATION_ID"] for f in j["features"]}
    if len(yeni_id):
        t = y.loc[yeni_id, "Tellcordia ID"].astype(str).str.strip()
        p["yeni_binada_onemap_cekimi_gereken"] = int((~t.isin(om_id) & ~pd.Index(yeni_id).isin(om_loc)).sum())
    p["ozdes_mi"] = bool(len(yeni_id) == 0 and len(silinen) == 0 and not degisen)
    return p


# ----------------------------------------------------------------------------- OneMap değişim alanları
def onemap_profili() -> dict:
    j = json.load(open(ONEMAP, encoding="utf-8"))
    a = pd.DataFrame([f["a"] for f in j["features"]])
    ref = pd.Timestamp(j.get("extracted_at", "2026-09-19")).tz_localize(None)
    son = pd.to_datetime(a["LAST_EDITED_DATE"], unit="ms")
    return {
        "kaynak_katman": j.get("source"), "cekim": j.get("extracted_at"), "bina": int(len(a)),
        "degisim_alani": "LAST_EDITED_DATE (epoch ms, %d dolu)" % int(a["LAST_EDITED_DATE"].notna().sum()),
        "son_duzenleme_yas": {
            "30_gun": int((son >= ref - pd.Timedelta(days=30)).sum()),
            "90_gun": int((son >= ref - pd.Timedelta(days=90)).sum()),
            "365_gun": int((son >= ref - pd.Timedelta(days=365)).sum()),
        },
        "CREATED_DATE_dolu": int(a["CREATED_DATE"].notna().sum()),
        "ISDELETED": dagilim(a["ISDELETED"]),
        "EDIT_DURUM": dagilim(a["EDIT_DURUM"]),
        "MAINTAIN": dagilim(a["MAINTAIN"]),
        "TEKNOLOJI": dagilim(a["TEKNOLOJI"]),
        "BAYI_NO": dagilim(a["BAYI_NO"]),
        "BESLEYEN_BAT_dolu": int(a["BESLEYEN_BAT"].notna().sum()),
        "alan_sayisi": int(a.shape[1]),
    }


# ----------------------------------------------------------------------------- kullanıcının takip çalışma kitabı
def takip_kitabi(master: pd.DataFrame) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(TUR_YENI, read_only=True)
    p: dict = {"sayfalar": {ws.title: {"satir": ws.max_row, "kolon": ws.max_column} for ws in wb.worksheets}}
    wb.close()
    loc_set = set(lokasyon_norm(master["location_id"]).dropna())
    for sayfa in ["TICKET", "GUZERGAH"]:
        d = pd.read_excel(TUR_YENI, sheet_name=sayfa, dtype=str)
        d = d.dropna(how="all")
        lok = lokasyon_norm(d["Lokasyon"]) if "Lokasyon" in d.columns else pd.Series(dtype=str)
        p[sayfa] = {
            "satir": int(len(d)), "kolonlar": [c for c in d.columns if not str(c).startswith("Unnamed")],
            "Konu": dagilim(d.get("Konu", pd.Series(dtype=str))), "Durum": dagilim(d.get("Durum", pd.Series(dtype=str))),
            "Kanal": dagilim(d.get("Kanal", pd.Series(dtype=str))),
            "ticket_dolu": int(d["Ticket"].notna().sum()) if "Ticket" in d.columns else None,
            "lokasyon_dolu": int(lok.notna().sum()), "lokasyon_tekil": int(lok.nunique()),
            "master_eslesme_orani": yuvarla(lok.dropna().isin(loc_set).mean()) if lok.notna().any() else None,
            "baslangic_araligi": [str(pd.to_datetime(d["Baslangic"], errors="coerce").min()),
                                  str(pd.to_datetime(d["Baslangic"], errors="coerce").max())],
        }
    d = pd.read_excel(TUR_YENI, sheet_name="ALTYAPI", dtype=str).dropna(how="all")
    p["ALTYAPI"] = {"satir": int(len(d)), "kolonlar": [c for c in d.columns if not str(c).startswith("Unnamed")],
                    "Kanal": dagilim(d["Kanal"]) if "Kanal" in d.columns else {}}
    return p


def ps26_yayin_riski() -> dict:
    """PS26 deposu data.xlsx'i JSON'a çevirip GitHub'a itiyor (guncelle.bat). Yalnız dosya adı/boyut."""
    if not PS26_DATA.exists():
        return {}
    return {
        "json_dosyalari": {f.name: f.stat().st_size for f in sorted(PS26_DATA.glob("*.json"))},
        "not": "guncelle.bat: convert.py → git add data → git push; GitHub Pages ile yayın. İçerikte tur raporu (ORIGN), "
               "müşteri no'lu TICKET/GUZERGAH/ALTYAPI sayfaları var. Pages siteleri depo özel olsa bile herkese açıktır "
               "(Enterprise Cloud erişim denetimi hariç).",
    }


# ----------------------------------------------------------------------------- görev tipi sözlüğü
# Kanonik tip: (kod, aile, btk_onceligi). Aile: baglanti · tv · ses · cihaz · iade · evrak · ucret · btk_sikayet ·
# sikayet · kurulum · diger. BTK önceliği kullanıcı kuralıdır ("bağlantı ve TV problemleri BTK'ya sayar");
# 'kanal_sikayeti' ve 'arama_problemi' için teyit gerekir (acik_soru=True).
KANONIK = {
    "Bağlantı Problemleri": ("BAGLANTI_PROBLEMI", "baglanti", True, False),
    "Arıza Bildirim": ("DOPING_ARIZA", "baglanti", True, True),
    "IP TV Arıza": ("TV_ARIZA", "tv", True, False),
    "Kanal Şikayeti": ("KANAL_SIKAYETI", "tv", True, True),
    "Arama Problemi": ("ARAMA_PROBLEMI", "ses", True, True),
    "BTK Şikayet": ("BTK_SIKAYET", "btk_sikayet", True, False),
    "BTK / Mahkeme Şikayet": ("BTK_MAHKEME_SIKAYET", "btk_sikayet", True, False),
    "Modem Değişikliği": ("MODEM_DEGISIKLIGI", "cihaz", False, False),
    "Superbox Modem Değişikliği": ("SUPERBOX_MODEM_DEGISIKLIGI", "cihaz", False, False),
    "STB Cihaz Değişikliği": ("STB_DEGISIKLIGI", "cihaz", False, False),
    "Cihaz İade Bekleniyor": ("CIHAZ_IADE", "iade", False, False),
    "Cihaz Geri Alım": ("CIHAZ_GERI_ALIM", "iade", False, False),
    "Turksat Cihaz İade Bekleniyor": ("TURKSAT_CIHAZ_IADE", "iade", False, False),
    "Sosyal Destek Evrak Toplama": ("SOSYAL_DESTEK_EVRAK", "evrak", False, False),
    "Turksat Evrak Toplama": ("TURKSAT_EVRAK", "evrak", False, False),
    "Ücretlendirilecek Servisler": ("TEKNIK_SERVIS_UCRET", "ucret", False, False),
    "Donanım Teslimat Şikayet": ("DONANIM_TESLIMAT_SIKAYET", "sikayet", False, False),
    "Bayi Kanal İnceleme": ("BAYI_KANAL_INCELEME", "sikayet", False, False),
    "Kurulum Taskı Ürememiş": ("KURULUM_TASKI_UREMEMIS", "diger", False, False),
    "Fiber gelsin": ("FIBER_GELSIN", "diger", False, False),
    "Kampanya Tanımlama Problemleri": ("KAMPANYA_TANIMLAMA", "diger", False, False),
    "XDSL Boş Port Yok": ("XDSL_BOS_PORT_YOK", "diger", False, False),
    "Nakil/Numara Değişikliği": ("NAKIL_NUMARA_DEGISIKLIGI", "diger", False, False),
    "BDH Problem Çözüm": ("BDH_PROBLEM_COZUM", "diger", False, False),
    "Fatura İtiraz Bildirimi": ("FATURA_ITIRAZ", "diger", False, False),
    "Quiknet Kurulum Talebi": ("FIBER_KURULUM", "kurulum", False, False),
    "IP TV Kurulum": ("TV_KURULUM", "kurulum", False, False),
    "SuperBox Kurulum": ("SUPERBOX_KURULUM", "kurulum", False, False),
    "Kurulum ve cihaz gönderim": ("KURULUM_CIHAZ_GONDERIM", "kurulum", False, False),
    "TT Fiber Kurulum Talebi": ("TT_FIBER_KURULUM", "kurulum", False, False),
    # Adı Fox'ta 'Kurulum' içeriyor, BOSS'ta içermiyor → kural hangi ada uygulanırsa sonuç değişir (açık soru)
    "İkinci Donanım Kurulum": ("IKINCI_DONANIM", "kurulum", False, True),
    "TV Yan Oda Kurulum": ("TV_YAN_ODA", "kurulum", False, True),
}


def gorev_sozlugu(fox: pd.DataFrame, boss: pd.DataFrame) -> list[dict]:
    j = boss.merge(fox, left_on="Task No", right_on="Akış No", suffixes=("_b", "_f"))
    boss_adlari = j.groupby("Task Adı_f")["Task Adı_b"].agg(lambda s: sorted(set(s))).to_dict()
    sl = fox.groupby("Task Adı")["Hedef SL"].agg(lambda s: int(s[s > 0].mode().iloc[0]) if (s > 0).any() else 0).to_dict()
    adet = fox["Task Adı"].value_counts().to_dict()
    yalniz_fox = set(fox.loc[~fox["Akış No"].isin(boss["Task No"]), "Task Adı"]) - set(boss_adlari)
    satirlar = []
    for fox_ad in sorted(set(fox["Task Adı"])):
        kod, aile, btk, soru = KANONIK.get(fox_ad, (None, "tanimsiz", False, True))
        b_adlar = boss_adlari.get(fox_ad, [])
        kural_fox = kurulum_mu(fox_ad)
        kural_boss = [kurulum_mu(x) for x in b_adlar]
        satirlar.append({
            "fox_adi": fox_ad, "boss_adlari": b_adlar, "kod": kod, "aile": aile,
            "kurulum_kurali_fox_adina": kural_fox,
            "kurulum_kurali_boss_adina": (all(kural_boss) if kural_boss else None),
            "kural_celiskili": bool(kural_boss) and any(k != kural_fox for k in kural_boss),
            "sorumlu_birim_onerisi": "yeni_musteri_kurulum" if aile == "kurulum" else "operasyon",
            "btk_onceligi": btk, "turkcell_hedef_sl_saat": sl.get(fox_ad, 0),
            "yalniz_fox_ta": fox_ad in yalniz_fox, "acik_adet": int(adet.get(fox_ad, 0)),
            "teyit_gerekir": soru,
        })
    return satirlar


def main() -> int:
    for akis in (sys.stdout, sys.stderr):
        try:
            akis.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, OSError, ValueError):
            pass
    CIKTI.mkdir(parents=True, exist_ok=True)
    boss, fox, master = oku()
    sonuc = {
        "uretim": pd.Timestamp.now().strftime("%Y-%m-%d %H:%M"),
        "gizlilik": "yalnız toplu sayımlar; kişisel veri yok",
        "boss": boss_profili(boss, master),
        "fox": fox_profili(fox, boss),
        "tur_raporu_farki": tur_farki(),
        "onemap": onemap_profili(),
        "takip_kitabi": takip_kitabi(master),
        "ps26_yayin": ps26_yayin_riski(),
        "master": {"bina": int(len(master)), "location_id_dolu": int(_kod(master["location_id"]).notna().sum()),
                   "location_id_tekil": int(_kod(master["location_id"]).nunique())},
    }
    yol = CIKTI / "entegrasyon_profil.json"
    yol.write_text(json.dumps(sonuc, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    sozluk = gorev_sozlugu(fox, boss)
    (CIKTI / "gorev_sozlugu.json").write_text(json.dumps(
        {"aciklama": "Fox→BOSS görev adı eşlemesi + kanonik tip, kurulum kuralı, BTK önceliği, Turkcell hedef SL (Fox)",
         "kaynak": "2026-09-29 Fox açık+askı ve BOSS açık dışa aktarımları", "tipler": sozluk},
        ensure_ascii=False, indent=1), encoding="utf-8")
    b, f, t =sonuc["boss"], sonuc["fox"]["uzlastirma"], sonuc["tur_raporu_farki"]
    print(f"BOSS {b['satir']} · Lokasyon dolu {b['lokasyon']['dolu']} · eşleşen {b['lokasyon']['master_location_id_eslesen_satir']} "
          f"({b['lokasyon']['master_location_id_eslesme_orani_tum']:.1%} tüm)")
    print(f"Fox {f['fox_toplam']} · ortak {f['ortak']} · yalnız Fox {f['yalniz_fox']} · yalnız BOSS {f['yalniz_boss']} · "
          f"Fox→BOSS medyan {f['fox_to_boss_gecikme_dk'].get('medyan')} dk")
    print(f"Tur: yeni {t['yeni_bina']} · kalkan {t['kalkan_bina']} · değişen alan {t['ortak_binada_degisen_alan']}")
    print(f"→ {yol}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
