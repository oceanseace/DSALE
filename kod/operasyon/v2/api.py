"""v2 iş emri uçları (spec §5.3.2–§5.3.6, EK-3, EK-4, EK-10).

``saha/api.py`` (WP-A) bir ``Bagimliliklar`` (izin, kullanici, baglanti) verir ve ``yonlendiriciler(b)``'nin
döndürdüğü yönlendiricileri bağlar — döngüsel içe aktarma yok. Motor hataları (``hatalar.V2Hata``) burada tek yerde
``{"hata": "<Türkçe cümle>", "kod": "<kod>", …}`` gövdesine çevrilir.

Müşteri bilgisi yalnız ``gorunum`` üzerinden çıkar; döndüğü her yanıt ``erisim_kaydi``'na yazılır (F21).
"""
from __future__ import annotations

import datetime as dt
import io
import json
import sqlite3
from contextlib import asynccontextmanager
from typing import Any, Callable

from fastapi import APIRouter, Depends, Header, HTTPException, Query, Request
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict, Field
from starlette.concurrency import run_in_threadpool

import yollar
from operasyon import is_emri as ie

from . import (akis, aktarim, aski, boss, boss_detay, boss_ham, boss_yazma, dagitim, gorunum, isaret, izleme, kesinti,
               kisi, kurallar, obek, siralama, sozluk, takip, tasks_tercih, zaman)
from .hatalar import V2Hata
from .islem import islem


# Bu modülün yeni uçlarının izinleri (EK-12). ``saha/api.py`` bağlarken ``yetki.ROTA_IZNI``'na ekler
# (``rota_izni_ekle``); sözleşme tablosundaki satırlar ezilmez. Eylemler yetki.IZINLER'de tanımlıdır.
ROTA_IZNI: dict[tuple[str, str], str | tuple[str, ...]] = {
    ("GET", "/api/isler/kurallar"): "is.liste",
    ("PUT", "/api/isler/kurallar"): "veri.yonet",
    ("POST", "/api/isler/{is_no}/arama"): "is.duzenle",
    ("POST", "/api/isler/{is_no}/btk-sikayet"): "is.duzenle",
    ("POST", "/api/isler/{is_no}/tcs"): "is.duzenle",
    ("POST", "/api/isler/{is_no}/bizim-durum"): "is.duzenle",
    ("GET", "/api/kesinti"): "is.duzenle",
    ("POST", "/api/kesinti"): "is.duzenle",
    ("GET", "/api/kesinti/aday"): "is.duzenle",
    ("POST", "/api/kesinti/{kesinti_id}/bitti"): "is.duzenle",
    # Spec §6.10 "Bugün çalışmıyor" anahtarının yazma ucu (okuma: GET /api/isler/teknikler.bugun_yok).
    ("PUT", "/api/isler/teknikler/{teknik_id}/bugun-yok"): "kapasite.yaz",
    # İş türüne göre ayır + kuralla ofis çalışanına dağıt (Mevcut Müşteri → Tablo)
    ("GET", "/api/dagitim"): "is.ata",
    ("PUT", "/api/dagitim/kurallar"): "is.ata",
    ("POST", "/api/dagitim/uygula"): "is.ata",
    ("POST", "/api/dagitim/ata"): "is.ata",
    # Görünümler: adlandırılmış task başlığı süzgeçleri (+ sorumlu)
    ("GET", "/api/isler/gorunumler"): "is.liste",
    ("PUT", "/api/isler/gorunumler"): "is.ata",
    # TASKS/MAP kişisel tercihleri (Excel düzeni, kayıtlı düzenler, ortak süzgeç) + unvan varsayılanı (Administrator)
    ("GET", "/api/isler/tercih"): "is.liste",
    ("PUT", "/api/isler/tercih"): "is.liste",
    ("PUT", "/api/isler/tercih/unvan"): "is.liste",
    # BOSS canlı eşitleme (eklenti): eşleme koduyla ya da rapor yükleme izniyle
    ("POST", "/api/boss/esitle"): "is.yukle",
    ("POST", "/api/boss/gunluk"): "is.yukle",
    ("GET", "/api/boss/akis"): "is.yukle",
    ("GET", "/api/boss/durum"): "is.liste",
    ("GET", "/api/boss/eslesme-kodu"): "is.ata",
    ("POST", "/api/boss/eslesme-kodu/yenile"): "is.ata",
    # BOSS iş ayrıntısı (isteğe bağlı, eklenti okur) + BOSS ekip listesi
    ("POST", "/api/isler/{is_no}/boss-detay"): "is.liste",
    ("GET", "/api/isler/{is_no}/boss-detay"): "is.liste",
    ("POST", "/api/boss/detay-oncelik"): "is.liste",
    ("GET", "/api/boss/istekler"): "is.yukle",
    ("POST", "/api/boss/esitle-iste"): "is.liste",
    ("GET", "/api/boss/izler"): "is.yukle",
    ("POST", "/api/boss/iletisim"): "is.yukle",
    ("POST", "/api/boss/detay"): "is.yukle",
    ("POST", "/api/boss/detay-onden"): "is.yukle",
    ("GET", "/api/boss/ekipler"): "is.ata",
    # BOSS ekip GPS konumları (eklenti BOSS "GPS Takip" ucunu okur; MAP "Ekipler" katmanı) — ekip_konum.py
    ("POST", "/api/boss/ekip-konum"): "is.yukle",
    ("GET", "/api/isler/ekip-konumlari"): "is.liste",
    # TASKS tablosu: satır rengi
    ("POST", "/api/isler/{is_no}/isaret"): "is.duzenle",
    # TASKS tablosu: hücre dolgu rengi (Excel gibi yalnız seçili hücreler)
    ("POST", "/api/isler/hucre-renk"): "is.duzenle",
    # BOSS'a yazma (genel anahtar varsayılan KAPALI): eklenti yoklar + sonucu bildirir; ekranda önizleme ve askı
    # nedenleri. Genel anahtar ve tek iş denemesi YALNIZ Administrator'a (yetkinlik.yonet).
    ("GET", "/api/boss/yazilacaklar"): "is.yukle",
    ("POST", "/api/boss/yazma-sonucu"): "is.yukle",
    ("GET", "/api/isler/{is_no}/boss-yazma-onizleme"): "is.liste",
    ("GET", "/api/boss/aski-nedenleri"): "is.duzenle",
    ("GET", "/api/isler/{is_no}/boss-aski-nedenleri"): "is.duzenle",
    ("PUT", "/api/boss/yazma-anahtari"): "yetkinlik.yonet",
    ("POST", "/api/boss/yazma-dene"): "yetkinlik.yonet",
    # Toplu atamanın BOSS hâli (TASKS/MAP "Ata…" özeti): iş no → gidiyor / tamam / hata (nedeni) / kapalı
    ("POST", "/api/boss/atama-durumlari"): "is.ata",
    # BOSS'a yazılamayan toplu atamaları yeniden kuyruğa al (her iş tek başına istekte)
    ("POST", "/api/boss/atama-yeniden"): "is.ata",
    # BOSS yazma kuyruğu (hata ayıklama; kişisel veri yok): son 7 günün kayıtları + tek kayıt "Yeniden dene"
    ("GET", "/api/boss/yazmalar"): "is.ata",
    ("POST", "/api/boss/yazmalar/{kayit_id}/yeniden"): "is.ata",
    # Toplu askı (TASKS/MAP "Askıya al…"): seçili işlerin BOSS askı nedenleri + uygula
    ("POST", "/api/isler/toplu-aski/nedenler"): "is.duzenle",
    ("POST", "/api/isler/toplu-aski"): "is.duzenle",
    # İlişkili işler birlikte (09.10): seçili işlerin aynı müşterideki öbür açık işleri (Ata/Askı formu) ve
    # "İlişkili iş farklı ekipte/atanmamış — hepsini aynı ekibe ata" (Pano/TASKS)
    ("POST", "/api/isler/iliskili/isler"): "is.liste",
    ("GET", "/api/isler/iliskili/uyumsuz"): "is.liste",
    ("POST", "/api/isler/iliskili/hizala"): "is.ata",
    # Eşitleme değişiklik akışı (son 50 eşitlemenin yeni/değişen/kapanan task noları; bellekte)
    ("GET", "/api/boss/farklar"): "is.liste",
}


# Canlı olay akışı: bellekte halka arabellek (diske yazılmaz, şişmez)
from collections import deque as _deque
_AKIS: "_deque[dict]" = _deque(maxlen=500)
_AKIS_NO = [0]
# Eşitleme farkları (değişiklik akışı): son 50 eşitleme, yalnız task noları (kişisel veri yok); bellekte
_FARKLAR: "_deque[dict]" = _deque(maxlen=50)
_FARK_NO = [0]


def _fark_ekle(sonuc: dict | None, govde: dict) -> None:
    d = (sonuc or {}).get("degisiklik")
    if not isinstance(d, dict):
        return
    _FARK_NO[0] += 1
    _FARKLAR.append({"no": _FARK_NO[0], "zaman": zaman.metin(), "mod": (sonuc or {}).get("mod"),
                     "tam": bool(govde.get("tam")), "yeni": list(d.get("yeni") or []),
                     "degisen": list(d.get("degisen") or []), "kapanan": list(d.get("kapanan") or [])})


def _akis_ekle(olay: dict) -> None:
    _AKIS_NO[0] += 1
    _AKIS.append({"no": _AKIS_NO[0], "zaman": zaman.metin(), **{k: v for k, v in olay.items() if k != "no"}})


def _json_varsayilan(v):
    """json.dumps'ın tanımadığı değerler (motor yalnız metin/sayı döndürür; bu yalnız güvenlik ağı)."""
    if isinstance(v, dt.datetime):
        return zaman.metin(v)
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, (set, frozenset, tuple)):
        return list(v)
    if isinstance(v, sqlite3.Row):
        return dict(v)
    if isinstance(v, bytes):
        return v.decode("utf-8", "replace")
    raise TypeError(f"JSON'a çevrilemeyen değer: {type(v).__name__}")


class HizliJSON(Response):
    """Büyük liste yanıtları doğrudan ``json.dumps`` ile yazılır.

    FastAPI dönen sözlüğü önce ``jsonable_encoder``'dan geçirir: 437 işlik ``GET /api/isler``'de bu tek başına
    ≈80 ms (her alanda isinstance/dataclass denetimi). Motorun yanıtı zaten JSON'a hazırdır (metin, sayı, bool,
    None, liste, sözlük); spec §5.6 bütçesi (< 150 ms) bu yüzden bu yolla tutulur. Biçim Starlette'in
    ``JSONResponse``'uyla aynıdır (UTF-8, ``allow_nan=False``).
    """
    media_type = "application/json"

    def render(self, content) -> bytes:
        return json.dumps(content, ensure_ascii=False, allow_nan=False, separators=(",", ":"),
                          default=_json_varsayilan).encode("utf-8")


def _hata(e: V2Hata) -> HTTPException:
    return HTTPException(status_code=e.durum, detail=e.govde())


def _cagir(fn: Callable, *a, **kw):
    try:
        return fn(*a, **kw)
    except V2Hata as e:
        raise _hata(e)


def _fabrika():
    from saha import db
    return db.baglan()


def _erisim(conn: sqlite3.Connection, k: dict | None, eylem: str, is_no: str | None = None, adet: int | None = None,
            ip: str | None = None) -> None:
    """KVKK erişim kaydı: iş ayrıntısı kişi+iş+gün başına bir satır; listeler adetle."""
    if not k:
        return
    try:
        with islem(conn):
            conn.execute("INSERT OR IGNORE INTO erisim_kaydi (zaman, gun, kullanici_id, kullanici_ad, eylem, is_no, adet, ip) "
                         "VALUES (?,?,?,?,?,?,?,?)", (zaman.metin(), zaman.metin()[:10], k["id"], k.get("ad"), eylem,
                                                      is_no, adet, ip))
    except sqlite3.Error:
        pass


def _ip(istek: Request | None) -> str | None:
    return istek.client.host if istek is not None and istek.client else None


def _yonetim_kaydi(conn, k, eylem, hedef, ozet) -> None:
    try:
        with islem(conn):
            conn.execute("INSERT INTO yonetim_kaydi (zaman, kullanici_id, kullanici_ad, eylem, hedef, ozet) "
                         "VALUES (?,?,?,?,?,?)", (zaman.metin(), (k or {}).get("id"), (k or {}).get("ad"), eylem, hedef,
                                                  json.dumps(ozet, ensure_ascii=False) if ozet is not None else None))
    except sqlite3.Error:
        pass


# ============================================================================= gövdeler
class _Model(BaseModel):
    model_config = ConfigDict(extra="ignore")


class Randevu(_Model):
    bas: str | None = None
    bit: str | None = None
    teyitli: bool = False


class AtaGirdi(_Model):
    teknik_id: int
    randevu: Randevu | None = None
    yine_de: bool = False
    notu: str | None = Field(default=None, max_length=1000)
    surum: int | None = None


class OneriOnayGirdi(_Model):
    isler: dict[str, int]


class TopluAtama(_Model):
    is_no: str
    teknik_id: int
    randevu: Randevu | None = None
    surum: int | None = None
    yine_de: bool = False          # yalnız iş çekmecesinin tek işi (altyapı engeline rağmen); toplu panel göndermez


class DagitimKurallariGirdi(_Model):
    kurallar: list[dict[str, Any]]


class GorunumlerGirdi(_Model):
    gorunumler: list[dict[str, Any]] = Field(max_length=200)


class UnvanDuzenGirdi(_Model):
    unvan_id: int
    duzen: dict[str, Any] | None = None


class DagitimAtaGirdi(_Model):
    task_nolar: list[str] = Field(max_length=5000)
    kisi_id: int | None = None


class TopluAtaGirdi(_Model):
    atamalar: list[TopluAtama] = Field(max_length=500)
    parca: str | None = Field(default=None, max_length=80)


class AtamaDurumGirdi(_Model):
    is_nolar: list[str] = Field(max_length=500)


class TopluAskiNedenGirdi(_Model):
    is_nolar: list[str] = Field(max_length=2000)


class IliskiliGirdi(_Model):
    is_nolar: list[str] = Field(max_length=500)


class IliskiliHizalaGirdi(_Model):
    is_nolar: list[str] | None = Field(default=None, max_length=500)


class TopluAskiGirdi(_Model):
    is_nolar: list[str] = Field(min_length=1, max_length=500)
    neden: str = Field(min_length=1, max_length=120)
    alt: str | None = Field(default=None, max_length=160)
    bitis: str = Field(min_length=10, max_length=25)
    aciklama: str | None = Field(default=None, max_length=500)


class RandevuGirdi(_Model):
    bas: str | None = None
    bit: str | None = None
    teyitli: bool = False
    notu: str | None = Field(default=None, max_length=1000)
    surum: int | None = None


class DurumGirdi(_Model):
    yeni: str
    neden: str | None = Field(default=None, max_length=300)
    evde_miydi: bool | None = None
    sonuc_kodu: str | None = Field(default=None, max_length=40)
    uyanma: str | None = None
    canli_test: bool | None = None
    randevu: Randevu | None = None
    teknik_id: int | None = None
    notu: str | None = Field(default=None, max_length=1000)
    surum: int | None = None
    istemci_id: str | None = Field(default=None, max_length=80)
    zaman: str | None = None


class TeshisGirdi(_Model):
    sonuc: str
    canli_test: bool = False
    randevu: Randevu | None = None
    surum: int | None = None


class YeniTicket(_Model):
    konu: str
    ticket_no: str | None = None
    detay: str | None = Field(default=None, max_length=2000)


class TicketGirdi(_Model):
    ticket_id: int | None = None
    yeni: YeniTicket | None = None
    teknisyen_gonderme: bool = True
    surum: int | None = None


class ObekGirdi(_Model):
    obek_id: int | None = None
    surum: int | None = None


class MahalleGirdi(_Model):
    mahalle_id: int
    surum: int | None = None


class NotGirdi(_Model):
    metin: str = Field(min_length=1, max_length=1000)
    istemci_id: str | None = Field(default=None, max_length=80)


class IletisimGirdi(_Model):
    musteri_tel: str | None = None
    surum: int | None = None


class KopyaGirdi(_Model):
    alan: str = "musteri_no"


class BossIslendiGirdi(_Model):
    islendi: bool = True
    surum: int | None = None


class BossBaglaGirdi(_Model):
    boss_task_no: str
    surum: int | None = None


class GeriAlGirdi(_Model):
    olay_id: int


class YeniIsGirdi(_Model):
    task_adi: str = Field(min_length=1, max_length=120)
    bina_serial: str | None = None
    adres: str | None = Field(default=None, max_length=500)
    mahalle_id: int | None = None
    musteri_adi: str | None = Field(default=None, max_length=200)
    musteri_no: str | None = Field(default=None, max_length=40)
    aciklama: str | None = Field(default=None, max_length=1000)
    istemci_id: str | None = Field(default=None, max_length=80)


class DagitGirdi(_Model):
    obek_id: int
    teknik_idler: list[int]
    isler: dict[str, int] | None = None


class BossEkipEsleGirdi(_Model):
    boss_ekip: str = Field(min_length=1, max_length=200)
    teknik_id: int


class GorduGirdi(_Model):
    is_nolar: list[str]


class SiraGirdi(_Model):
    is_no: str
    yeni_sira: int
    neden: str = "diger"
    istemci_id: str | None = Field(default=None, max_length=80)


class AktarimUygulaGirdi(_Model):
    mod: str = "tam"


class AramaGirdi(_Model):
    sonuc: str
    sure_sn: int | None = Field(default=None, ge=0, le=3600)
    webphone: bool = True
    sms: bool = False
    yeni_numara: bool = False
    notu: str | None = Field(default=None, max_length=1000)
    istemci_id: str | None = Field(default=None, max_length=80)


class BtkSikayetGirdi(_Model):
    tarih: str | None = None
    reopen: bool = False
    kaldir: bool = False
    surum: int | None = None


class BizimDurumGirdi(_Model):
    deger: str | None = None
    surum: int | None = None


class IsaretGirdi(_Model):
    renk: str | None = None


class HucreRenkGirdi(_Model):
    hucreler: list[dict] = Field(default_factory=list)
    renk: str | None = None


class TcsGirdi(_Model):
    tarih: str
    teyitli: bool = False
    surum: int | None = None


class KesintiGirdi(_Model):
    metin: str = Field(min_length=1, max_length=5000)


class IzlemeGirdi(_Model):
    acik: bool | None = None
    klasorler: list[str] | None = None
    desenler: list[str] | None = None


class ObekYeniGirdi(_Model):
    ad: str = Field(min_length=1, max_length=80)
    surum: int | None = None


class ObekGuncelleGirdi(_Model):
    ad: str | None = Field(default=None, max_length=80)
    renk: int | None = None
    sahip_id: int | None = None
    yedek_id: int | None = None
    surum: int | None = None


class IlceSecimi(_Model):
    il: str | None = None
    ilce: str


class ObekMahalleGirdi(_Model):
    mahalle_idler: list[int] | None = None
    tum_ilce: list[IlceSecimi] | None = None
    tasi: bool = False
    surum: int | None = None


class ObekCikarGirdi(_Model):
    refler: list[str]
    surum: int | None = None


class ObekGeriAlGirdi(_Model):
    olay_id: int
    surum: int | None = None


class ObekBolGirdi(_Model):
    k: int = Field(ge=2, le=10)
    onizle: bool = True
    surum: int | None = None


class MahalleYeniGirdi(_Model):
    il: str
    ilce: str
    ad: str = Field(min_length=1, max_length=80)
    benzerine_ragmen: bool = False


class EsadGirdi(_Model):
    il: str | None = None
    ilce: str | None = None
    esad: str = Field(min_length=1, max_length=80)
    mahalle_id: int


class KapasiteGirdi(_Model):
    tarih: str | None = None
    aktif_teknik: int | None = Field(default=None, ge=0, le=200)


class BugunYokGirdi(_Model):
    tarih: str | None = None
    yok: bool


class BossYazmaAnahtarGirdi(_Model):
    acik: bool


class BossYazmaDeneGirdi(_Model):
    is_no: str = Field(min_length=1, max_length=40)
    eylem: str = Field(min_length=1, max_length=20)
    aski_neden_id: str | None = Field(default=None, max_length=64)
    aski_alt_neden_id: str | None = Field(default=None, max_length=64)


class IsAyarGirdi(_Model):
    btk_durduran_aski: list[str] | None = None
    son24_askida_durur: bool | None = None
    kutlamalar: bool | None = None
    aski_nedenleri: list[dict] | None = None
    boss_task_url: str | None = Field(default=None, max_length=300)
    musteri_tel: str | None = None
    dilimler: str | None = Field(default=None, max_length=120)
    teknik_kapasite: int | None = Field(default=None, ge=1, le=60)
    filtre_istisnalari: list[str] | None = None
    boss_yazma: str | None = None
    durum_renkleri: dict[str, str] | None = None
    boss_yazma_not: str | None = None
    boss_yazma_atama: str | None = None
    # otomatik planlama (09.10; yalnız Administrator): ilişkili işleri dahil et, ticket otomatik askı, çözümde geri ata
    iliskili_dahil: bool | None = None
    ticket_otomatik_aski: bool | None = None
    ticket_cozum_yeniden_ata: bool | None = None
    ticket_cozum_slot: str | None = Field(default=None, max_length=20)
    ticket_aski_alt_neden: str | None = Field(default=None, max_length=160)   # eski (NETWORK); yok sayılır
    boss_yazma_paralel: int | None = Field(default=None, ge=1, le=10)       # eklentinin eşzamanlı BOSS yazması


# ============================================================================= yardımcılar
# MAP nokta rengi: durum → renk (5 renk). TEK tablo; arayüzde aynısı cekirdek/is/durumRenk.ts DURUM_RENK_VARSAYILAN.
# Administrator Ayarlar'dan değiştirir (PUT /api/isler/ayarlar {durum_renkleri}); kayıt ayar('durum_renkleri').
DURUM_RENKLERI = ("mor", "kirmizi", "sari", "gri", "mavi")
DURUM_RENK_VARSAYILAN = {
    "askiya_alindi": "mor", "ertelendi": "kirmizi", "merkeze_gonderildi": "kirmizi", "altyapi": "kirmizi",
    "randevu_alindi": "sari", "acik_randevulu": "sari", "acik": "gri",
    "basladi": "mavi", "konum_paylasildi": "mavi",
}


def durum_renkleri(conn: sqlite3.Connection) -> dict[str, str]:
    """Kayıtlı eşleme varsayılanın üstüne (bilinmeyen anahtar/renk yok sayılır)."""
    r = conn.execute("SELECT deger FROM ayar WHERE anahtar='durum_renkleri'").fetchone()
    try:
        kayit = json.loads(r[0]) if r and r[0] else {}
    except ValueError:
        kayit = {}
    out = dict(DURUM_RENK_VARSAYILAN)
    if isinstance(kayit, dict):
        out.update({k_: v for k_, v in kayit.items() if k_ in DURUM_RENK_VARSAYILAN and v in DURUM_RENKLERI})
    return out


def is_ayarlari(conn: sqlite3.Connection) -> dict:
    """``IsAyarlari`` — ekranın ihtiyaç duyduğu ayarlar (herkese değil, iş listesine izni olana)."""
    def ayar(a, v=None):
        r = conn.execute("SELECT deger FROM ayar WHERE anahtar=?", (a,)).fetchone()
        return r[0] if r and r[0] not in (None, "") else v
    mesai = str(ayar("mesai", "08:00-20:00")).split("-")
    return {"dilimler": [{"bas": f"{a:02d}:00", "bit": f"{b:02d}:00"} for a, b in siralama.dilimler(conn)],
            "mesai": {"bas": mesai[0].strip(), "bit": (mesai[1] if len(mesai) > 1 else "20:00").strip()},
            "musteri_tel": "acik" if ayar("musteri_tel", "kapali") == "acik" else "kapali",
            "boss_task_url": ayar("boss_task_url"),
            "aski_nedenleri": aski.nedenler(conn), "son24_askida_durur": aski.son24_durur(conn),
            "btk_durduran_aski": json.loads(ayar("btk_durduran_aski", "null") or "null") or aski.VARSAYILAN_KURAL,
            "kutlamalar": str(ayar("kutlamalar", "acik")) == "acik",
            "kapanis_nedenleri": [{"kod": k_, "metin": v} for k_, v in akis.KAPANIS_NEDENLERI.items()],
            "teknik_kapasite": siralama.varsayilan_kapasite(conn),
            "boss_yazma": "acik" if boss_yazma.yazma_acik(conn) else "kapali",
            "boss_yazma_not": "acik" if boss_yazma.eylem_acik(conn, "not") else "kapali",
            "boss_yazma_atama": "acik" if boss_yazma.eylem_acik(conn, "atama") else "kapali",
            "durum_renkleri": durum_renkleri(conn), **otomatik_ayarlar(conn)}


def otomatik_ayarlar(conn: sqlite3.Connection) -> dict:
    """Otomatik planlama anahtarları (Ayarlar; Administrator değiştirir)."""
    from saha import ticket_boss_aski as tba

    from . import iliskili
    bas, bit = tba.cozum_slotu(conn)
    return {"iliskili_dahil": iliskili.dahil_varsayilan(conn),
            "ticket_otomatik_aski": tba.acik_mi(conn, tba.AYAR_ASKI),
            "ticket_cozum_yeniden_ata": tba.acik_mi(conn, tba.AYAR_YENIDEN_ATA),
            "ticket_cozum_slot": f"{bas}-{bit}",
            # bütün askıların varsayılanı (tek, toplu, MAP, ticket): ABONE > Aboneye Ulaşılamadı
            "varsayilan_aski": boss_yazma.varsayilan_aski(),
            "boss_yazma_paralel": boss_yazma.paralel(conn)}


def boss_ekip_listesi(conn: sqlite3.Connection) -> dict:
    h = kisi.boss_ekip_haritasi(conn)
    sayim: dict[str, list] = {}
    for (ad,) in conn.execute("SELECT boss_ekip FROM is_emri WHERE boss_ekip IS NOT NULL AND atanan_id IS NULL AND "
                              "durum IN ('triyaj','bekliyor','randevulu')"):
        a = ie.anahtar(ad)
        if a and a not in h:
            sayim.setdefault(a, [ad, 0])[1] += 1
    kisiler = kisi.teknikler(conn)
    eslesmis = [{"boss_ekip": t["boss_ekip"], "kisi": kisi.ozet(t)} for t in kisiler.values() if t.get("boss_ekip")]
    return {"eslesmemis": sorted(({"boss_ekip": v[0], "is": v[1]} for v in sayim.values()), key=lambda x: -x["is"]),
            "eslesmis": sorted(eslesmis, key=lambda x: x["boss_ekip"])}


def teknik_listesi(conn: sqlite3.Connection, tarih: str | None = None) -> list[dict]:
    gun = tarih or zaman.simdi().strftime("%Y-%m-%d")
    try:
        yok = set(json.loads((conn.execute("SELECT deger FROM ayar WHERE anahtar=?", (f"bugun_yok_{gun}",)).fetchone()
                              or ["[]"])[0] or "[]"))
    except ValueError:
        yok = set()
    obekler: dict[int, list] = {}
    for oid, ad, sahip, yedek in conn.execute("SELECT id, ad, sahip_id, yedek_id FROM obek WHERE aktif=1 ORDER BY ad"):
        for t in (sahip, yedek):
            if t:
                obekler.setdefault(t, []).append({"id": oid, "ad": ad})
    vars_ = siralama.varsayilan_kapasite(conn)
    out = []
    for tid, t in sorted(kisi.teknikler(conn).items(), key=lambda x: x[1]["ad"]):
        r = conn.execute("SELECT SUM(durum IN ('atandi','yolda','sahada')), SUM(substr(COALESCE(cozum_zamani,''),1,10)=?), "
                         "SUM(durum IN ('atandi','yolda','sahada') AND serit='BTK') FROM is_emri WHERE atanan_id=?",
                         (gun, tid)).fetchone()
        out.append({**kisi.ozet(t), "etiket": t.get("etiket") or [], "boss_ekip": t.get("boss_ekip"),
                    "kapasite": siralama.kisi_kapasitesi(conn, t, vars_),
                    "bugun": {"atanan": r[0] or 0, "biten": r[1] or 0, "btk": r[2] or 0},
                    "obekler": obekler.get(tid, []), "bugun_yok": tid in yok})
    return out


def bugun_yok_yaz(conn: sqlite3.Connection, teknik_id: int, tarih: str | None, yok: bool, k: dict) -> dict:
    """Spec §6.10 "Bugün çalışmıyor": ``ayar.bugun_yok_<gün>`` JSON listesine teknik eklenir/çıkarılır.

    Yalnız o günü etkiler (ertesi gün anahtar kendiliğinden boştur); ``teknik_listesi`` okur.
    """
    gun = tarih or zaman.simdi().strftime("%Y-%m-%d")
    try:
        dt.date.fromisoformat(gun)
    except ValueError:
        raise V2Hata("Tarih AAAA-AA-GG olmalı.", kod="alan_eksik", durum=422)
    if teknik_id not in kisi.teknikler(conn):
        raise V2Hata("Seçilen kişi aktif bir teknik görevli değil.", kod="teknik_gecersiz", durum=422)
    anahtar = f"bugun_yok_{gun}"
    with islem(conn):
        satir = conn.execute("SELECT deger FROM ayar WHERE anahtar=?", (anahtar,)).fetchone()
        try:
            liste = {int(x) for x in json.loads((satir[0] if satir else None) or "[]")}
        except (ValueError, TypeError):
            liste = set()
        (liste.add if yok else liste.discard)(teknik_id)
        conn.execute("INSERT INTO ayar(anahtar, deger) VALUES(?,?) ON CONFLICT(anahtar) DO UPDATE SET deger=excluded.deger",
                     (anahtar, json.dumps(sorted(liste))))
    _yonetim_kaydi(conn, k, "teknik_bugun_yok", str(teknik_id), {"tarih": gun, "yok": bool(yok)})
    return {"teknik_id": teknik_id, "tarih": gun, "bugun_yok": bool(yok)}


def boss_giden_listesi(conn: sqlite3.Connection) -> list[dict]:
    """EK-4: eklentinin (ya da operatörün) BOSS'a birebir yazacağı değerler."""
    b = gorunum.Baglam(conn, None, is_nolar=[])
    out = []
    for r in conn.execute("SELECT * FROM is_emri WHERE boss_bekleyen IS NOT NULL AND boss_islendi IS NULL AND "
                          "durum NOT IN ('kapandi') ORDER BY guncelleme"):
        g = gorunum.boss_giden(conn, r, b)
        if g:
            out.append(g)
    return out


def excel(conn: sqlite3.Connection, k: dict, kova: str, obek_id: int | None) -> tuple[bytes, int]:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    satirlar, b = gorunum.satirlari_oku(conn, k, kova, 30)
    isler = [gorunum.satir(conn, r, k, b) for r in satirlar]
    if obek_id:
        isler = [x for x in isler if (x.get("obek") or {}).get("id") == obek_id]
    musteri = b.ofis
    kol = ["Task No", "Kaynak", "Task Adı", "Şerit", "Durum", "Öbek", "İl", "İlçe", "Mahalle"]
    if musteri:
        kol += ["Müşteri Adı", "Müşteri No", "Adres"]
    kol += ["Açılış", "24 saat", "Kalan (dk)", "Randevu", "Teknisyen", "BOSS Ekip", "Ticket", "Rozetler"]
    adresler = {r["is_no"]: r["adres"] for r in satirlar} if musteri else {}
    wb = Workbook()
    ws = wb.active
    ws.title = "İşler"
    ws.append(kol)
    for c in ws[1]:
        c.font = Font(bold=True)
    for x in isler:
        rv = x.get("randevu")
        satir = [x["is_no"], "Bayi" if x["kaynak"] == "bayi" else "BOSS", x["task_adi"], x["serit"],
                 akis.ETIKET[x["durum"]], (x.get("obek") or {}).get("ad"), x["il"], x["ilce"], x["mahalle"]]
        if musteri:
            satir += [x.get("musteri_adi"), x.get("musteri_no"), adresler.get(x["is_no"])]
        satir += [x["acilis"], x["son24"], x["kalan_dk"], f"{rv['bas'][:16]}–{rv['bit'][11:16]}" if rv else None,
                  (x.get("atanan") or {}).get("ad"), x.get("boss_ekip"),
                  f"#{x['ticket']['id']} {x['ticket']['konu']}" if x.get("ticket") else None, ", ".join(x["rozetler"])]
        ws.append(satir)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for j, ad in enumerate(kol, start=1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = min(40, max(10, len(ad) + 4))
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue(), len(isler)


def _dagit_gruplari(conn: sqlite3.Connection, obek_id: int, teknik_idler: list[int]) -> list[dict]:
    """Ekiplere dağıt: öbeğin atanmamış açık işleri yakınlığa göre dengeli parçalara (aynı bina bölünmez)."""
    import numpy as np
    from saha import ayarlar
    from operasyon import yakinlik
    if not teknik_idler:
        raise V2Hata("Teknisyen seçin.", kod="alan_eksik", durum=422)
    for t in teknik_idler:
        akis.teknik_denetle(conn, t)
    isler = [dict(r) for r in conn.execute(
        "SELECT is_no, lat, lon, bina_serial, serit, surum FROM is_emri WHERE COALESCE(obek_elle_id, obek_id)=? AND "
        "atanan_id IS NULL AND durum IN ('bekliyor','randevulu') ORDER BY is_no", (obek_id,))]
    if not isinstance(teknik_idler, list):
        teknik_idler = list(teknik_idler)
    etiket = yakinlik.dengeli_bol([x["lat"] if x["lat"] is not None else np.nan for x in isler],
                                  [x["lon"] if x["lon"] is not None else np.nan for x in isler], len(teknik_idler),
                                  grup=[x["bina_serial"] or f"{x['lat']},{x['lon']}" for x in isler]) \
        if isler else []
    gruplar = []
    for i, tid in enumerate(teknik_idler, start=1):
        secili = [x for x, e in zip(isler, etiket) if int(e) == i or (int(e) == 0 and i == 1)]
        km = 0.0
        lat, lon = ayarlar.OFIS["lat"], ayarlar.OFIS["lon"]
        kalan = [x for x in secili if x["lat"] is not None]
        while kalan:                                   # en yakın komşu turu: kaba yol tahmini
            d = [float(yakinlik.yol_km_tahmini(lat, lon, x["lat"], x["lon"])) for x in kalan]
            j = int(np.argmin(d))
            km += d[j]
            lat, lon = kalan[j]["lat"], kalan[j]["lon"]
            kalan.pop(j)
        gruplar.append({"teknik_id": tid, "is_nolar": [x["is_no"] for x in secili],
                        "btk": sum(1 for x in secili if x["serit"] == "BTK"), "km": round(km, 1)})
    return gruplar


def _dagit_uygula(conn, obek_id: int, teknik_idler: list[int], surumler: dict[str, int] | None, k) -> dict:
    gruplar = _dagit_gruplari(conn, obek_id, teknik_idler)
    o = obek.getir(conn, obek_id)
    simdi = zaman.simdi()
    yazilan, atlanan = [], []
    with islem(conn):
        for i, g in enumerate(gruplar, start=1):
            doluluk: dict[str, int] = {}
            for no in g["is_nolar"]:
                r = conn.execute("SELECT surum FROM is_emri WHERE is_no=?", (no,)).fetchone()
                if surumler is not None and (no not in surumler or r is None or int(surumler[no]) != r[0]):
                    atlanan.append({"is_no": no, "kod": "guncel_degil"})
                    continue
                bas, bit = siralama.dilim_oner(conn, g["teknik_id"], simdi, haric=no, ek_doluluk=doluluk)
                doluluk[bas[:13]] = doluluk.get(bas[:13], 0) + 1
                conn.execute("UPDATE is_emri SET oneri_teknik_id=?, oneri_bas=?, oneri_bit=?, oneri_neden=?, parca=?, "
                             "parca_tarih=? WHERE is_no=?", (g["teknik_id"], bas, bit, f"Ekiplere dağıtıldı · {o['ad']}-{i}",
                                                            f"{o['ad']}-{i}", simdi.strftime("%Y-%m-%d"), no))
                yazilan.append(no)
    return {"gruplar": gruplar, "yazilan": yazilan, "atlanan": atlanan}


def ayar_yaz(conn: sqlite3.Connection, girdi: IsAyarGirdi, k) -> dict:
    degisen = {}
    with islem(conn):
        def yaz(a, v):
            conn.execute("INSERT INTO ayar(anahtar, deger) VALUES(?,?) ON CONFLICT(anahtar) DO UPDATE SET "
                         "deger=excluded.deger", (a, v))
            degisen[a] = True
        if girdi.btk_durduran_aski is not None:
            yaz("btk_durduran_aski", json.dumps([x for x in girdi.btk_durduran_aski if x and x.strip()], ensure_ascii=False))
        if girdi.son24_askida_durur is not None:
            yaz("son24_askida_durur", "true" if girdi.son24_askida_durur else "false")
        if girdi.kutlamalar is not None:
            yaz("kutlamalar", "acik" if girdi.kutlamalar else "kapali")
        if girdi.aski_nedenleri is not None:
            temiz = [{"kod": str(n.get("kod") or "")[:30], "metin": str(n.get("metin") or "")[:120]}
                     for n in girdi.aski_nedenleri if n.get("metin")]
            yaz("aski_nedenleri", json.dumps(temiz, ensure_ascii=False))
        if girdi.boss_task_url is not None:
            yaz("boss_task_url", girdi.boss_task_url.strip())
        if girdi.musteri_tel is not None:
            if girdi.musteri_tel not in ("acik", "kapali"):
                raise V2Hata("musteri_tel 'acik' ya da 'kapali' olmalı.", kod="alan_eksik", durum=422)
            yaz("musteri_tel", girdi.musteri_tel)
        if girdi.dilimler is not None:
            yaz("dilimler", girdi.dilimler.strip())
        if girdi.teknik_kapasite is not None:
            yaz("teknik_kapasite", str(girdi.teknik_kapasite))
        if girdi.filtre_istisnalari is not None:
            yaz("filtre_istisnalari", json.dumps([ie.sade(x) for x in girdi.filtre_istisnalari if ie.sade(x)]))
        if girdi.boss_yazma is not None:
            if girdi.boss_yazma not in ("acik", "kapali"):
                raise V2Hata("boss_yazma 'acik' ya da 'kapali' olmalı.", kod="alan_eksik", durum=422)
            degisir = girdi.boss_yazma != ("acik" if boss_yazma.yazma_acik(conn) else "kapali")
            if degisir and not (k or {}).get("administrator"):   # genel anahtar yalnız Administrator'da
                raise V2Hata("BOSS'a yazma anahtarını yalnız Administrator değiştirir.", kod="yasak", durum=403)
            yaz("boss_yazma", girdi.boss_yazma)
        otomatik = {a: getattr(girdi, a) for a in ("iliskili_dahil", "ticket_otomatik_aski", "ticket_cozum_yeniden_ata",
                                                    "ticket_cozum_slot", "boss_yazma_paralel")
                    if getattr(girdi, a) is not None}
        if otomatik and not (k or {}).get("administrator"):
            raise V2Hata("Otomatik planlama ayarlarını yalnız Administrator değiştirir.", kod="yasak", durum=403)
        for a in ("iliskili_dahil", "ticket_otomatik_aski", "ticket_cozum_yeniden_ata"):
            if a in otomatik:
                yaz(a, "acik" if otomatik[a] else "kapali")
        if "ticket_cozum_slot" in otomatik:
            from saha import ticket_boss_aski as tba
            sl = tba.slot_coz(otomatik["ticket_cozum_slot"])
            if sl is None:
                raise V2Hata("Dilim \"08:00-10:00\" biçiminde olmalı (06:00–22:00).", kod="alan_eksik", durum=422)
            yaz("ticket_cozum_slot", f"{sl[0]}-{sl[1]}")
        if "boss_yazma_paralel" in otomatik:                                # 1–10 (model sınırı)
            yaz("boss_yazma_paralel", str(max(1, min(boss_yazma.PARALEL_EN_COK, int(otomatik["boss_yazma_paralel"])))))
        for eylem, ad in (("not", "not"), ("atama", "atama (randevu)")):     # eylemin kendi anahtarı
            deger = getattr(girdi, f"boss_yazma_{eylem}")
            if deger is None:
                continue
            if deger not in ("acik", "kapali"):
                raise V2Hata(f"boss_yazma_{eylem} 'acik' ya da 'kapali' olmalı.", kod="alan_eksik", durum=422)
            degisir = deger != ("acik" if boss_yazma.eylem_acik(conn, eylem) else "kapali")
            if degisir and not (k or {}).get("administrator"):
                raise V2Hata(f"BOSS'a {ad} yazma anahtarını yalnız Administrator değiştirir.", kod="yasak", durum=403)
            yaz(f"boss_yazma_{eylem}", deger)
        if girdi.durum_renkleri is not None:
            if not (k or {}).get("administrator"):
                raise V2Hata("Durum renklerini yalnız Administrator değiştirir.", kod="yasak", durum=403)
            yanlis = {a: v for a, v in girdi.durum_renkleri.items()
                      if a not in DURUM_RENK_VARSAYILAN or v not in DURUM_RENKLERI}
            if yanlis:
                raise V2Hata(f"Tanınmayan durum/renk: {', '.join(sorted(yanlis))}", kod="alan_eksik", durum=422)
            yaz("durum_renkleri", json.dumps({**DURUM_RENK_VARSAYILAN, **girdi.durum_renkleri}, ensure_ascii=False))
        if "btk_durduran_aski" in degisen:
            aski.kurali_yeniden_uygula(conn)
    _yonetim_kaydi(conn, k, "ayar", "is_emri", {"anahtarlar": sorted(degisen)})
    return is_ayarlari(conn)


def kurallari_yaz(conn: sqlite3.Connection, girdi: dict, k) -> dict:
    """PUT /api/isler/kurallar: ``{anahtar: değer | null}`` (null → varsayılana dön). Biçim denetlenir; değişen kural
    açık işlere yeniden uygulanır (şerit/hedef saatleri, askı durdurma). Yönetim kaydına yalnız anahtar adları yazılır."""
    if not isinstance(girdi, dict) or not girdi:
        raise V2Hata("Değiştirilecek kuralı gönderin.", kod="alan_eksik", durum=422)
    temiz: dict[str, object] = {}
    for a, v in girdi.items():
        if a not in kurallar.VARSAYILAN:
            raise V2Hata(f"Tanınmayan kural: {a}", kod="alan_eksik", durum=422)
        try:
            temiz[a] = None if v is None else kurallar.dogrula(a, v)
        except (ValueError, TypeError) as e:
            raise V2Hata(str(e), kod="alan_eksik", durum=422)
    with islem(conn):
        for a, v in temiz.items():
            if v is None:
                conn.execute("DELETE FROM ayar WHERE anahtar=?", (a,))
            else:
                conn.execute("INSERT INTO ayar(anahtar, deger) VALUES(?,?) ON CONFLICT(anahtar) DO UPDATE SET "
                             "deger=excluded.deger", (a, json.dumps(v, ensure_ascii=False)))
        if {"btk_durduran_aski", "aski_gecerlilik"} & set(temiz):
            aski.kurali_yeniden_uygula(conn)
        if {"serit_kurallari", "hedef_saatleri"} & set(temiz):
            akis.hedefleri_yeniden_hesapla(conn)
    _yonetim_kaydi(conn, k, "ayar", "is_kurallari", {"anahtarlar": sorted(temiz)})
    if {"oncelikli_isler", "serit_kurallari", "hedef_saatleri"} & set(temiz):
        siralama.oneri_hesapla(conn)
    return kurallar.hepsi(conn)


async def _dosya(istek: Request) -> tuple[bytes, str, str | None]:
    """Multipart (``dosya`` + ``dosya_zamani``) ya da ham gövde (X-Dosya-Adi, X-Dosya-Zamani başlıkları)."""
    uzunluk = int(istek.headers.get("content-length") or 0)
    if uzunluk > aktarim.AZAMI_BOYUT + 1024 * 64:
        raise HTTPException(413, {"hata": "Dosya çok büyük (en çok 25 MB).", "kod": "buyuk_dosya"})
    tur = (istek.headers.get("content-type") or "").lower()
    if tur.startswith("multipart/form-data"):
        form = await istek.form()
        f = form.get("dosya") or next((v for v in form.values() if hasattr(v, "read")), None)
        if f is None or not hasattr(f, "read"):
            raise HTTPException(422, {"hata": "Formda 'dosya' alanı yok.", "kod": "alan_eksik"})
        veri = await f.read()
        return veri, getattr(f, "filename", "") or "rapor.xlsx", (form.get("dosya_zamani") or None)
    from urllib.parse import unquote
    return await istek.body(), unquote(istek.headers.get("x-dosya-adi", "") or "rapor.xlsx"), \
        istek.headers.get("x-dosya-zamani")


# ============================================================================= yaşam döngüsü (EK-10)
def _temizligi_calistir() -> None:
    """Tek seferlik "tek iş havuzu" temizliği (``temizlik.py``): yalnız canlı veride (ya da
    ``SAHA_TEK_HAVUZ_TEMIZLIGI=1``); ``=0`` kapatır. Sonuç canlı olay akışına yazılır."""
    import os
    import sys
    secim = (os.environ.get("SAHA_TEK_HAVUZ_TEMIZLIGI") or "").strip()
    if secim == "0":
        return
    if secim != "1":
        import yollar
        if not yollar.canli_veri_mi():
            return
    ana = sys.modules.get("saha.api")
    if ana is not None and getattr(ana, "GUNCELLEME_BEKLIYOR", False):
        return
    from saha import db
    from . import temizlik
    yol = db.db_yolu()
    if not yol.exists():
        return
    conn = db.baglan(yol)
    try:
        if temizlik.yapildi_mi(conn):
            _akis_ekle({"kaynak": "sunucu", "adim": "temizlik", "durum": "zaten_yapildi"})
            return
        sonuc = temizlik.calistir(conn, yol, sozluk.obek_dosyasi())
    finally:
        conn.close()
    if sonuc is not None:
        _akis_ekle({"kaynak": "sunucu", "adim": "temizlik", "durum": "yapildi", **sonuc})


def _ekip_uzlasmasi() -> dict | None:
    """Tek seferlik açılış uzlaşması (``ekip_uzlas.calistir``): uygulama ataması BOSS ekibine uyar; BOSS'ta ekibi
    boş iş (askı/merkeze/ertelendi) ekipten alınır. Yalnız canlı veride (ya da ``SAHA_EKIP_UZLASMA=1``); ``=0``
    kapatır. Sonuç (sayılar + yedek) canlı olay akışına yazılır."""
    import os
    secim = (os.environ.get("SAHA_EKIP_UZLASMA") or "").strip()
    if secim == "0":
        return None
    if secim != "1":
        import yollar
        if not yollar.canli_veri_mi():
            return None
    from saha import db
    from . import ekip_uzlas
    yol = db.db_yolu()
    if not yol.exists():
        return None
    conn = db.baglan(yol)
    try:
        if ekip_uzlas.yapildi_mi(conn):
            _akis_ekle({"kaynak": "sunucu", "adim": "ekip_uzlasma", "durum": "zaten_yapildi"})
            return None
        sonuc = ekip_uzlas.calistir(conn, yol)
    finally:
        conn.close()
    if sonuc is not None:
        _akis_ekle({"kaynak": "sunucu", "adim": "ekip_uzlasma", "durum": "yapildi", **sonuc})
    return sonuc


def _izlemeyi_baslat():
    """Rapor klasörü izleyicisi: yalnız iş tabloları kuruluysa ve sunucu güncelleme beklemiyorsa (sürüm uyumsuzsa
    v2 uçları zaten 503'tür). Testlerde, geliştirme/karalama verisinde ve ``SAHA_KLASOR_IZLEME=0`` iken
    ``izleme.baslat`` iş parçacığı açmaz (``izleme.surec_izni``)."""
    import logging
    import sys
    from saha import db
    ana = sys.modules.get("saha.api")
    if ana is not None and getattr(ana, "GUNCELLEME_BEKLIYOR", False):
        return None
    yol = db.db_yolu()
    if not yol.exists():                   # veritabanı yoksa izlenecek bir şey de yok (dosya burada YARATILMAZ)
        return None
    conn = db.baglan(yol)
    try:
        if conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='is_emri'").fetchone() is None:
            return None
    finally:
        conn.close()
    iz = izleme.baslat()
    if iz is not None:
        logging.getLogger("operasyon.izleme").info("Rapor klasörü izleniyor (15 sn'de bir; ayardan kapatılır).")
    return iz


@asynccontextmanager
async def _yasam(_uygulama):
    """Sunucunun kendi yaşam döngüsünün İÇİNDE koşar (FastAPI yönlendirici yaşam döngülerini iç içe birleştirir):
    önce ``saha/api.py`` açılışı (sürüm denetimi), sonra klasör izleme; kapanışta ters sırayla."""
    import logging
    try:
        _temizligi_calistir()
    except Exception as e:     # noqa: BLE001 — temizlik düşerse sunucu yine açılır; bir sonraki açılışta yeniden denenir
        logging.getLogger("operasyon.temizlik").exception("Tek havuz temizliği yapılamadı")
        _akis_ekle({"kaynak": "sunucu", "adim": "temizlik", "durum": "hata", "hata": str(e)[:200]})
    try:                       # tek seferlik: Excel'den gelen ticket'lar silinir (saha/ticket_temizlik.py, yedekli)
        from saha import ticket_temizlik
        sonuc = ticket_temizlik.acilista_calistir()
        if sonuc is not None:
            _akis_ekle({"kaynak": "sunucu", "adim": "ticket_sifirlama", "durum": "yapildi", **sonuc})
    except Exception as e:     # noqa: BLE001 — düşerse sunucu yine açılır; bir sonraki açılışta yeniden denenir
        logging.getLogger("saha.ticket_temizlik").exception("Ticket defteri sıfırlanamadı")
        _akis_ekle({"kaynak": "sunucu", "adim": "ticket_sifirlama", "durum": "hata", "hata": str(e)[:200]})
    try:                       # tek seferlik: uygulama ataması BOSS ekibine uydurulur (ekip_uzlas.py, yedekli)
        _ekip_uzlasmasi()
    except Exception as e:     # noqa: BLE001 — düşerse sunucu yine açılır; bir sonraki açılışta yeniden denenir
        logging.getLogger("operasyon.ekip_uzlas").exception("Ekip uzlaşması yapılamadı")
        _akis_ekle({"kaynak": "sunucu", "adim": "ekip_uzlasma", "durum": "hata", "hata": str(e)[:200]})
    iz = None
    try:
        iz = _izlemeyi_baslat()
    except Exception:          # noqa: BLE001 — izleme açılamazsa sunucu yine açılır; elle yükleme çalışır
        logging.getLogger("operasyon.izleme").exception("Rapor klasörü izleme başlatılamadı")
    try:
        yield
    finally:
        if iz is not None:
            izleme.durdur()


# ============================================================================= yönlendiriciler
def yonlendiriciler(b) -> list[APIRouter]:
    """``b``: saha.yetki.Bagimliliklar (izin, kullanici, baglanti)."""
    r = APIRouter(lifespan=_yasam)
    izin, baglanti = b.izin, b.baglanti

    # ------------------------------------------------------------------ Dağıtım (iş türüne göre, ofis çalışanına)
    @r.get("/api/dagitim")
    def dagitim_liste(istek: Request, hepsi: bool = False, k: dict = Depends(izin("is.ata")),
                      conn: sqlite3.Connection = Depends(baglanti)):
        y = _cagir(dagitim.liste, conn, hepsi)
        if y["satirlar"]:
            _erisim(conn, k, "dagitim_liste", adet=len(y["satirlar"]), ip=_ip(istek))
        return HizliJSON(y)

    @r.put("/api/dagitim/kurallar")
    def dagitim_kurallar(girdi: DagitimKurallariGirdi, k: dict = Depends(izin("is.ata")),
                         conn: sqlite3.Connection = Depends(baglanti)):
        y = _cagir(dagitim.kurallari_yaz, conn, girdi.kurallar, k)
        _yonetim_kaydi(conn, k, "ayar", "dagitim_kurallari", {"kural": len(y)})
        return {"kurallar": y}

    @r.post("/api/dagitim/uygula")
    def dagitim_uygula(k: dict = Depends(izin("is.ata")), conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(dagitim.uygula, conn, k)

    @r.post("/api/dagitim/ata")
    def dagitim_ata(girdi: DagitimAtaGirdi, k: dict = Depends(izin("is.ata")),
                    conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(dagitim.elle_ata, conn, girdi.task_nolar, girdi.kisi_id, k)

    @r.get("/api/isler/gorunumler")
    def isler_gorunumler(k: dict = Depends(izin("is.liste")), conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(dagitim.gorunumler, conn)

    @r.put("/api/isler/gorunumler")
    def isler_gorunumler_yaz(girdi: GorunumlerGirdi, k: dict = Depends(izin("is.ata")),
                             conn: sqlite3.Connection = Depends(baglanti)):
        y = _cagir(dagitim.gorunumleri_yaz, conn, girdi.gorunumler, k)
        _yonetim_kaydi(conn, k, "ayar", "gorunumler", {"gorunum": len(y["gorunumler"])})
        return y

    @r.get("/api/isler/tercih")
    def isler_tercih(k: dict = Depends(izin("is.liste")), conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(tasks_tercih.oku, conn, k)

    @r.put("/api/isler/tercih")
    def isler_tercih_yaz(girdi: dict[str, Any], k: dict = Depends(izin("is.liste")),
                         conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(tasks_tercih.yaz, conn, k, girdi)

    @r.put("/api/isler/tercih/unvan")
    def isler_tercih_unvan(girdi: UnvanDuzenGirdi, k: dict = Depends(izin("is.liste")),
                           conn: sqlite3.Connection = Depends(baglanti)):
        y = _cagir(tasks_tercih.unvan_yaz, conn, k, girdi.unvan_id, girdi.duzen)
        _yonetim_kaydi(conn, k, "ayar", "unvan_duzeni", {"unvan_id": girdi.unvan_id, "kaldir": girdi.duzen is None})
        return y

    # ------------------------------------------------------------------ İşler
    @r.get("/api/isler")
    def isler(istek: Request, kova: str = Query(default="acik", pattern="^(acik|biten)$"),
              gun: int = Query(default=7, ge=1, le=90), ek: int = Query(default=0, ge=0, le=1),
              k: dict = Depends(izin("is.liste")), conn: sqlite3.Connection = Depends(baglanti)):
        _uyananlar(conn)
        y = gorunum.isler_yaniti(conn, k, kova, gun)
        _liste_telefonlari(conn, k, y)
        if ek:
            y["ek_etiketler"] = boss_ham.ekle(conn, y["isler"], _ek_kisisel(conn, k))
        adet = sum(1 for x in y["isler"] if "musteri_adi" in x or "musteri_no" in x)
        if adet:
            _erisim(conn, k, "is_liste", adet=adet, ip=_ip(istek))
        return HizliJSON(y)

    @r.get("/api/isler/degisim")
    def degisim(imlec: int = Query(default=0, ge=0), bekle: int = Query(default=0, ge=0, le=25),
                ek: int = Query(default=0, ge=0, le=1),
                k: dict = Depends(izin("is.liste")), conn: sqlite3.Connection = Depends(baglanti)):
        y = gorunum.degisim(conn, k, imlec)
        _liste_telefonlari(conn, k, y)          # değişen satırlar da telefonla gelsin (yoksa tablodaki numara silinir)
        if ek:
            y["ek_etiketler"] = boss_ham.ekle(conn, y["isler"], _ek_kisisel(conn, k))
        return HizliJSON(y)

    @r.get("/api/isler/excel")
    def isler_excel(istek: Request, kova: str = Query(default="acik", pattern="^(acik|biten)$"),
                    obek_id: int | None = None, k: dict = Depends(izin("is.excel")),
                    conn: sqlite3.Connection = Depends(baglanti)):
        veri, adet = excel(conn, k, kova, obek_id)
        _erisim(conn, k, "excel", adet=adet, ip=_ip(istek))
        ad = f"isler-{zaman.simdi():%Y%m%d-%H%M}.xlsx"
        return Response(veri, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{ad}"', "Cache-Control": "no-store"})

    @r.get("/api/isler/ayarlar")
    def isler_ayarlar(k: dict = Depends(izin("is.liste")), conn: sqlite3.Connection = Depends(baglanti)):
        return is_ayarlari(conn)

    @r.put("/api/isler/ayarlar")
    def isler_ayarlar_yaz(girdi: IsAyarGirdi, k: dict = Depends(izin("veri.yonet")),
                          conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(ayar_yaz, conn, girdi, k)

    @r.get("/api/isler/boss-ekip")
    def boss_ekip(k: dict = Depends(izin("is.ata")), conn: sqlite3.Connection = Depends(baglanti)):
        return boss_ekip_listesi(conn)

    @r.post("/api/isler/boss-ekip/esle")
    def boss_ekip_esle(girdi: BossEkipEsleGirdi, k: dict = Depends(izin("is.ata")),
                       conn: sqlite3.Connection = Depends(baglanti)):
        izinli = getattr(b, "izinli", None)
        if not (izinli(k, "ekip.teknikler") if izinli else kisi.izinli(conn, k, "ekip.teknikler")):
            raise HTTPException(403, {"hata": "Bu bölüm görevinize kapalı.", "kod": "yasak"})
        return _cagir(akis.boss_ekip_esle, conn, girdi.boss_ekip, girdi.teknik_id, k)

    @r.get("/api/isler/teknikler")
    def teknikler(tarih: str | None = None, k: dict = Depends(izin("ekip.teknikler")),
                  conn: sqlite3.Connection = Depends(baglanti)):
        return teknik_listesi(conn, tarih)

    @r.put("/api/isler/teknikler/{teknik_id}/bugun-yok")
    def teknik_bugun_yok(teknik_id: int, girdi: BugunYokGirdi, k: dict = Depends(izin("kapasite.yaz")),
                         conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(bugun_yok_yaz, conn, teknik_id, girdi.tarih, girdi.yok, k)

    @r.post("/api/isler/oneri-onayla")
    def oneri_onayla(girdi: OneriOnayGirdi, k: dict = Depends(izin("is.ata")),
                     conn: sqlite3.Connection = Depends(baglanti)):
        return akis.oneri_onayla(conn, girdi.isler, k)

    @r.post("/api/isler/toplu-ata")
    def toplu_ata(girdi: TopluAtaGirdi, k: dict = Depends(izin("is.ata")), conn: sqlite3.Connection = Depends(baglanti)):
        return akis.toplu_ata(conn, [{"is_no": a.is_no, "teknik_id": a.teknik_id, "surum": a.surum, "yine_de": a.yine_de,
                                      "randevu": a.randevu.model_dump() if a.randevu else None} for a in girdi.atamalar],
                              k, girdi.parca)

    @r.post("/api/isler/toplu/{toplu_id}/geri-al")
    def toplu_geri_al(toplu_id: str, k: dict = Depends(izin("is.ata")), conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.toplu_geri_al, conn, toplu_id, k)

    @r.post("/api/isler/dagit/onizle")
    def dagit_onizle(girdi: DagitGirdi, k: dict = Depends(izin("is.ata")), conn: sqlite3.Connection = Depends(baglanti)):
        return {"gruplar": _cagir(_dagit_gruplari, conn, girdi.obek_id, girdi.teknik_idler)}

    @r.post("/api/isler/dagit/uygula")
    def dagit_uygula(girdi: DagitGirdi, k: dict = Depends(izin("is.ata")), conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(_dagit_uygula, conn, girdi.obek_id, girdi.teknik_idler, girdi.isler, k)

    @r.get("/api/isler/kurallar")
    def isler_kurallar(k: dict = Depends(izin("is.liste")), conn: sqlite3.Connection = Depends(baglanti)):
        return kurallar.hepsi(conn)

    @r.put("/api/isler/kurallar")
    def isler_kurallar_yaz(girdi: dict[str, Any], k: dict = Depends(izin("veri.yonet")),
                           conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(kurallari_yaz, conn, girdi, k)

    @r.post("/api/isler", status_code=201)
    def yeni_is(girdi: YeniIsGirdi, k: dict = Depends(izin("is.olustur")), conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.bayi_olustur, conn, girdi.model_dump(), k)

    # /api/isler/{is_no}'dan ÖNCE (yoksa "ekip-konumlari" iş no sanılır)
    @r.get("/api/isler/ekip-konumlari")
    def isler_ekip_konumlari(k: dict = Depends(izin("is.liste")), conn: sqlite3.Connection = Depends(baglanti)):
        """MAP "Ekipler" katmanı (BOSS GPS): {ekipler: [{ekip, ad, lat, lon, zaman, yas_dk, teknik_id?}],
        kaynak_goruldu, son_alim}. Kaynak: eklenti → POST /api/boss/ekip-konum (ekip_konum.py)."""
        from . import ekip_konum
        return ekip_konum.liste(conn)

    @r.get("/api/isler/{is_no}")
    def is_ayrinti(is_no: str, istek: Request, k: dict = Depends(izin("is.liste")),
                   conn: sqlite3.Connection = Depends(baglanti)):
        a = _cagir(gorunum.ayrinti, conn, is_no, k)
        if "musteri_adi" in a or "musteri_no" in a:
            _erisim(conn, k, "is_ayrinti", is_no=is_no, ip=_ip(istek))
        return HizliJSON(a)

    @r.post("/api/isler/{is_no}/ata")
    def ata(is_no: str, girdi: AtaGirdi, k: dict = Depends(izin("is.ata")), conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.ata, conn, is_no, girdi.teknik_id, k,
                      randevu=girdi.randevu.model_dump() if girdi.randevu else None, surum=girdi.surum,
                      yine_de=girdi.yine_de, notu=girdi.notu)

    @r.post("/api/isler/{is_no}/randevu")
    def randevu(is_no: str, girdi: RandevuGirdi, k: dict = Depends(izin("is.ata", "is.randevu")),
                conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.randevu, conn, is_no, k, girdi.bas, girdi.bit, girdi.teyitli, surum=girdi.surum,
                      notu=girdi.notu)

    @r.post("/api/isler/{is_no}/durum")
    def durum(is_no: str, girdi: DurumGirdi, k: dict = Depends(izin("is.saha", "is.duzenle")),
              conn: sqlite3.Connection = Depends(baglanti)):
        alanlar = {a: v for a, v in girdi.model_dump(exclude={"yeni", "surum", "istemci_id", "zaman"}).items()
                   if v is not None}
        return _cagir(akis.gecis, conn, is_no, girdi.yeni, k, surum=girdi.surum, istemci_id=girdi.istemci_id,
                      zaman_=girdi.zaman, **alanlar)

    @r.post("/api/isler/{is_no}/teshis")
    def teshis(is_no: str, girdi: TeshisGirdi, k: dict = Depends(izin("is.duzenle")),
               conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.teshis, conn, is_no, k, girdi.sonuc, canli_test=girdi.canli_test,
                      randevu_=girdi.randevu.model_dump() if girdi.randevu else None, surum=girdi.surum)

    @r.post("/api/isler/{is_no}/ticket")
    def ticket_bagla(is_no: str, girdi: TicketGirdi, k: dict = Depends(izin("is.duzenle")),
                     conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.ticket_bagla, conn, is_no, k, ticket_id=girdi.ticket_id,
                      yeni=girdi.yeni.model_dump() if girdi.yeni else None,
                      teknisyen_gonderme=girdi.teknisyen_gonderme, surum=girdi.surum)

    @r.delete("/api/isler/{is_no}/ticket")
    def ticket_ayir(is_no: str, surum: int | None = None, k: dict = Depends(izin("is.duzenle")),
                    conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.ticket_ayir, conn, is_no, k, surum)

    @r.post("/api/isler/{is_no}/obek")
    def is_obek(is_no: str, girdi: ObekGirdi, k: dict = Depends(izin("is.duzenle")),
                conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.obek_ata, conn, is_no, girdi.obek_id, k, girdi.surum)

    @r.post("/api/isler/{is_no}/mahalle")
    def is_mahalle(is_no: str, girdi: MahalleGirdi, k: dict = Depends(izin("is.duzenle")),
                   conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.mahalle_sec, conn, is_no, girdi.mahalle_id, k, girdi.surum)

    @r.post("/api/isler/{is_no}/not", status_code=201)
    def is_not(is_no: str, girdi: NotGirdi, k: dict = Depends(izin("is.duzenle", "is.saha")),
               conn: sqlite3.Connection = Depends(baglanti)):
        return {"olay_id": _cagir(akis.not_ekle, conn, is_no, girdi.metin, k, girdi.istemci_id)}

    @r.patch("/api/isler/{is_no}/iletisim")
    def iletisim(is_no: str, girdi: IletisimGirdi, istek: Request, k: dict = Depends(izin("is.duzenle")),
                 conn: sqlite3.Connection = Depends(baglanti)):
        a = _cagir(akis.iletisim, conn, is_no, girdi.musteri_tel, k, girdi.surum)
        _erisim(conn, k, "tel_goster", is_no=is_no, ip=_ip(istek))
        return a

    @r.post("/api/isler/{is_no}/kopya", status_code=204)
    def kopya(is_no: str, girdi: KopyaGirdi, istek: Request, k: dict = Depends(izin("is.musteri")),
              conn: sqlite3.Connection = Depends(baglanti)):
        r_ = conn.execute("SELECT * FROM is_emri WHERE is_no=?", (is_no,)).fetchone()
        if r_ is None:
            raise _hata(akis._H().IsYok())
        _cagir(akis.kapsam_denetle, conn, k, r_)
        _erisim(conn, k, "kopya_musteri_no", is_no=is_no, ip=_ip(istek))
        return Response(status_code=204)

    @r.post("/api/isler/{is_no}/arama")
    def arama(is_no: str, girdi: AramaGirdi, k: dict = Depends(izin("is.duzenle")),
              conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.arama_kaydet, conn, is_no, k, girdi.sonuc, sure_sn=girdi.sure_sn, webphone=girdi.webphone,
                      sms=girdi.sms, yeni_numara=girdi.yeni_numara, notu=girdi.notu, istemci_id=girdi.istemci_id)

    @r.post("/api/isler/{is_no}/btk-sikayet")
    def btk_sikayet(is_no: str, girdi: BtkSikayetGirdi, k: dict = Depends(izin("is.duzenle")),
                    conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.btk_sikayet_kaydet, conn, is_no, k, tarih=girdi.tarih, reopen=girdi.reopen,
                      kaldir=girdi.kaldir, surum=girdi.surum)

    @r.post("/api/isler/{is_no}/bizim-durum")
    def bizim_durum(is_no: str, girdi: BizimDurumGirdi, k: dict = Depends(izin("is.duzenle")),
                    conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.bizim_durum_yaz, conn, is_no, k, girdi.deger, surum=girdi.surum)

    @r.post("/api/isler/hucre-renk")
    def is_hucre_renk(girdi: HucreRenkGirdi, k: dict = Depends(izin("is.duzenle")),
                      conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(isaret.hucre_yaz, conn, girdi.hucreler, girdi.renk, k)

    @r.post("/api/isler/{is_no}/isaret")
    def is_isaret(is_no: str, girdi: IsaretGirdi, k: dict = Depends(izin("is.duzenle")),
                  conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(isaret.yaz, conn, is_no, girdi.renk, k)

    @r.post("/api/boss/atama-durumlari")
    def boss_atama_durumlari(girdi: AtamaDurumGirdi, k: dict = Depends(izin("is.ata")),
                             conn: sqlite3.Connection = Depends(baglanti)):
        return HizliJSON({"durumlar": boss_yazma.atama_durumlari(conn, girdi.is_nolar)})

    @r.post("/api/boss/atama-yeniden")
    def boss_atama_yeniden(girdi: AtamaDurumGirdi, k: dict = Depends(izin("is.ata")),
                           conn: sqlite3.Connection = Depends(baglanti)):
        """Toplu atama özetindeki "Yeniden dene": BOSS'a yazılamayan atamalar aynı girdiyle yeniden kuyruğa (tekli)."""
        s = boss_yazma.yeniden_dene(conn, girdi.is_nolar, k)
        _akis_ekle({"kaynak": "sunucu", "adim": "atama_yeniden", "yeniden": len(s["yeniden"]),
                    "atlanan": len(s["atlanan"])})
        return s

    @r.get("/api/boss/yazmalar")
    def boss_yazmalar(durum: str | None = None, tur: str | None = None, ara: str | None = None,
                      k: dict = Depends(izin("is.ata")), conn: sqlite3.Connection = Depends(baglanti)):
        """BOSS yazma kuyruğu: son 7 günün kayıtları (zaman, eylem, task no, durum, mesaj, eksik alan; kişisel veri
        yok). Süzgeç: durum, tur (atama/aski/not), ara (task no)."""
        return HizliJSON(boss_yazma.kuyruk(conn, durum=durum, tur=tur, ara=ara))

    @r.post("/api/boss/yazmalar/{kayit_id}/yeniden")
    def boss_yazma_yeniden(kayit_id: int, k: dict = Depends(izin("is.ata")),
                           conn: sqlite3.Connection = Depends(baglanti)):
        """Kuyruk ekranında tek kaydı yeniden dene (hata / zaman aşımı / eksik alan)."""
        s = _cagir(boss_yazma.kayit_yeniden, conn, kayit_id, k)
        _akis_ekle({"kaynak": "sunucu", "adim": "yazma_yeniden", "kayit": kayit_id, "yeni": s.get("yeni_id")})
        return s

    @r.post("/api/isler/toplu-aski/nedenler")
    def toplu_aski_nedenleri(girdi: TopluAskiNedenGirdi, k: dict = Depends(izin("is.duzenle")),
                             conn: sqlite3.Connection = Depends(baglanti)):
        return boss_yazma.toplu_nedenler(conn, girdi.is_nolar)

    @r.post("/api/isler/toplu-aski")
    def toplu_aski(girdi: TopluAskiGirdi, k: dict = Depends(izin("is.duzenle")),
                   conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.toplu_aski, conn, girdi.is_nolar, k, neden=girdi.neden, alt=girdi.alt, bitis=girdi.bitis,
                      aciklama=girdi.aciklama)

    @r.post("/api/isler/iliskili/isler")
    def iliskili_isler(girdi: IliskiliGirdi, k: dict = Depends(izin("is.liste")),
                       conn: sqlite3.Connection = Depends(baglanti)):
        from . import iliskili
        return iliskili.iliskili_isler(conn, girdi.is_nolar)

    @r.get("/api/isler/iliskili/uyumsuz")
    def iliskili_uyumsuz(k: dict = Depends(izin("is.liste")), conn: sqlite3.Connection = Depends(baglanti)):
        from . import iliskili
        return iliskili.uyumsuzlar(conn)

    @r.post("/api/isler/iliskili/hizala")
    def iliskili_hizala(girdi: IliskiliHizalaGirdi, k: dict = Depends(izin("is.ata")),
                        conn: sqlite3.Connection = Depends(baglanti)):
        from . import iliskili
        return _cagir(iliskili.hizala, conn, k, girdi.is_nolar)

    @r.get("/api/isler/{is_no}/boss-yazma-onizleme")
    def boss_yazma_onizleme(is_no: str, k: dict = Depends(izin("is.liste")),
                            conn: sqlite3.Connection = Depends(baglanti)):
        _cagir(gorunum.ayrinti, conn, is_no, k)          # işi görebiliyor mu (teknik: yalnız kendi işi)
        return HizliJSON(boss_yazma.onizleme(conn, is_no))

    @r.post("/api/isler/{is_no}/tcs")
    def tcs(is_no: str, girdi: TcsGirdi, k: dict = Depends(izin("is.duzenle")),
            conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.tcs_gir, conn, is_no, k, girdi.tarih, girdi.teyitli, surum=girdi.surum)

    @r.post("/api/isler/{is_no}/boss-islendi")
    def boss_islendi(is_no: str, girdi: BossIslendiGirdi, k: dict = Depends(izin("is.ata")),
                     conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.boss_islendi, conn, is_no, girdi.islendi, k, girdi.surum)

    @r.post("/api/isler/{is_no}/boss-bagla")
    def boss_bagla(is_no: str, girdi: BossBaglaGirdi, k: dict = Depends(izin("is.ata")),
                   conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.boss_bagla, conn, is_no, girdi.boss_task_no, k, girdi.surum)

    @r.post("/api/isler/{is_no}/geri-al")
    def geri_al(is_no: str, girdi: GeriAlGirdi, k: dict = Depends(izin("is.ata", "is.duzenle", "is.saha")),
                conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.geri_al, conn, is_no, girdi.olay_id, k)

    # ------------------------------------------------------------------ Teknik
    @r.get("/api/islerim")
    def islerim(k: dict = Depends(izin("is.saha")), conn: sqlite3.Connection = Depends(baglanti)):
        return HizliJSON(gorunum.islerim(conn, k))

    @r.post("/api/islerim/gordu", status_code=204)
    def gordu(girdi: GorduGirdi, k: dict = Depends(izin("is.saha")), conn: sqlite3.Connection = Depends(baglanti)):
        akis.teknik_gordu(conn, k, girdi.is_nolar[:200])
        return Response(status_code=204)

    @r.post("/api/islerim/sira")
    def sira(girdi: SiraGirdi, k: dict = Depends(izin("is.saha")), conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.sira_degistir, conn, k, girdi.is_no, girdi.yeni_sira, girdi.neden, girdi.istemci_id)

    # ------------------------------------------------------------------ Aktarım
    @r.post("/api/aktarim")
    async def aktarim_yukle(istek: Request, k: dict = Depends(izin("is.yukle"))):
        veri, ad, dosya_zamani = await _dosya(istek)
        try:
            return await run_in_threadpool(aktarim.aktar, _fabrika, veri, ad, dosya_zamani, k, "surukle")
        except V2Hata as e:
            raise _hata(e)

    @r.post("/api/aktarim/{aktarim_id}/uygula")
    def aktarim_uygula(aktarim_id: int, girdi: AktarimUygulaGirdi, k: dict = Depends(izin("is.yukle"))):
        return _cagir(aktarim.uygula, _fabrika, aktarim_id, girdi.mod, k)

    @r.post("/api/aktarim/{aktarim_id}/vazgec")
    def aktarim_vazgec(aktarim_id: int, k: dict = Depends(izin("is.yukle"))):
        return _cagir(aktarim.vazgec, _fabrika, aktarim_id, k)

    @r.get("/api/aktarim")
    def aktarim_gecmis(limit: int = Query(default=30, ge=1, le=200), k: dict = Depends(izin("is.yukle")),
                       conn: sqlite3.Connection = Depends(baglanti)):
        return aktarim.gecmis(conn, limit)

    @r.get("/api/aktarim/izleme")
    def izleme_durum(k: dict = Depends(izin("is.yukle")), conn: sqlite3.Connection = Depends(baglanti)):
        return izleme.durum(conn)

    @r.put("/api/aktarim/izleme")
    def izleme_yaz(girdi: IzlemeGirdi, k: dict = Depends(izin("is.yukle")),
                   conn: sqlite3.Connection = Depends(baglanti)):
        # Klasör/desen değiştirmek sunucunun okuyacağı yeri değiştirir: yalnız veri yönetimi yetkisiyle.
        if (girdi.klasorler is not None or girdi.desenler is not None) and not kisi.izinli(conn, k, "veri.yonet"):
            raise HTTPException(403, {"hata": "Klasörleri yalnız yönetici değiştirebilir.", "kod": "yasak"})
        izleme.ayar_yaz(conn, acik=girdi.acik, klasorler=girdi.klasorler, desenler=girdi.desenler)
        _yonetim_kaydi(conn, k, "ayar", "klasor_izleme", {"acik": girdi.acik,
                                                          "klasor_degisti": girdi.klasorler is not None})
        return izleme.durum(conn)

    # ------------------------------------------------------------------ Öbekler ve mahalleler
    @r.get("/api/obekler")
    def obekler(k: dict = Depends(izin("obek.oku")), conn: sqlite3.Connection = Depends(baglanti)):
        return HizliJSON(obek.liste(conn))

    @r.post("/api/obekler", status_code=201)
    def obek_olustur(girdi: ObekYeniGirdi, k: dict = Depends(izin("obek.duzenle")),
                     conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(obek.olustur, conn, girdi.ad, k, girdi.surum)

    @r.patch("/api/obekler/{obek_id}")
    def obek_guncelle(obek_id: int, girdi: ObekGuncelleGirdi, k: dict = Depends(izin("obek.duzenle")),
                      conn: sqlite3.Connection = Depends(baglanti)):
        alanlar = girdi.model_dump(exclude={"surum"}, exclude_unset=True)
        return _cagir(obek.guncelle, conn, obek_id, k, girdi.surum, **alanlar)

    @r.post("/api/obekler/{obek_id}/mahalle")
    def obek_mahalle(obek_id: int, girdi: ObekMahalleGirdi, k: dict = Depends(izin("obek.duzenle")),
                     conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(obek.mahalle_ekle, conn, obek_id, k, girdi.surum, mahalle_idler=girdi.mahalle_idler or (),
                      tum_ilce=[t.model_dump() for t in girdi.tum_ilce or ()], tasi=girdi.tasi)

    @r.post("/api/obekler/{obek_id}/mahalle/cikar")
    def obek_cikar(obek_id: int, girdi: ObekCikarGirdi, k: dict = Depends(izin("obek.duzenle")),
                   conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(obek.mahalle_cikar, conn, obek_id, girdi.refler, k, girdi.surum)

    @r.delete("/api/obekler/{obek_id}")
    def obek_sil(obek_id: int, surum: int | None = None, k: dict = Depends(izin("obek.duzenle")),
                 conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(obek.sil, conn, obek_id, k, surum)

    @r.post("/api/obekler/geri-al")
    def obek_geri_al(girdi: ObekGeriAlGirdi, k: dict = Depends(izin("obek.duzenle")),
                     conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(obek.geri_al, conn, girdi.olay_id, k, girdi.surum)

    @r.post("/api/obekler/{obek_id}/bol")
    def obek_bol(obek_id: int, girdi: ObekBolGirdi, k: dict = Depends(izin("obek.duzenle")),
                 conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(obek.bol, conn, obek_id, girdi.k, girdi.onizle, k, girdi.surum)

    @r.get("/api/mahalleler")
    def mahalleler(q: str | None = Query(default=None, max_length=80), il: str | None = None, ilce: str | None = None,
                   limit: int = Query(default=50, ge=1, le=500), k: dict = Depends(izin("obek.oku")),
                   conn: sqlite3.Connection = Depends(baglanti)):
        return HizliJSON(sozluk.ara(conn, q, il, ilce, limit))

    @r.get("/api/ilceler")
    def ilceler(k: dict = Depends(izin("obek.oku")), conn: sqlite3.Connection = Depends(baglanti)):
        return sozluk.ilceler(conn)

    @r.post("/api/mahalleler", status_code=201)
    def mahalle_ekle(girdi: MahalleYeniGirdi, k: dict = Depends(izin("mahalle.ekle")),
                     conn: sqlite3.Connection = Depends(baglanti)):
        def ekle():
            with islem(conn):
                m = sozluk.ekle(conn, girdi.il, girdi.ilce, girdi.ad, k, girdi.benzerine_ragmen)
                r_ = sozluk.getir(conn, m["id"])
                # "Yeni mahalle" dendi: o yazımla Kontrol'de bekleyen işler yeniden çözülür
                nolar = [x[0] for x in conn.execute(
                    "SELECT is_no FROM is_emri WHERE durum<>'kapandi' AND il_k=? AND ilce_k=? AND mahalle_k=?",
                    (r_["il_k"], r_["ilce_k"], r_["mahalle_k"]))]
                degisen = obek.yeniden_coz(conn, nolar, k)
            if degisen:
                siralama.oneri_hesapla(conn, degisen)
            return {"mahalle": m, "etkilenen_is": len(degisen)}
        return _cagir(ekle)

    @r.post("/api/mahalleler/esad", status_code=201)
    def mahalle_esad(girdi: EsadGirdi, k: dict = Depends(izin("mahalle.ekle")),
                     conn: sqlite3.Connection = Depends(baglanti)):
        def esad():
            with islem(conn):
                e = sozluk.esad_ekle(conn, girdi.il, girdi.ilce, girdi.esad, girdi.mahalle_id, k)
                nolar = [x[0] for x in conn.execute(
                    "SELECT is_no FROM is_emri WHERE durum<>'kapandi' AND il_k=? AND ilce_k=? AND mahalle_k=?",
                    (e["il_k"], e["ilce_k"], e["esad_k"]))]
                degisen = obek.yeniden_coz(conn, nolar, k)
            if degisen:
                siralama.oneri_hesapla(conn, degisen)
            return {**e, "etkilenen_is": len(degisen)}
        return _cagir(esad)

    @r.post("/api/mahalleler/yukle")
    async def mahalle_yukle(istek: Request, k: dict = Depends(izin("mahalle.yukle"))):
        import pandas as pd
        veri, ad, _ = await _dosya(istek)

        def isle():
            conn = _fabrika()
            try:
                if ad.lower().endswith(".csv"):
                    df = pd.read_csv(io.BytesIO(veri), dtype=str)
                else:
                    df = pd.read_excel(io.BytesIO(veri), dtype=str)
                with islem(conn):
                    s = sozluk.yukle(conn, df, k)
                _yonetim_kaydi(conn, k, "mahalle_yukle", None, {"eklenen": s["eklenen"]})
                return s
            finally:
                conn.close()
        try:
            return await run_in_threadpool(isle)
        except V2Hata as e:
            raise _hata(e)
        except (ValueError, OSError):
            raise HTTPException(422, {"hata": "Dosya okunamadı (xlsx ya da csv: il, ilce, mahalle).", "kod": "alan_eksik"})

    # ------------------------------------------------------------------ Takip
    @r.get("/api/takip")
    def takip_(gun: int = Query(default=7, ge=1, le=90), k: dict = Depends(izin("takip.oku")),
               conn: sqlite3.Connection = Depends(baglanti)):
        return HizliJSON(takip.takip(conn, zaman.simdi(), gun, tam=kisi.izinli(conn, k, "takip.tam")))

    @r.get("/api/takip/satis")
    def takip_satis(gun: int = Query(default=7, ge=1, le=90), k: dict = Depends(izin("takip.tam")),
                    conn: sqlite3.Connection = Depends(baglanti)):
        return takip.satis(conn, zaman.simdi(), gun)

    @r.get("/api/takip/erisim")
    def takip_erisim(gun: int = Query(default=7, ge=1, le=366), k: dict = Depends(izin("takip.tam")),
                     conn: sqlite3.Connection = Depends(baglanti)):
        return takip.erisim(conn, zaman.simdi(), gun)

    @r.get("/api/takip/yonetim")
    def takip_yonetim(limit: int = Query(default=100, ge=1, le=1000), k: dict = Depends(izin("takip.tam")),
                      conn: sqlite3.Connection = Depends(baglanti)):
        return takip.yonetim(conn, limit)

    @r.put("/api/takip/kapasite")
    def takip_kapasite(girdi: KapasiteGirdi, k: dict = Depends(izin("kapasite.yaz")),
                       conn: sqlite3.Connection = Depends(baglanti)):
        gun = girdi.tarih or zaman.simdi().strftime("%Y-%m-%d")
        try:
            dt.date.fromisoformat(gun)
        except ValueError:
            raise HTTPException(422, {"hata": "Tarih AAAA-AA-GG olmalı.", "kod": "alan_eksik"})
        with islem(conn):
            if girdi.aktif_teknik is None:
                conn.execute("DELETE FROM ayar WHERE anahtar=?", (f"kapasite_{gun}",))
            else:
                conn.execute("INSERT INTO ayar(anahtar, deger) VALUES(?,?) ON CONFLICT(anahtar) DO UPDATE SET "
                             "deger=excluded.deger", (f"kapasite_{gun}", str(girdi.aktif_teknik)))
        siralama.oneri_hesapla(conn)
        return siralama.mod(conn, zaman.simdi())

    # ------------------------------------------------------------------ EK-12.6: kesintiye duyarlı sevk
    @r.get("/api/kesinti")
    def kesinti_liste(k: dict = Depends(izin("is.duzenle")), conn: sqlite3.Connection = Depends(baglanti)):
        return kesinti.liste(conn)

    @r.post("/api/kesinti", status_code=201)
    def kesinti_ekle(girdi: KesintiGirdi, k: dict = Depends(izin("is.duzenle")),
                     conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(kesinti.ekle, conn, girdi.metin, k)

    @r.get("/api/kesinti/aday")
    def kesinti_aday(k: dict = Depends(izin("is.duzenle")), conn: sqlite3.Connection = Depends(baglanti)):
        return kesinti.adaylar(conn, zaman.simdi(), musteri=kisi.izinli(conn, k, "is.musteri"))

    @r.post("/api/kesinti/{kesinti_id}/bitti")
    def kesinti_bitti(kesinti_id: int, k: dict = Depends(izin("is.duzenle")),
                      conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(kesinti.bitir, conn, kesinti_id, k)

    # ------------------------------------------------------------------ EK-4: BOSS giden kutusu
    @r.get("/api/boss/giden")
    def boss_giden(k: dict = Depends(izin("is.ata")), conn: sqlite3.Connection = Depends(baglanti)):
        return boss_giden_listesi(conn)

    @r.post("/api/boss/giden/{is_no}/islendi")
    def boss_giden_islendi(is_no: str, k: dict = Depends(izin("is.ata")), conn: sqlite3.Connection = Depends(baglanti)):
        return _cagir(akis.boss_islendi, conn, is_no, True, k)

    # ------------------------------------------------------------------ BOSS canlı eşitleme (eklenti)
    kullanici_al, izinli_mi = b.kullanici, b.izinli

    def _esitle_kimlik(conn: sqlite3.Connection = Depends(baglanti),
                       authorization: str | None = Header(default=None),
                       x_eslesme_kodu: str | None = Header(default=None)) -> dict:
        """Eklenti: ``X-Eslesme-Kodu``; yoksa rapor yükleme izni olan oturum (is.yukle)."""
        if x_eslesme_kodu:
            if boss.kod_dogru(x_eslesme_kodu):
                return {"id": None, "ad": "BOSS Yardımcısı", "eklenti": True}
            raise HTTPException(401, {"hata": "Eşleme kodu yanlış. Uygulamadaki BOSS bağlantısı ekranındaki kodu "
                                              "girin.", "kod": "eslesme_yanlis"})
        k = kullanici_al(conn, authorization)
        if not izinli_mi(k, "is.yukle"):
            raise HTTPException(403, {"hata": "Bu bölüm görevinize kapalı.", "kod": "yasak"})
        return k

    _esitle_kimlik.eylemler = ("is.yukle",)

    @r.post("/api/boss/gunluk")
    async def boss_gunluk(istek: Request, k: dict = Depends(_esitle_kimlik)):
        """Eklentinin tanı olayı (değer içermez: adım, uç yolu, HTTP kodu). Diske yazılmaz; bellekte son 500 olay.
        Tek istisna ``aski_nedenleri``: BOSS'un askı nedeni listesi (kişisel veri değil) saklanır, akışa sayısı yazılır."""
        try:
            govde = await istek.json()
        except ValueError:
            govde = {}
        govde = govde if isinstance(govde, dict) else {}
        if govde.get("adim") in ("aski_nedenleri", "aski_alt_nedenleri"):
            ek_ = govde.get("ek")
            n = await run_in_threadpool(_nedenleri_sakla, govde.get("yol"), ek_.get("veri") if isinstance(ek_, dict) else None)
            govde = {"adim": govde.get("adim"), "yol": govde.get("yol"), "sayi": n, "surum": govde.get("surum")}
        elif govde.get("adim") == "deger_ornegi" and str(govde.get("yol") or "").endswith("/task/suspend"):
            ek_ = govde.get("ek")
            await run_in_threadpool(_aski_ornegi_sakla, govde.get("yol"),
                                    ek_.get("govde") if isinstance(ek_, dict) else None)
        _akis_ekle({"kaynak": "eklenti", **govde})
        return {"ok": True}

    @r.get("/api/boss/yazilacaklar")
    def boss_yazilacaklar(k: dict = Depends(_esitle_kimlik), conn: sqlite3.Connection = Depends(baglanti)):
        """Eklenti 2–3 sn'de bir sorar: BOSS'a yazılacaklar. ``yazma='kapali'`` (varsayılan) iken eklenti yalnız
        önizleme der, BOSS'a hiçbir istek göndermez."""
        y = boss_yazma.verilecekler(conn)
        for i in y["istekler"]:
            _akis_ekle({"kaynak": "sunucu", "adim": "yazma_verildi", "tur": i["tur"], "kayit": len(i["kayitlar"]),
                        "yaz": i["yaz"], "deneme": bool(i.get("deneme"))})
        for yol in y.get("okunacaklar") or ():
            _akis_ekle({"kaynak": "sunucu", "adim": "aski_nedeni_istendi", "yol": yol})
        return HizliJSON(y)

    @r.post("/api/boss/yazma-sonucu")
    async def boss_yazma_sonucu(istek: Request, k: dict = Depends(_esitle_kimlik)):
        """Eklentinin yürütme sonucu (önizleme / gönderildi / hata; yanıt gövdesi gelmez)."""
        try:
            govde = await istek.json()
        except ValueError:
            raise HTTPException(422, {"hata": "Gövde JSON olmalı.", "kod": "alan_eksik"})

        def isle():
            conn = _fabrika()
            try:
                return _cagir(boss_yazma.sonuc_al, conn, govde)
            finally:
                conn.close()
        s = await run_in_threadpool(isle)
        g_ = govde if isinstance(govde, dict) else {}
        _akis_ekle({"kaynak": "sunucu", "adim": "yazma_sonucu", "durum": g_.get("durum"), "http": g_.get("http"),
                    "return_kodu": g_.get("return_kodu"), "guncellenen": s["guncellenen"]})
        return s

    @r.put("/api/boss/yazma-anahtari")
    def boss_yazma_anahtari(girdi: BossYazmaAnahtarGirdi, k: dict = Depends(izin("yetkinlik.yonet")),
                            conn: sqlite3.Connection = Depends(baglanti)):
        """BOSS'a yazma genel anahtarı (yalnız Administrator). Açıkken yalnız DOĞRULANMIŞ şablonlar gönderilir."""
        s = boss_yazma.yazma_ayarla(conn, girdi.acik)
        _yonetim_kaydi(conn, k, "boss_yazma_anahtari", "boss", {"acik": girdi.acik})
        _akis_ekle({"kaynak": "sunucu", "adim": "yazma_anahtari", "acik": girdi.acik})
        return s

    @r.post("/api/boss/yazma-dene")
    def boss_yazma_dene(girdi: BossYazmaDeneGirdi, k: dict = Depends(izin("yetkinlik.yonet")),
                        conn: sqlite3.Connection = Depends(baglanti)):
        """Tek iş denemesi (yalnız Administrator): bu işin bu eylemdeki son kaydı, genel anahtar kapalıyken bile
        eklentiden BİR KEZ gerçekten gönderilir. Sonuç: GET /api/isler/{is_no}/boss-yazma-onizleme ve /api/boss/durum."""
        s = _cagir(boss_yazma.deneme_baslat, conn, girdi.is_no, girdi.eylem, k, girdi.aski_neden_id,
                   girdi.aski_alt_neden_id)
        _yonetim_kaydi(conn, k, "boss_yazma_dene", girdi.is_no, {"eylem": girdi.eylem, "kayit_id": s["kayit_id"]})
        _akis_ekle({"kaynak": "sunucu", "adim": "yazma_deneme", "tur": girdi.eylem, "kayit": s["kayit_id"]})
        return s

    @r.get("/api/isler/{is_no}/boss-aski-nedenleri")
    def is_boss_aski_nedenleri(is_no: str, k: dict = Depends(izin("is.duzenle")),
                               conn: sqlite3.Connection = Depends(baglanti)):
        """Bu işin BOSS askı nedenleri (işin kategori + ürün koduna göre; yoksa eklentiden okunur)."""
        return boss_yazma.is_nedenleri(conn, is_no)

    @r.get("/api/boss/farklar")
    def boss_farklar(k: dict = Depends(izin("is.liste"))):
        """Değişiklik akışı: son 50 eşitlemenin yeni/değişen/kapanan task noları (en yenisi önde)."""
        return HizliJSON({"farklar": list(reversed(_FARKLAR))})

    @r.get("/api/boss/aski-nedenleri")
    def boss_aski_nedenleri(k: dict = Depends(izin("is.duzenle")), conn: sqlite3.Connection = Depends(baglanti)):
        """BOSS'un askı nedeni listesi (kanca BOSS'ta askı penceresi açılınca yakalar): [{id, ad, alt}]."""
        return boss_yazma.nedenler(conn)

    @r.get("/api/boss/akis")
    async def boss_akis(istek: Request, k: dict = Depends(_esitle_kimlik)):
        """Canlı olay akışı (text/event-stream): eklenti ve eşitleme olayları anlık; yalnız yerel ağ."""
        from fastapi.responses import StreamingResponse
        import asyncio

        async def uret():
            son = 0
            ayrac = "\n\n"
            for o in list(_AKIS):
                son = o["no"]
                yield f"data: {json.dumps(o, ensure_ascii=False)}{ayrac}"
            while not await istek.is_disconnected():
                yeni = [o for o in list(_AKIS) if o["no"] > son]
                for o in yeni:
                    son = o["no"]
                    yield f"data: {json.dumps(o, ensure_ascii=False)}{ayrac}"
                if not yeni:
                    yield ": canli" + ayrac
                await asyncio.sleep(0.5)
        return StreamingResponse(uret(), media_type="text/event-stream")

    @r.post("/api/boss/esitle")
    async def boss_esitle(istek: Request, k: dict = Depends(_esitle_kimlik)):
        try:
            govde = await istek.json()
        except ValueError:
            raise HTTPException(422, {"hata": "Gövde JSON olmalı.", "kod": "alan_eksik"})
        if not isinstance(govde, dict):
            raise HTTPException(422, {"hata": "Gövde JSON nesnesi olmalı.", "kod": "alan_eksik"})
        try:
            sonuc = await run_in_threadpool(boss.esitle, _fabrika, govde)
            _fark_ekle(sonuc, govde)
            oz = (sonuc or {}).get("ozet") or {}
            _akis_ekle({"kaynak": "sunucu", "adim": "esitle", "mod": (sonuc or {}).get("mod"),
                        "gelen": len(govde.get("gorevler") or []), "tam": govde.get("tam"),
                        "yeni": oz.get("yeni", 0), "degisen": oz.get("degisen", 0), "kapanan": oz.get("kapanan", 0),
                        "yeniden_acilan": ((sonuc or {}).get("fark") or {}).get("yeniden_acilan", oz.get("yeniden_acilan", 0)),
                        "rapor_temizligi": (sonuc or {}).get("rapor_temizligi"),
                        "iletisim_gerekli": len((sonuc or {}).get("iletisim") or []),
                        "detay_gerekli": len((sonuc or {}).get("detay_onden") or []),
                        "ozet": {kk: vv for kk, vv in (sonuc or {}).items() if isinstance(vv, (int, str, bool))}})
            return sonuc
        except V2Hata as e:
            raise _hata(e)

    @r.get("/api/boss/izler")
    def boss_izler(k: dict = Depends(_esitle_kimlik), conn: sqlite3.Connection = Depends(baglanti)):
        """Artımlı eşitleme: sunucunun açık bildiği görev no → parmak izi (kişisel veri yok)."""
        return HizliJSON(boss.izler(conn))

    @r.post("/api/boss/iletisim")
    async def boss_iletisim(istek: Request, k: dict = Depends(_esitle_kimlik)):
        """Eklentinin önden okuduğu müşteri iletişim listeleri → telefon (ayar açıkken) + sıradaki işler."""
        if int(istek.headers.get("content-length") or 0) > 8 * 1024 * 1024:
            raise HTTPException(413, {"hata": "Gövde çok büyük.", "kod": "buyuk_dosya"})
        try:
            govde = await istek.json()
        except ValueError:
            raise HTTPException(422, {"hata": "Gövde JSON olmalı.", "kod": "alan_eksik"})
        s = await run_in_threadpool(_cagir, boss.iletisim_yaz, _fabrika, govde)
        _akis_ekle({"kaynak": "sunucu", "adim": "iletisim", "okunan": s["okunan"], "telefon": s["telefon"],
                    "hata": s["hata"], "iletisim_tamam": s["iletisim_tamam"], "iletisim_toplam": s["iletisim_toplam"]})
        return s

    @r.post("/api/boss/detay-onden")
    async def boss_detay_onden(istek: Request, k: dict = Depends(_esitle_kimlik)):
        """Eklentinin önden okuduğu iş ayrıntıları (yeni/değişen açık işler) → kalıcı depo + sıradaki işler."""
        if int(istek.headers.get("content-length") or 0) > 16 * 1024 * 1024:
            raise HTTPException(413, {"hata": "Gövde çok büyük.", "kod": "buyuk_dosya"})
        try:
            govde = await istek.json()
        except ValueError:
            raise HTTPException(422, {"hata": "Gövde JSON olmalı.", "kod": "alan_eksik"})
        s = await run_in_threadpool(_cagir, boss_detay.onden_yaz, _fabrika, govde)
        _akis_ekle({"kaynak": "sunucu", "adim": "detay_onden", "okunan": s["okunan"], "hata": s["hata"],
                    "detay_tamam": s["detay_tamam"], "detay_toplam": s["detay_toplam"],
                    "detay_kalan_dk": s["detay_kalan_dk"]})
        return s

    @r.get("/api/boss/durum")
    def boss_durum(k: dict = Depends(izin("is.liste")), conn: sqlite3.Connection = Depends(baglanti)):
        return boss.durum(conn)

    @r.get("/api/boss/eslesme-kodu")
    def boss_eslesme_kodu(k: dict = Depends(izin("is.ata"))):
        return {"kod": boss.eslesme_kodu()}

    @r.post("/api/boss/eslesme-kodu/yenile")
    def boss_eslesme_yenile(k: dict = Depends(izin("is.ata")), conn: sqlite3.Connection = Depends(baglanti)):
        _yonetim_kaydi(conn, k, "boss_eslesme_yenile", "boss", None)
        return {"kod": boss.eslesme_kodu(yenile=True)}

    # ------------------------------------------------------------------ BOSS iş ayrıntısı (isteğe bağlı)
    @r.post("/api/isler/{is_no}/boss-detay")
    def boss_detay_iste(is_no: str, istek: Request, yenile: bool = False, k: dict = Depends(izin("is.liste")),
                        conn: sqlite3.Connection = Depends(baglanti)):
        _cagir(gorunum.ayrinti, conn, is_no, k)          # işi görebiliyor mu (teknik: yalnız kendi işi)
        s = _cagir(boss_detay.iste, conn, is_no, yenile)
        _akis_ekle({"kaynak": "sunucu", "adim": "detay_istek", "is_no": is_no, "durum": s["durum"]})
        return s

    @r.get("/api/isler/{is_no}/boss-detay")
    def boss_detay_oku(is_no: str, istek: Request, k: dict = Depends(izin("is.liste")),
                       conn: sqlite3.Connection = Depends(baglanti)):
        _cagir(gorunum.ayrinti, conn, is_no, k)
        musteri = kisi.izinli(conn, k, "is.musteri")
        d = boss_detay.oku(is_no, musteri, conn)
        if d["durum"] == "hazir" and not d["taze"] and not d["yenileniyor"]:
            try:                                          # eski kayıt hemen döner; arkada tazelenir
                d["yenileniyor"] = bool(boss_detay.iste(conn, is_no).get("istek_id"))
            except V2Hata:
                pass
        if d["durum"] == "hazir" and musteri:
            _erisim(conn, k, "boss_detay", is_no=is_no, ip=_ip(istek))
        return HizliJSON(d)

    @r.get("/api/boss/istekler")
    async def boss_istekler(bekle: float = Query(default=0, ge=0, le=25), k: dict = Depends(_esitle_kimlik)):
        """Eklentinin sorduğu okunacak iş ayrıntıları (yalnız iş no + BOSS iç kimliği). ``bekle`` > 0: uzun yoklama —
        istek yoksa en çok o kadar saniye bekletilir, kullanıcı bir iş açar açmaz yanıt döner (2,5 sn yoklama yok)."""
        import asyncio
        import time as _t
        son = _t.monotonic() + bekle
        esitle = False
        while True:
            ist = boss_detay.istekler_ver()
            esitle = esitle or boss_detay.esitle_ver()
            if ist or esitle or _t.monotonic() >= son:
                break
            await boss_detay.istek_bekle(min(0.5, max(0.0, son - _t.monotonic())))
        for i in ist:
            _akis_ekle({"kaynak": "sunucu", "adim": "detay_verildi", "is_no": i["is_no"]})
        if esitle:
            _akis_ekle({"kaynak": "sunucu", "adim": "esitle_istendi"})
        return {"istekler": ist, **({"esitle": True} if esitle else {})}

    @r.post("/api/boss/esitle-iste")
    def boss_esitle_iste(k: dict = Depends(izin("is.liste"))):
        """MAP/TASKS ⟳: eklentiden hemen bir BOSS eşitlemesi iste (en çok 20 sn'de bir). Eşitleme bitince değişen
        satırlar değişim yoklamasıyla gelir."""
        return boss_detay.esitle_iste()

    @r.post("/api/boss/detay-oncelik")
    def boss_detay_oncelik(govde: dict, k: dict = Depends(izin("is.liste"))):
        """TASKS'ın o an ekranda gösterdiği (süzülmüş + sıralanmış) ilk işler: önden okuma önce bunları okur."""
        return boss_detay.oncelik_yaz((govde or {}).get("is_nolar"))

    @r.post("/api/boss/detay")
    async def boss_detay_al(istek: Request, k: dict = Depends(_esitle_kimlik),
                            conn: sqlite3.Connection = Depends(baglanti)):
        if int(istek.headers.get("content-length") or 0) > 8 * 1024 * 1024:
            raise HTTPException(413, {"hata": "Ayrıntı çok büyük.", "kod": "buyuk_dosya"})
        try:
            govde = await istek.json()
        except ValueError:
            raise HTTPException(422, {"hata": "Gövde JSON olmalı.", "kod": "alan_eksik"})
        s = _cagir(boss_detay.yanit_al, govde, conn)
        # Ayrıntıdan gelen telefon işe de yazılır: bir kez açılan işin numarası listede hazır durur.
        try:
            no = str((govde or {}).get("is_no") or "")
            kayit = boss_detay._ONBELLEK.get(no) if s["durum"] == "hazir" else None
            tel = next((t for x in ((kayit or {}).get("sekmeler", {}).get("iletisim") or []) for t in (x.get("telefonlar") or [])), None)
            rakam = "".join(ch for ch in str(tel or "") if ch.isdigit())[-10:]
            if len(rakam) == 10 and ayar_acik(conn):
                conn.execute("UPDATE is_emri SET musteri_tel=? WHERE is_no=? AND COALESCE(musteri_tel,'')=''", (rakam, no))
                conn.commit()
        except (sqlite3.Error, AttributeError, TypeError):
            pass
        _akis_ekle({"kaynak": "sunucu", "adim": "detay_" + s["durum"], "is_no": str((govde or {}).get("is_no") or ""),
                    "sekme": s.get("sekme"), "hata": s.get("hata") if s["durum"] == "hazir" else str(govde.get("hata"))[:120]})
        return s

    @r.get("/api/boss/ekipler")
    def boss_ekipler(k: dict = Depends(izin("is.ata")), conn: sqlite3.Connection = Depends(baglanti)):
        return boss_detay.ekipler(conn)

    @r.post("/api/boss/ekip-konum")
    async def boss_ekip_konum(istek: Request, k: dict = Depends(_esitle_kimlik)):
        """Eklentinin okuduğu BOSS GPS yanıtı (ham JSON; biçim hoşgörüyle çözülür) → ekip başına son konum.
        Akışa yalnız sayılar yazılır."""
        from . import ekip_konum
        if int(istek.headers.get("content-length") or 0) > 8 * 1024 * 1024:
            raise HTTPException(413, {"hata": "Gövde çok büyük.", "kod": "buyuk_dosya"})
        try:
            govde = await istek.json()
        except ValueError:
            raise HTTPException(422, {"hata": "Gövde JSON olmalı.", "kod": "alan_eksik"})

        def isle():
            conn = _fabrika()
            try:
                return ekip_konum.yaz(conn, govde if isinstance(govde, dict) else {"veri": govde})
            finally:
                conn.close()
        s = await run_in_threadpool(isle)
        g_ = govde if isinstance(govde, dict) else {}
        _akis_ekle({"kaynak": "sunucu", "adim": "ekip_konum", "tur": str(g_.get("kaynak") or "")[:20] or None,
                    "bulunan": s["bulunan"], "guncellenen": s["guncellenen"], "toplam": s["toplam"],
                    **({"hata": s["hata"]} if s.get("hata") else {})})
        return s

    return [r]


_SON_UYANMA = {"an": 0.0}


def _ek_kisisel(conn, k) -> bool:
    """?ek=1: BOSS ham alanlarının kişisel olanları (müşteri adı/no, adres) yalnız ofiste ve is.musteri izniyle
    (teknik kendi işinde de yalnız kısa ad görür; ham alanda kısaltma yok)."""
    return kisi.izinli(conn, k, "is.musteri") and kisi.ofis_mu(conn, k)


def _aski_ornegi_sakla(yol, govde) -> bool:
    conn = _fabrika()
    try:
        return boss_yazma.aski_ornek_yaz(conn, yol, govde)
    except sqlite3.Error:
        return False
    finally:
        conn.close()


def _nedenleri_sakla(yol, veri) -> int:
    conn = _fabrika()
    try:
        return boss_yazma.nedenleri_yaz(conn, yol, veri)
    except sqlite3.Error:
        return 0
    finally:
        conn.close()


def ayar_acik(conn) -> bool:
    r = conn.execute("SELECT deger FROM ayar WHERE anahtar='musteri_tel'").fetchone()
    return bool(r) and r[0] == "acik"


def _liste_telefonlari(conn, k, y: dict) -> None:
    """İş listesi satırlarına saklanan (BOSS'tan önden okunan) telefon: yalnız ``is.musteri`` izni ve
    ``ayar.musteri_tel='acik'`` iken; değilse alan hiç eklenmez."""
    isler = y.get("isler") or []
    if not isler or not kisi.izinli(conn, k, "is.musteri"):
        return
    r = conn.execute("SELECT deger FROM ayar WHERE anahtar='musteri_tel'").fetchone()
    if not r or r[0] != "acik":
        return
    nolar = [x["is_no"] for x in isler if isinstance(x, dict) and x.get("is_no")]
    tel = {}
    for i in range(0, len(nolar), 900):
        parca = nolar[i:i + 900]
        tel.update(conn.execute(f"SELECT is_no, musteri_tel FROM is_emri WHERE is_no IN ({','.join('?' * len(parca))})",
                                parca).fetchall())
    for x in isler:
        if isinstance(x, dict) and x.get("is_no") in tel:
            x["musteri_tel"] = tel[x["is_no"]]


def _uyananlar(conn) -> None:
    """Uyanma zamanı gelen askılar (dakikada en çok bir kez; liste okunurken)."""
    import time
    if time.time() - _SON_UYANMA["an"] < 60:
        return
    _SON_UYANMA["an"] = time.time()
    try:
        donen = akis.uyananlari_isle(conn, zaman.simdi())
        if donen:
            siralama.oneri_hesapla(conn, donen)
    except sqlite3.Error:
        pass
