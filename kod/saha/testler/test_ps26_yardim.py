"""PS26 testlerinin ortak yardımcıları (Ek-8): sentetik data.xlsx, fotoğraf arşivi, küçük resimler.

Test yoktur; ``test_ps26*.py`` dosyaları buradan içe aktarır. Bütün müşteri numaraları, adlar ve satıcılar
UYDURMADIR (9000xxxx); gerçek çalışma kitabı yalnız ``test_ps26.py::test_gercek_calisma_kitabi_kopyasi``
içinde, geçici klasördeki KOPYASIYLA ve yalnız sayılarla kullanılır.
"""
from __future__ import annotations

import datetime as dt
import io
import os
import zipfile
from pathlib import Path

import yollar

# Gerçek PS26 klasörü (salt okunur; testler yalnız kopyasını kullanır). Başka makinede yoksa test atlanır.
GERCEK_PS26 = Path(os.environ.get("SAHA_PS26_KAYNAK") or yollar.PS26)

BASLIK = ["Ticket", "Baslangic", "Lokasyon", "Site", "Konu", "Durum", "Musteri", "Kanal", "Detay"]


def resim(tur: str = "JPEG", renk=(200, 40, 40), boyut=(640, 480)) -> bytes:
    """Gerçek (Pillow ile açılabilen) küçük bir fotoğraf."""
    from PIL import Image

    tampon = io.BytesIO()
    Image.new("RGB", boyut, renk).save(tampon, tur)
    return tampon.getvalue()


def binalar(conn, adet: int = 4, bolge: int | None = 1) -> list[dict]:
    """location_id'si olan, istenen bölgedeki binalar (lokasyon eşleşmesi için)."""
    kosul = "location_id IS NOT NULL AND location_id <> '' AND pasif=0"
    if bolge is not None:
        kosul += f" AND bolge={int(bolge)}"
    return [dict(r) for r in conn.execute(
        f"SELECT bina_serial, location_id, bolge, site_adi FROM bina WHERE {kosul} ORDER BY bina_serial LIMIT ?",
        (adet,)).fetchall()]


def kitap_yaz(yol: Path, ticket: list[list] | None = None, guzergah: list[list] | None = None,
              altyapi: list[list] | None = None, bos_guzergah: int = 50, pvt: bool = True) -> Path:
    """PS26 data.xlsx biçiminde çalışma kitabı: TICKET, GUZERGAH (çoğu satır boş), ALTYAPI, PVT (+ LOCS)."""
    from openpyxl import Workbook

    kitap = Workbook()
    kitap.remove(kitap.active)
    locs = kitap.create_sheet("LOCS")
    locs.append(["Bina Serial Number", "Location Id"])
    if ticket is not None:
        s = kitap.create_sheet("TICKET")
        s.append(BASLIK + [None, None])
        for r in ticket:
            s.append(r)
    if guzergah is not None:
        s = kitap.create_sheet("GUZERGAH")
        s.append(BASLIK)
        for r in guzergah:
            s.append(r)
        for _ in range(bos_guzergah):
            s.append([None] * len(BASLIK))
        s.append([None, None, None, None, None, 140239, None, None, None])     # yalnız "Durum" dolu: boş sayılır
    if altyapi is not None:
        s = kitap.create_sheet("ALTYAPI")
        s.append(["Musteri", "Kanal", "Baslangic", "Satici", "Bölge", None])
        for r in altyapi:
            s.append(r)
    if pvt:
        s = kitap.create_sheet("PVT")
        s.append(["Satır Etiketleri", "Say Satici"])
    kitap.save(yol)
    return yol


def zip_yaz(yol: Path, dosyalar: dict[str, bytes]) -> Path:
    yol.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(yol, "w") as zf:
        for ad, icerik in dosyalar.items():
            zf.writestr(ad, icerik)
    return yol


def gun(n: int) -> dt.datetime:
    """Bugünden n gün önce (Excel hücresi gibi datetime)."""
    return dt.datetime.combine(dt.date.today() - dt.timedelta(days=n), dt.time())


def sentetik(conn, klasor: Path) -> dict:
    """Standart sentetik PS26 klasörü: data.xlsx + imgs/*.zip. Dönüş: beklenen sayılar ve kimlikler."""
    b = binalar(conn, 4)
    ticket = [
        [10133762171, gun(20), b[0]["location_id"], "SITE A", "SİNYAL", "ÇÖZÜLDÜ", 90000001, "GLOBAL", "kontrol edildi"],
        [None, gun(10), b[1]["location_id"], "SITE B", "EK SP", "AÇIK", "90000002 / 90000003", "DEHA", None],
        [553311, gun(5), b[2]["location_id"] + "-1", "SITE C", "SİNYAL", "HATA", 90000004, "ARIZA", None],
        [None, gun(3), "BILINMEYEN-LOK", "SITE D", "SİNYAL", "AÇIK", 90000005, "TOPTAN", None],
    ]
    guzergah = [
        [None, gun(30), b[3]["location_id"], "SITE E", "GÜZERGAH", "YOK", 90000011, "DEHA", None],
        [None, gun(29), b[0]["location_id"], "SITE A", "GÜZERGAH", "YOK", 90000012, "GLOBAL", None],
        [None, gun(28), "BILINMEYEN-2", "SITE F", "GÜZERGAH", "YOK", 90000013, "DEHA", None],
    ]
    altyapi = [
        [90000021, "GLOBAL", gun(40), "Deneme Satıcı Bir", "DEMİRTAŞ", None],
        [90000022, "DEHA", gun(12), "Deneme Satıcı Bir", "GÖRÜKLE", "port yok"],
        [90000023, "TOPTAN", gun(2), "Deneme Satıcı İki", "BALAT", None],
    ]
    klasor.mkdir(parents=True, exist_ok=True)
    kitap_yaz(klasor / "data.xlsx", ticket, guzergah, altyapi)
    imgs = klasor / "imgs"
    j1, j2, p1 = resim("JPEG", (10, 120, 200)), resim("JPEG", (220, 180, 20)), resim("PNG", (40, 160, 90))
    zip_yaz(imgs / "90000001.zip", {"IMG_1.jpeg": j1, "alt/IMG_2.jpg": j2, "Thumbs.db": b"\x00" * 64})
    zip_yaz(imgs / "90000003.zip", {"ekran.png": p1})                      # çok müşterili alan: 2. numara
    zip_yaz(imgs / "553311.zip", {"a.jpg": j2})                            # ticket no ile
    zip_yaz(imgs / "90000012.zip", {"g.jpg": j1})                          # GÜZERGAH satırı (TICKET'ta yok)
    zip_yaz(imgs / "12345678.zip", {"x.jpg": j1})                          # hiçbir satırla eşleşmez
    return {"binalar": b, "ticket": len(ticket), "guzergah": len(guzergah), "altyapi": len(altyapi),
            "foto": 5, "zip": 5, "eslesen_zip": 4, "alinmayan": 1}
