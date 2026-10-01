"""PS26 içe aktarımı (Ek-8): önce önizleme, etkisiz (iki kez = aynı), ticket kuralları ``ticket_aktar`` ile aynı,
GÜZERGAH yalnız dolu satırlar, ALTYAPI → ``altyapi_bekleyen``, fotoğraflar arşivden ticket'a.

Sentetik çalışma kitabı her makinede koşar; gerçek ``PS26/data.xlsx`` varsa KOPYASIYLA da denenir (yalnız sayılar,
kaynağın özeti değişmez). Kişisel veri hiçbir iddia mesajına yazılmaz.
"""
from __future__ import annotations

import hashlib
import shutil
import sqlite3
from pathlib import Path

import pytest

from saha import ek_dosya, ps26_aktar, ticket_aktar

from .test_ps26_yardim import GERCEK_PS26, kitap_yaz, sentetik, zip_yaz


def _sayim(conn: sqlite3.Connection) -> dict:
    return {t: conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
            for t in ("ticket", "ticket_gecmis", "altyapi_bekleyen", "ticket_ek", "yonetim_kaydi")}


def _diskteki_fotolar() -> int:
    kok = ek_dosya.ek_kok() / "ticket"
    return sum(1 for p in kok.rglob("*") if p.is_file() and p.parent.name != "kucuk") if kok.exists() else 0


@pytest.fixture
def ps26(conn, tmp_path):
    klasor = tmp_path / "PS26"
    beklenen = sentetik(conn, klasor)
    return klasor, beklenen


# ============================================================ önizleme, uygulama, etkisizlik
def test_onizleme_hicbir_sey_yazmaz_uygulama_iki_kez_ayni(conn, ps26):
    klasor, b = ps26
    okuma = ps26_aktar.oku(klasor / "data.xlsx", imgs=klasor / "imgs")
    assert (len(okuma.ticket), len(okuma.guzergah), len(okuma.altyapi)) == (b["ticket"], b["guzergah"], b["altyapi"])
    once = _sayim(conn)

    s = ps26_aktar.aktar(conn, okuma, kuru=True)
    assert _sayim(conn) == once and _diskteki_fotolar() == 0          # önizleme: ne satır ne dosya
    assert s["ticket"]["eklenecek"] == b["ticket"]
    assert s["guzergah"]["eklenecek"] == b["guzergah"]
    assert s["altyapi"]["eklenecek"] == b["altyapi"]
    assert s["foto"]["eklenecek"] == b["foto"] and s["foto"]["eslesen_zip"] == b["eslesen_zip"]
    assert s["foto"]["eslesmeyen_zip"] == 1 and s["foto"]["alinmayan"] == b["alinmayan"]

    s1 = ps26_aktar.aktar(conn, okuma, kuru=False)
    assert s1["ticket"]["eklenen"] == b["ticket"] and s1["guzergah"]["eklenen"] == b["guzergah"]
    assert s1["altyapi"]["eklenen"] == b["altyapi"] and s1["foto"]["eklenen"] == b["foto"]
    sonra = _sayim(conn)
    assert sonra["ticket"] == once["ticket"] + b["ticket"] + b["guzergah"]
    assert sonra["altyapi_bekleyen"] == b["altyapi"] and sonra["ticket_ek"] == b["foto"]
    assert _diskteki_fotolar() == b["foto"]

    s2 = ps26_aktar.aktar(conn, okuma, kuru=False)                    # ikinci kez: hiçbir şey eklenmez
    for bolum in ("ticket", "guzergah", "altyapi", "foto"):
        assert s2[bolum].get("eklenen", 0) == 0, bolum
        assert s2[bolum].get("guncellenen", 0) == 0, bolum
    assert s2["ticket"]["zaten_var"] == b["ticket"] and s2["foto"]["zaten_var"] == b["foto"]
    assert {a: v for a, v in _sayim(conn).items() if a != "yonetim_kaydi"} == \
           {a: v for a, v in sonra.items() if a != "yonetim_kaydi"}
    assert _diskteki_fotolar() == b["foto"]
    # İkinci önizleme de "yapılacak bir şey yok" der.
    assert ps26_aktar.aktar(conn, okuma, kuru=True)["eklenecek_toplam"] == 0


def test_ticket_kurallari_ticket_aktar_ile_ayni(conn, ps26):
    """TICKET sayfası ``ticket_aktar`` ile birebir aynı anahtarla yazılır: ardından eski araç hiçbir şey eklemez."""
    klasor, b = ps26
    ps26_aktar.aktar(conn, ps26_aktar.oku(klasor / "data.xlsx"), kuru=False)
    s = ticket_aktar.aktar(conn, ticket_aktar.satirlari_oku(klasor / "data.xlsx"), kuru=True)
    assert s.get("eklenecek", 0) == 0 and s["zaten_var"] == b["ticket"]
    # Lokasyon → bina: sondaki "-1" eki atılarak eşleşir; tanınmayan lokasyon binasız ama yine aktarılır.
    bagli = conn.execute("SELECT COUNT(*) FROM ticket WHERE tur IS NULL AND bina_serial IS NOT NULL").fetchone()[0]
    assert bagli == b["ticket"] - 1
    assert conn.execute("SELECT COUNT(*) FROM ticket WHERE tur IS NULL AND kaynak='excel'").fetchone()[0] == b["ticket"]


def test_guzergah_yalniz_dolu_satirlar_tur_ve_durum(conn, ps26):
    klasor, b = ps26
    s = ps26_aktar.aktar(conn, ps26_aktar.oku(klasor / "data.xlsx"), kuru=False)
    assert s["guzergah"]["okunan"] == b["guzergah"]                  # 50 boş + yalnız "Durum"u dolu satır alınmadı
    satirlar = conn.execute("SELECT konu, durum, tur, kaynak, ticket_no FROM ticket WHERE tur='guzergah'").fetchall()
    assert len(satirlar) == b["guzergah"]
    assert {(r["konu"], r["durum"], r["kaynak"]) for r in satirlar} == {("GÜZERGAH", "AÇIK", "excel")}  # 'YOK' → açık
    assert all(r["ticket_no"] is None for r in satirlar)


def test_uygulamada_degisen_kayda_excel_dokunmaz(conn, ps26, istemci, kisi):
    klasor, b = ps26
    ps26_aktar.aktar(conn, ps26_aktar.oku(klasor / "data.xlsx"), kuru=False)
    _, op = kisi("operasyon")
    # Uygulamada: bir ticket çözüldü, bir altyapı satırı kuruldu.
    tid = conn.execute("SELECT id FROM ticket WHERE tur IS NULL AND durum='AÇIK' ORDER BY id LIMIT 1").fetchone()[0]
    assert istemci.patch(f"/api/ticket/{tid}", headers=op, json={"durum": "ÇÖZÜLDÜ"}).status_code == 200
    aid = conn.execute("SELECT id FROM altyapi_bekleyen ORDER BY id LIMIT 1").fetchone()[0]
    assert istemci.patch(f"/api/tablolar/altyapi/{aid}", headers=op, json={"durum": "kuruldu"}).status_code == 200

    # Excel'de: aynı satırlar değişmiş gibi (satıcı adı, ticket durumu), bir altyapı satırı silinmiş.
    from openpyxl import load_workbook

    kitap = load_workbook(klasor / "data.xlsx")
    for satir in kitap["TICKET"].iter_rows(min_row=2):
        satir[5].value = "KAPATILDI"
    alt = kitap["ALTYAPI"]
    for satir in alt.iter_rows(min_row=2):
        satir[3].value = "Deneme Satıcı Üç"
    alt.delete_rows(alt.max_row)
    kitap.save(klasor / "data.xlsx")

    s = ps26_aktar.aktar(conn, ps26_aktar.oku(klasor / "data.xlsx"), kuru=False)
    assert s["ticket"]["uygulamada_degismis"] == 1                    # uygulamada değişene Excel dokunmaz
    assert conn.execute("SELECT durum FROM ticket WHERE id=?", (tid,)).fetchone()[0] == "ÇÖZÜLDÜ"
    assert conn.execute("SELECT durum FROM altyapi_bekleyen WHERE id=?", (aid,)).fetchone()[0] == "kuruldu"
    assert s["altyapi"]["guncellenen"] == b["altyapi"] - 1            # satıcı adı tazelendi, durum ezilmedi
    assert s["altyapi"]["excelde_yok"] == 1                            # silinmez, sayılır
    assert conn.execute("SELECT COUNT(*) FROM altyapi_bekleyen").fetchone()[0] == b["altyapi"]


def test_pvt_turetilir(conn, ps26):
    klasor, _ = ps26
    ps26_aktar.aktar(conn, ps26_aktar.oku(klasor / "data.xlsx"), kuru=False)
    ozet = {r["satici"]: r for r in ps26_aktar.satici_ozeti(conn)}
    assert ozet["Deneme Satıcı Bir"]["bekliyor"] == 2 and ozet["Deneme Satıcı İki"]["toplam"] == 1


# ============================================================ fotoğraf eşleşmesi
def test_fotolar_dogru_ticketa_ve_yalniz_resimler(conn, ps26):
    klasor, _ = ps26
    ps26_aktar.aktar(conn, ps26_aktar.oku(klasor / "data.xlsx", imgs=klasor / "imgs"), kuru=False)

    def ek_sayisi(musteri_parcasi: str) -> int:
        return conn.execute("SELECT COUNT(*) FROM ticket_ek e JOIN ticket t ON t.id=e.ticket_id "
                            "WHERE t.musteri LIKE ?", (f"%{musteri_parcasi}%",)).fetchone()[0]

    assert ek_sayisi("90000001") == 2          # arşivde 2 resim + Thumbs.db (alınmaz), alt klasördeki de alındı
    assert ek_sayisi("90000003") == 1          # "90000002 / 90000003": ikinci numara da eşleşir
    assert ek_sayisi("90000004") == 1          # arşivin adı ticket no (553311)
    assert ek_sayisi("90000012") == 1          # TICKET'ta yoksa GÜZERGAH satırına
    # Diskteki ad içerikten türetilir, arşivdeki yol ("alt/…") kullanılmaz.
    for r in conn.execute("SELECT ticket_id, sha256, dosya_adi FROM ticket_ek"):
        yol = ek_dosya.dosya_yolu(r[0], r[1], r[2])
        assert yol.exists() and yol.parent == ek_dosya.ticket_dizini(r[0])
        assert hashlib.sha256(yol.read_bytes()).hexdigest() == r[1]
        assert "/" not in r[2] and "\\" not in r[2]


def test_fotograf_klasoru_yoksa_atlanir(conn, ps26):
    klasor, b = ps26
    s = ps26_aktar.aktar(conn, ps26_aktar.oku(klasor / "data.xlsx", imgs=klasor / "yok"), kuru=False)
    assert s["foto"] == {"klasor": False} and s["ticket"]["eklenen"] == b["ticket"]


def test_bozuk_arsiv_ve_buyuk_girdi(conn, tmp_path, ps26):
    klasor, _ = ps26
    (klasor / "imgs" / "90000005.zip").write_bytes(b"PK\x03\x04 bu bir zip degil")
    s = ps26_aktar.aktar(conn, ps26_aktar.oku(klasor / "data.xlsx", imgs=klasor / "imgs"), kuru=False)
    assert s["foto"]["bozuk_zip"] == 1
    # 10 MB'ı aşan girdi arşivden okunmaz.
    buyuk = zip_yaz(tmp_path / "b" / "90000001.zip", {"dev.jpg": b"\xff\xd8\xff" + b"0" * (ek_dosya.EN_BUYUK + 10)})
    alinan, alinmayan = ek_dosya.arsiv_dosyalari(buyuk)
    assert alinan == [] and alinmayan == 1


# ============================================================ hatalar, komut satırı
def test_ps26_olmayan_dosya_taninmaz(tmp_path):
    from openpyxl import Workbook

    k = Workbook()
    k.active.title = "Sayfa1"
    k.save(tmp_path / "rapor.xlsx")
    with pytest.raises(ps26_aktar.Ps26Hatasi) as h:
        ps26_aktar.oku(tmp_path / "rapor.xlsx")
    assert h.value.kod == "rapor_tanimadi"
    (tmp_path / "bozuk.xlsx").write_bytes(b"bu bir excel degil")
    with pytest.raises(ps26_aktar.Ps26Hatasi):
        ps26_aktar.oku(tmp_path / "bozuk.xlsx")


def test_eksik_sayfalar_sorun_degil(conn, tmp_path):
    kitap_yaz(tmp_path / "yalniz_altyapi.xlsx", altyapi=[[90000031, "DEHA", None, "Deneme", "BALAT", None]])
    s = ps26_aktar.aktar(conn, ps26_aktar.oku(tmp_path / "yalniz_altyapi.xlsx"), kuru=False)
    assert s["ticket"]["okunan"] == 0 and s["guzergah"]["okunan"] == 0 and s["altyapi"]["eklenen"] == 1


def test_komut_satiri_once_onizleme_sonra_yedekli_uygulama(db_yolu, ps26, capsys):
    klasor, b = ps26
    c = sqlite3.connect(db_yolu)
    once = c.execute("SELECT COUNT(*) FROM ticket").fetchone()[0]
    c.close()
    assert ps26_aktar.main([str(klasor / "data.xlsx"), "--db", str(db_yolu), "--kuru"]) == 0
    cikti = capsys.readouterr().out
    assert "ÖNİZLEME" in cikti and "90000001" not in cikti          # yalnız sayılar
    c = sqlite3.connect(db_yolu)
    assert c.execute("SELECT COUNT(*) FROM ticket").fetchone()[0] == once
    c.close()
    assert ps26_aktar.main([str(klasor / "data.xlsx"), "--db", str(db_yolu)]) == 0
    cikti = capsys.readouterr().out
    assert "Yedek alındı" in cikti and "90000" not in cikti
    assert list((db_yolu.parent / "yedek" / "aktarim").glob("saha-aktarim-*.db"))
    c = sqlite3.connect(db_yolu)
    assert c.execute("SELECT COUNT(*) FROM ticket").fetchone()[0] == once + b["ticket"] + b["guzergah"]
    assert c.execute("SELECT COUNT(*) FROM ticket_ek").fetchone()[0] == b["foto"]
    c.close()
    assert (db_yolu.parent / "ek" / "ticket").is_dir()                   # fotoğraflar --db'nin yanında


# ============================================================ API
def _gonder(istemci, baslik, klasor: Path, kuru: bool, zipler: bool = True):
    dosyalar = [("dosya", ("data.xlsx", (klasor / "data.xlsx").read_bytes(),
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))]
    if zipler:
        dosyalar += [("zip", (z.name, z.read_bytes(), "application/zip")) for z in sorted((klasor / "imgs").glob("*.zip"))]
    return istemci.post(f"/api/ps26/aktar?kuru={'true' if kuru else 'false'}", headers=baslik, files=dosyalar)


def test_aktarim_ucu_onizleme_uygulama_ve_yetki(istemci, kisi, conn, ps26, db_yolu):
    klasor, b = ps26
    _, op = kisi("operasyon")
    y = _gonder(istemci, op, klasor, kuru=True)
    assert y.status_code == 200, y.text
    v = y.json()
    assert v["kuru"] is True and v["ticket"]["eklenecek"] == b["ticket"] and v["foto"]["eklenecek"] == b["foto"]
    assert conn.execute("SELECT COUNT(*) FROM altyapi_bekleyen").fetchone()[0] == 0
    y = _gonder(istemci, op, klasor, kuru=False)
    assert y.status_code == 200, y.text
    v = y.json()
    assert v["altyapi"]["eklenen"] == b["altyapi"] and v["foto"]["eklenen"] == b["foto"]
    assert v["yedek"] and "/" not in v["yedek"] and "\\" not in v["yedek"]   # yalnız dosya adı
    assert _gonder(istemci, op, klasor, kuru=False).json()["eklenecek_toplam"] == 0
    # Yanıtta kişisel veri yok (müşteri numaraları sentetik ama biçim aynı).
    assert "9000000" not in y.text
    for rol in ("satisci", "teknik"):
        _, h = kisi(rol, bolge=1 if rol == "satisci" else None)
        assert _gonder(istemci, h, klasor, kuru=True).status_code == 403
    assert istemci.post("/api/ps26/aktar", files={"dosya": ("x.xlsx", b"x")}).status_code == 401


def test_aktarim_ucu_tanimadigi_dosya(istemci, kisi, tmp_path):
    _, op = kisi("operasyon")
    y = istemci.post("/api/ps26/aktar?kuru=true", headers=op, files={"dosya": ("x.xlsx", b"excel degil")})
    assert y.status_code == 422 and y.json()["kod"] == "rapor_tanimadi"
    y = istemci.post("/api/ps26/aktar?kuru=true", headers=op)
    assert y.status_code == 400 and y.json()["kod"] == "dosya_yok"


def test_sunucudaki_ps26_klasorunden(istemci, kisi, ps26, monkeypatch, tmp_path):
    klasor, b = ps26
    _, op = kisi("operasyon")
    monkeypatch.setenv("SAHA_PS26", str(tmp_path / "yok"))
    d = istemci.get("/api/ps26", headers=op).json()
    assert d["klasor"]["bulundu"] is False
    assert istemci.post("/api/ps26/aktar?klasor=true", headers=op).json()["kod"] == "klasor_yok"
    monkeypatch.setenv("SAHA_PS26", str(klasor))
    d = istemci.get("/api/ps26", headers=op).json()
    assert d["klasor"]["bulundu"] is True and d["klasor"]["zip"] == b["zip"] and d["son_aktarim"] is None
    y = istemci.post("/api/ps26/aktar?klasor=true&kuru=false", headers=op)
    assert y.status_code == 200 and y.json()["foto"]["eklenen"] == b["foto"]
    assert istemci.get("/api/ps26", headers=op).json()["son_aktarim"]


# ============================================================ gerçek çalışma kitabının KOPYASI
@pytest.mark.skipif(not (GERCEK_PS26 / "data.xlsx").is_file(), reason="PS26/data.xlsx bu makinede yok")
def test_gercek_calisma_kitabi_kopyasi(conn, tmp_path):
    """Gerçek data.xlsx + imgs: önizleme = uygulama; ikinci uygulama hiçbir şey eklemez; kaynak değişmez."""
    kaynak = GERCEK_PS26 / "data.xlsx"
    ozet_once = hashlib.sha256(kaynak.read_bytes()).hexdigest()
    kopya = tmp_path / "PS26"
    kopya.mkdir()
    shutil.copy2(kaynak, kopya / "data.xlsx")
    if (GERCEK_PS26 / "imgs").is_dir():
        shutil.copytree(GERCEK_PS26 / "imgs", kopya / "imgs")
    okuma = ps26_aktar.oku(kopya / "data.xlsx", imgs=kopya / "imgs")
    # Sayfa boyları PS26'nın bilinen hâline yakın (kullanıcı eklemeye devam ediyor: alt sınır).
    assert len(okuma.ticket) >= 190 and 40 <= len(okuma.guzergah) <= 500 and len(okuma.altyapi) >= 60

    on = ps26_aktar.aktar(conn, okuma, kuru=True)
    s1 = ps26_aktar.aktar(conn, okuma, kuru=False)
    for bolum in ("ticket", "guzergah", "altyapi", "foto"):
        assert on[bolum].get("eklenecek", 0) == s1[bolum].get("eklenen", 0), bolum
    assert s1["ticket"]["eklenen"] + s1["ticket"]["zaten_var"] == len(okuma.ticket)
    assert s1["guzergah"]["eklenen"] == len(okuma.guzergah)
    assert s1["altyapi"]["eklenen"] == len(okuma.altyapi)
    assert s1["ticket"]["binaya_baglanan"] >= 0.9 * len(okuma.ticket)
    if okuma.zipler:
        assert s1["foto"]["eslesen_zip"] >= 0.7 * len(okuma.zipler) and s1["foto"]["eklenen"] > 0
    sayim = _sayim(conn)

    s2 = ps26_aktar.aktar(conn, okuma, kuru=False)
    assert sum(s2[b].get("eklenen", 0) + s2[b].get("guncellenen", 0) for b in ("ticket", "guzergah", "altyapi", "foto")) == 0
    assert {a: v for a, v in _sayim(conn).items() if a != "yonetim_kaydi"} == \
           {a: v for a, v in sayim.items() if a != "yonetim_kaydi"}
    assert conn.execute("SELECT COUNT(*) FROM ticket WHERE tur='guzergah'").fetchone()[0] == len(okuma.guzergah)
    assert hashlib.sha256(kaynak.read_bytes()).hexdigest() == ozet_once        # kaynak salt okunur kaldı
