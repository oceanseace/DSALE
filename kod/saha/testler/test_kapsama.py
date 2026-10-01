"""Kapsama matematiği, günlük özet, harita dizileri ve Excel raporu."""
from __future__ import annotations

from .conftest import SATISCI_TEL


def test_baslangicta_hicbir_seye_dokunulmamis(istemci, yonetici):
    veri = istemci.get("/api/ozet/kapsama?kirilim=bolge", headers=yonetici).json()
    assert veri["toplam"] == 19706
    assert veri["dokunulan"] == 0
    assert veri["kalan"] == 19706
    assert veri["oran"] == 0.0
    assert veri["firsat"] == 189272
    assert veri["kalan_firsat"] == 189272
    assert len(veri["satirlar"]) == 8
    assert sum(s["toplam"] for s in veri["satirlar"]) == 19706


def test_ziyaret_kapsamayi_ilerletir(istemci, satisci1, yonetici):
    gorev = istemci.post("/api/gorev/olustur", headers=satisci1, json={"adet": 5}).json()
    hedefler = gorev["binalar"][:3]
    for i, b in enumerate(hedefler):
        istemci.post("/api/ziyaret", headers=satisci1, json={
            "offline_id": f"kapsama-{i}", "bina_serial": b["bina_serial"],
            "sonuc": "satis" if i == 0 else "ilgilenmedi", "satis_adedi": 1 if i == 0 else 0,
        })

    veri = istemci.get("/api/ozet/kapsama?kirilim=bolge", headers=yonetici).json()
    assert veri["dokunulan"] == 3
    assert veri["kalan"] == 19706 - 3
    assert veri["oran"] == round(3 / 19706, 4)
    assert veri["satis"] == 1
    # "Kalan boş kapı" artık SATILMAMIŞ kapı sayar. Eskiden binaya bir kez
    # dokunulunca o binanın BÜTÜN boş kapıları "halledilmiş" sayılıyor ve
    # ilerleme ciddi şekilde abartılıyordu (753 satışla ~38.000 kapı bitmiş
    # görünüyordu). Burada tek bir abonelik satıldı: kalan tam 1 azalır.
    assert veri["kalan_firsat"] == 189272 - 1
    # "dokunulmayan_firsat" eski ölçüyü ayrı bir alan olarak sürdürür.
    dusen = sum(b["firsat"] for b in hedefler)
    assert veri["dokunulmayan_firsat"] == 189272 - dusen
    # Üç ziyaretin üçünde de kapı açıldı: temas = dokunulan.
    assert veri["temas"] == 3

    bolge1 = next(s for s in veri["satirlar"] if s["ad"] == 1)
    assert bolge1["dokunulan"] == 3
    assert bolge1["toplam"] == bolge1["dokunulan"] + bolge1["kalan"]


def test_kapsama_kirilimlari(istemci, yonetici):
    for kirilim in ("bolge", "mahalle", "ilce"):
        veri = istemci.get(f"/api/ozet/kapsama?kirilim={kirilim}", headers=yonetici).json()
        assert veri["kirilim"] == kirilim
        assert sum(s["toplam"] for s in veri["satirlar"]) == veri["toplam"] == 19706
    assert istemci.get("/api/ozet/kapsama?kirilim=sokak", headers=yonetici).status_code == 422


def test_satisci_kendi_kapsamasini_gorur(istemci, satisci1):
    veri = istemci.get("/api/ozet/kapsama?kirilim=mahalle", headers=satisci1).json()
    assert 0 < veri["toplam"] < 19706
    assert len(veri["satirlar"]) > 1


def test_ozet_gun(istemci, satisci1, yonetici):
    gorev = istemci.post("/api/gorev/olustur", headers=satisci1, json={"adet": 6}).json()
    sonuclar = ["satis", "satis", "ilgilenmedi", "randevu", "evde_yok"]
    for i, (b, sonuc) in enumerate(zip(gorev["binalar"], sonuclar)):
        istemci.post("/api/ziyaret", headers=satisci1, json={
            "offline_id": f"gunozet-{i}", "bina_serial": b["bina_serial"], "sonuc": sonuc,
            "satis_adedi": 2 if sonuc == "satis" else 0,
        })

    veri = istemci.get("/api/ozet/gun", headers=yonetici).json()
    assert len(veri["satiscilar"]) == 8
    satir = next(s for s in veri["satiscilar"] if s["bolge"] == 1)
    assert satir["ziyaret"] == 5
    assert satir["satis"] == 2 and satir["satis_adedi"] == 4
    assert satir["ret"] == 1 and satir["randevu"] == 1 and satir["evde_yok"] == 1
    # Liste site bütünlüğü için birkaç bina taşabilir (bkz. rota.LISTE_TASMA):
    # kalan = listedeki bina - işlenen bina.
    assert satir["kalan"] == len(gorev["binalar"]) - 5
    assert satir["donusum"] == round(2 / 5, 4)

    assert veri["toplam"]["ziyaret"] == 5
    assert veri["toplam"]["satis_adedi"] == 4
    bos = istemci.get("/api/ozet/gun?tarih=2020-01-01", headers=yonetici).json()
    assert bos["toplam"]["ziyaret"] == 0


def test_harita_dizileri(istemci, satisci1, yonetici):
    kendi = istemci.get("/api/harita", headers=satisci1).json()
    assert kendi["adet"] == len(kendi["serial"]) == len(kendi["lat"]) == len(kendi["durum"])
    assert set(kendi["bolge"]) == {1}
    assert kendi["ofis"]["lat"] > 40

    hepsi = istemci.get("/api/harita", headers=yonetici).json()
    assert hepsi["adet"] == 19706
    assert set(hepsi["durum"]) == {"bekliyor"}


def test_bina_arama_ve_sayfalama(istemci, yonetici):
    ilk = istemci.get("/api/bina?limit=10", headers=yonetici).json()
    assert ilk["toplam"] == 19706 and len(ilk["binalar"]) == 10
    ikinci = istemci.get("/api/bina?limit=10&offset=10", headers=yonetici).json()
    assert {b["bina_serial"] for b in ilk["binalar"]}.isdisjoint(
        {b["bina_serial"] for b in ikinci["binalar"]})

    mahalle = ilk["binalar"][0]["mahalle"]
    suzulmus = istemci.get(f"/api/bina?mahalle={mahalle}&limit=5", headers=yonetici).json()
    assert all(b["mahalle"] == mahalle for b in suzulmus["binalar"])

    bekleyen = istemci.get("/api/bina?durum=bekliyor&limit=1", headers=yonetici).json()
    assert bekleyen["toplam"] == 19706
    assert istemci.get("/api/bina?durum=uydurma", headers=yonetici).json()["kod"] == "durum_gecersiz"


def test_gunluk_excel_raporu(istemci, satisci1, yonetici):
    gorev = istemci.post("/api/gorev/olustur", headers=satisci1, json={"adet": 3}).json()
    istemci.post("/api/ziyaret", headers=satisci1, json={
        "offline_id": "rapor-1", "bina_serial": gorev["binalar"][0]["bina_serial"],
        "sonuc": "satis", "satis_adedi": 1, "not": "Excel testi",
    })
    yanit = istemci.get("/api/dosya/rapor.xlsx", headers=yonetici)
    assert yanit.status_code == 200
    assert yanit.content[:2] == b"PK"           # xlsx = zip
    assert len(yanit.content) > 4000
    assert "spreadsheetml" in yanit.headers["content-type"]


def test_rapor_sayfalari(conn, tmp_path):
    from openpyxl import load_workbook

    from saha import rapor

    yol = rapor.gunluk_rapor_yaz(conn, "2026-09-21", tmp_path / "r.xlsx")
    wb = load_workbook(yol)
    assert wb.sheetnames == ["GUN", "ZIYARETLER", "KAPSAMA"]
    assert wb["GUN"]["A1"].value == "SAHA GÜNLÜK RAPORU"
    kapsama = wb["KAPSAMA"]
    bolgeler = [kapsama.cell(row=s, column=1).value for s in range(5, 13)]
    assert bolgeler == list(range(1, 9))
