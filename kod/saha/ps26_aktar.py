"""PS26 ``data.xlsx`` → Saha Sistemi (Ek-8): PS26'yı kaldırabilmek için tek aktarım.

    .venv/Scripts/python.exe -m saha.ps26_aktar ..\\PS26\\data.xlsx --kuru     # önizleme: ne olacak, hiçbir şey yazılmaz
    .venv/Scripts/python.exe -m saha.ps26_aktar ..\\PS26\\data.xlsx            # uygula (önce doğrulanmış yedek)
    .venv/Scripts/python.exe -m saha.ps26_aktar data.xlsx --db kopya.db --imgs ..\\PS26\\imgs

Ekrandan da çalışır: Tablolar › "PS26'dan aktar" (``saha/tablolar_uclari.py``, aynı fonksiyonlar).

Sayfalar:
  TICKET      → ticket defteri: ``ticket_aktar`` ile BİREBİR aynı kurallar (aynı fonksiyon; anahtar, eşitleme,
                "uygulamada değiştirilmiş ticket'a Excel dokunmaz", lokasyon → bina eşleşmesi).
  GUZERGAH    → ticket (``tur='guzergah'``, konu GÜZERGAH); yalnız dolu satırlar (16.912 satırın ≈66'sı).
                Bu sayfada "Durum" hep 'YOK' yazar (güzergah yok = sorun sürüyor) → AÇIK.
  ALTYAPI     → ``altyapi_bekleyen``: altyapı yüzünden kurulamayan satışlar. Durum (bekliyor · kuruldu · iptal)
                bundan sonra uygulamada tutulur; Excel yalnız satıcı/bölge/not'u tazeler, durumu asla ezmez.
  PVT         → saklanmaz; "satıcıya göre" sayım ekranda ``altyapi_bekleyen``'den türetilir.
  LOCS, ORIGN → alınmaz (bina tablosu zaten var).
  imgs/*.zip  → ticket fotoğrafları (``ek_dosya``). PS26'da arşivin adı çoğunlukla ticket satırındaki MÜŞTERİ
                numarasıdır (bazen ticket numarası): ikisine de bakılır; birden çok satır eşleşirse en yeni açılış.

Etkisizdir: iki kez çalıştırmak ikinci kez hiçbir şey eklemez. Çıktıda kişisel veri yoktur; yalnız sayılar.
"""
from __future__ import annotations

import argparse
import hashlib
import os
import re
import sqlite3
import sys
import time
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import yollar

from . import ayarlar, db, ek_dosya, ticket, ticket_aktar

SAYFA_TICKET, SAYFA_GUZERGAH, SAYFA_ALTYAPI = "TICKET", "GUZERGAH", "ALTYAPI"
ALTYAPI_DURUMLARI = ("bekliyor", "kuruldu", "iptal")
TUR_GUZERGAH = "guzergah"


class Ps26Hatasi(Exception):
    """Dosya PS26 çalışma kitabı değil ya da okunamadı. ``mesaj`` Türkçe."""

    def __init__(self, mesaj: str, kod: str = "rapor_tanimadi"):
        super().__init__(mesaj)
        self.mesaj, self.kod = mesaj, kod


# ============================================================================= klasör
def klasor_adaylari(conn: sqlite3.Connection | None = None) -> list[Path]:
    """PS26 klasörünün aranacağı yerler: ``SAHA_PS26`` · ayar ``ps26_klasoru`` · deponun yanı · GitHub Desktop."""
    adaylar: list[Path] = []
    if os.environ.get("SAHA_PS26"):
        return [Path(os.environ["SAHA_PS26"])]          # testler ve provalar: yalnız bu
    if conn is not None:
        try:
            ayar = db.ayar_oku(conn, "ps26_klasoru")
        except sqlite3.Error:
            ayar = None
        if ayar:
            adaylar.append(Path(ayar))
    adaylar += [yollar.PS26, Path.home() / "Documents" / "GitHub" / "PS26"]
    return list(dict.fromkeys(adaylar))


def klasor_bul(conn: sqlite3.Connection | None = None) -> Path | None:
    """İçinde ``data.xlsx`` olan ilk aday klasör; yoksa None."""
    for aday in klasor_adaylari(conn):
        if (aday / "data.xlsx").is_file():
            return aday
    return None


def arsivler(imgs: Path | None) -> list[Path] | None:
    """``imgs/*.zip`` (ada göre sıralı). Klasör yoksa None ("fotoğraf klasörü yok"), boşsa []."""
    if imgs is None or not Path(imgs).is_dir():
        return None
    return sorted(Path(imgs).glob("*.zip"), key=lambda p: p.name)


# ============================================================================= okuma
@dataclass
class Okuma:
    dosya_adi: str
    sayfalar: list[str]
    ticket: list[dict] = field(default_factory=list)
    guzergah: list[dict] = field(default_factory=list)
    altyapi: list[dict] = field(default_factory=list)
    guzergah_ham: int = 0                    # sayfadaki bütün satırlar (çoğu boş)
    zipler: list[Path] | None = None         # None = fotoğraf klasörü verilmedi


def _sayfa_adlari(dosya: Path) -> list[str]:
    from openpyxl import load_workbook

    try:
        kitap = load_workbook(dosya, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — zip değil, bozuk, şifreli…
        raise Ps26Hatasi("Bu dosya açılamadı. PS26 klasöründeki data.xlsx'i seçin.") from exc
    try:
        return list(kitap.sheetnames)
    finally:
        kitap.close()


def _altyapi_oku(dosya: Path) -> list[dict]:
    """ALTYAPI sayfası: Musteri · Kanal · Baslangic · Satici · Bölge (+ başlıksız not sütunu)."""
    import pandas as pd

    d = pd.read_excel(dosya, sheet_name=SAYFA_ALTYAPI, dtype=str)
    adlar = {ticket._sade(str(c)): c for c in d.columns}
    ekler = [c for c in d.columns if str(c).lower().startswith("unnamed")]

    def al(r: dict, *secenek: str) -> str:
        for s in secenek:
            if s in adlar:
                deger = ticket_aktar._metin(r.get(adlar[s]))
                if deger:
                    return deger
        return ""

    satirlar = []
    for i, r in enumerate(d.to_dict("records")):
        kayit = {
            "musteri": al(r, "MUSTERI", "MUSTERINO"),
            "kanal": ticket._tr_ust(al(r, "KANAL")),
            "baslangic": ticket._tarih(al(r, "BASLANGIC", "TARIH")) or "",
            "satici": al(r, "SATICI", "SATISTEMSILCISI", "SATISTEMSILCISIADSOYAD"),
            "bolge": al(r, "BOLGE"),
            "not": " · ".join(ticket_aktar._metin(r.get(c)) for c in ekler if ticket_aktar._metin(r.get(c))),
            "_satir": i + 2,
        }
        if not any(kayit[a] for a in ("musteri", "satici", "baslangic")):
            continue                               # tamamen boş satır
        satirlar.append(kayit)
    return satirlar


def oku(dosya: Path, imgs: Path | None = None, zipler: list[Path] | None = None) -> Okuma:
    """Çalışma kitabını okur (yazmaz). ``imgs`` klasörü ya da doğrudan ``zipler`` verilebilir."""
    dosya = Path(dosya)
    sayfalar = _sayfa_adlari(dosya)
    bilinen = {SAYFA_TICKET, SAYFA_GUZERGAH, SAYFA_ALTYAPI} & set(sayfalar)
    if not bilinen:
        raise Ps26Hatasi("Bu dosya PS26 data.xlsx değil: TICKET, GUZERGAH ya da ALTYAPI sayfası yok.")
    o = Okuma(dosya_adi=dosya.name, sayfalar=sayfalar)
    try:
        if SAYFA_TICKET in sayfalar:
            o.ticket = ticket_aktar.satirlari_oku(dosya, SAYFA_TICKET)
        if SAYFA_GUZERGAH in sayfalar:
            # Aynı sütunlar (Ticket, Baslangic, Lokasyon, Site, Konu, Durum, Musteri, Kanal, Detay); boş
            # satırları satirlari_oku zaten atar (yalnız "Durum"u dolu satır da boş sayılır).
            o.guzergah = ticket_aktar.satirlari_oku(dosya, SAYFA_GUZERGAH)
        if SAYFA_ALTYAPI in sayfalar:
            o.altyapi = _altyapi_oku(dosya)
    except Ps26Hatasi:
        raise
    except Exception as exc:  # noqa: BLE001 — sütun adı değişmiş, bozuk hücre…
        raise Ps26Hatasi(f"data.xlsx okunamadı ({type(exc).__name__}). Sayfa başlıkları değişmiş olabilir.") from exc
    o.zipler = zipler if zipler is not None else arsivler(imgs)
    return o


# ============================================================================= anahtarlar
def _anahtarli(satirlar: list[dict], onek: str) -> list[tuple[str, dict]]:
    """``ticket_aktar.aktar`` ile AYNI anahtar: ``<önek>:<temel>:<aynı satırın kaçıncısı>``."""
    gorulen: Counter = Counter()
    sonuc = []
    for r in satirlar:
        temel = ticket_aktar._anahtar(r)
        gorulen[temel] += 1
        sonuc.append((f"{onek}:{temel}:{gorulen[temel]}", r))
    return sonuc


def _altyapi_anahtarli(satirlar: list[dict]) -> list[tuple[str, dict]]:
    """Yalnız satış açılırken yazılan ve sonra değişmeyen alanlardan: başlangıç + müşteri + kanal."""
    gorulen: Counter = Counter()
    sonuc = []
    for r in satirlar:
        ham = "|".join((r["baslangic"], r["musteri"], r["kanal"])).upper()
        temel = hashlib.sha1(ham.encode("utf-8")).hexdigest()[:16]
        gorulen[temel] += 1
        sonuc.append((f"ps26:{temel}:{gorulen[temel]}", r))
    return sonuc


# ============================================================================= GUZERGAH
def _guzergah_durumu(ham: str) -> str:
    """GÜZERGAH sayfasında 'YOK' (güzergah yok) = sorun sürüyor → AÇIK; tanınan durum aynen."""
    return ticket.durum_coz(ham) or "AÇIK"


def guzergah_aktar(conn: sqlite3.Connection, satirlar: list[dict], kuru: bool = False) -> dict:
    sayac: Counter = Counter()
    for anahtar, r in _anahtarli(satirlar, TUR_GUZERGAH):
        bina = ticket.lokasyondan_bina(conn, r["Lokasyon"]) if r["Lokasyon"] else None
        veri = {
            "ticket_no": r["Ticket"], "acilis": r["Baslangic"],
            "konu": ticket.konu_coz(r["Konu"]) or "GÜZERGAH", "durum": _guzergah_durumu(r["Durum"]),
            "bina_serial": bina["bina_serial"] if bina else None,
            "location_id": (bina or {}).get("location_id") or r["Lokasyon"],
            "site": r["Site"].strip(), "musteri": r["Musteri"], "kanal": ticket._tr_ust(r["Kanal"]),
            "detay": " · ".join(x for x in (r["Detay"], r["_ek"]) if x), "tur": TUR_GUZERGAH,
        }
        if conn.execute("SELECT 1 FROM ticket WHERE aktarim_anahtari=?", (anahtar,)).fetchone():
            sayac[ticket_aktar._esitle(conn, anahtar, veri, kuru)] += 1
        elif kuru:
            sayac["eklenecek"] += 1
        else:
            tid, _ = ticket.olustur(conn, veri, None, kaynak="excel", aktarim_anahtari=anahtar)
            sayac["eklenen" if tid else "zaten_var"] += 1
        sayac["binaya_baglanan" if bina else "binasiz"] += 1
    if not kuru:
        conn.commit()
    return {"okunan": len(satirlar), **dict(sayac)}


# ============================================================================= ALTYAPI
def altyapi_aktar(conn: sqlite3.Connection, satirlar: list[dict], kuru: bool = False) -> dict:
    """``altyapi_bekleyen``'e yazar. Durum ve bina uygulamanındır: Excel yalnız satıcı/bölge'yi, not boşsa notu tazeler."""
    sayac: Counter = Counter()
    simdi = ayarlar.zaman_metni()
    gorulen: set[str] = set()
    for anahtar, r in _altyapi_anahtarli(satirlar):
        gorulen.add(anahtar)
        var = conn.execute("SELECT * FROM altyapi_bekleyen WHERE aktarim_anahtari=?", (anahtar,)).fetchone()
        if var:
            fark = {a: (d or None) for a, d in (("satici_ad", r["satici"]), ("bolge", r["bolge"]))
                    if (d or None) != (var[a] or None)}
            if r["not"] and not var["notu"] and var["durum"] == "bekliyor":
                fark["notu"] = r["not"]
            if not fark:
                sayac["zaten_var"] += 1
                continue
            if not kuru:
                atama = ", ".join(f"{a}=?" for a in fark)
                conn.execute(f"UPDATE altyapi_bekleyen SET {atama}, guncelleme=? WHERE id=?",
                             (*fark.values(), simdi, var["id"]))
            sayac["guncellenen"] += 1
        elif kuru:
            sayac["eklenecek"] += 1
        else:
            conn.execute(
                "INSERT INTO altyapi_bekleyen (musteri_no, kanal, baslangic, satici_ad, bolge, bina_serial, durum, "
                "notu, aktarim_anahtari, olusturma, guncelleme) VALUES (?,?,?,?,?,NULL,'bekliyor',?,?,?,?) "
                "ON CONFLICT(aktarim_anahtari) DO NOTHING",
                (r["musteri"] or None, r["kanal"] or None, r["baslangic"] or None, r["satici"] or None,
                 r["bolge"] or None, r["not"] or None, anahtar, simdi, simdi))
            sayac["eklenen"] += 1
    # Excel'den silinmiş (büyük olasılıkla kurulmuş) ama bizde hâlâ bekleyen satırlar: silinmez, sayılır.
    yer = ",".join("?" * len(gorulen))
    excelde_yok = conn.execute(
        "SELECT COUNT(*) FROM altyapi_bekleyen WHERE aktarim_anahtari LIKE 'ps26:%' AND durum='bekliyor'"
        + (f" AND aktarim_anahtari NOT IN ({yer})" if gorulen else ""), tuple(gorulen)).fetchone()[0]
    if not kuru:
        conn.commit()
    satici = Counter(r["satici"] for r in satirlar if r["satici"])
    return {"okunan": len(satirlar), **dict(sayac), "excelde_yok": int(excelde_yok), "satici_sayisi": len(satici)}


def satici_ozeti(conn: sqlite3.Connection) -> list[dict]:
    """PVT'nin yerine: satıcı başına bekleyen / kurulan / iptal sayısı (türetilir, saklanmaz)."""
    try:
        satirlar = conn.execute(
            "SELECT COALESCE(NULLIF(TRIM(satici_ad),''), '(Satıcı yok)') AS satici, "
            "SUM(durum='bekliyor') AS bekliyor, SUM(durum='kuruldu') AS kuruldu, SUM(durum='iptal') AS iptal, "
            "COUNT(*) AS toplam FROM altyapi_bekleyen GROUP BY 1 ORDER BY 2 DESC, 5 DESC, 1").fetchall()
    except sqlite3.OperationalError:
        return []
    return [dict(r) for r in satirlar]


# ============================================================================= fotoğraflar
def _rakamlar(metin: str) -> list[str]:
    return re.findall(r"\d+", metin or "")


def zip_eslesmesi(okuma: Okuma) -> tuple[dict[str, str], Counter]:
    """Arşiv adı → ticket aktarım anahtarı. Önce TICKET, bulunamazsa GUZERGAH satırları.

    Arşivin adı satırın ticket numarasıysa ya da müşteri alanındaki numaralardan biriyse eşleşir
    ("12345678 / 23456789" gibi çok müşterili alanlar dahil). Birden çok satır eşleşirse en yeni açılış.
    """
    sayac: Counter = Counter()
    adaylar: dict[str, list[tuple[str, str]]] = {}
    for onek, satirlar in (("excel", okuma.ticket), (TUR_GUZERGAH, okuma.guzergah)):
        yerel: dict[str, list[tuple[str, str]]] = {}
        for anahtar, r in _anahtarli(satirlar, onek):
            for no in {ticket_aktar._metin(r["Ticket"]), *_rakamlar(r["Musteri"])}:
                if no:
                    yerel.setdefault(no, []).append((r["Baslangic"], anahtar))
        for no, liste in yerel.items():
            adaylar.setdefault(no, liste)            # TICKET önce; GUZERGAH yalnız TICKET'ta yoksa
    eslesme = {}
    for z in okuma.zipler or []:
        liste = adaylar.get(z.stem.strip())
        if not liste:
            continue
        if len({a for _, a in liste}) > 1:
            sayac["birden_cok_eslesen_zip"] += 1
        eslesme[z.stem.strip()] = max(liste)[1]
    return eslesme, sayac


def foto_aktar(conn: sqlite3.Connection, okuma: Okuma, kuru: bool = False) -> dict:
    if okuma.zipler is None:
        return {"klasor": False}
    eslesme, sayac = zip_eslesmesi(okuma)
    sayac.update({"zip": len(okuma.zipler)})
    goruldu: set[tuple] = set()
    for z in okuma.zipler:
        anahtar = eslesme.get(z.stem.strip())
        if not anahtar:
            sayac["eslesmeyen_zip"] += 1
            continue
        sayac["eslesen_zip"] += 1
        try:
            dosyalar, alinmayan = ek_dosya.arsiv_dosyalari(z)
        except (zipfile.BadZipFile, OSError):
            sayac["bozuk_zip"] += 1
            continue
        sayac["alinmayan"] += alinmayan
        sayac["dosya"] += len(dosyalar)
        satir = conn.execute("SELECT id FROM ticket WHERE aktarim_anahtari=?", (anahtar,)).fetchone()
        tid = int(satir[0]) if satir else None
        for ad, icerik in dosyalar:
            if kuru or tid is None:
                try:
                    hazir, _ = ek_dosya.hazirla(icerik, ad)
                except ek_dosya.EkHatasi:
                    sayac["alinmayan"] += 1
                    continue
                imza = (tid or anahtar, hashlib.sha256(hazir).hexdigest())
                if imza in goruldu or (tid is not None and ek_dosya.var_mi(conn, tid, imza[1])):
                    sayac["zaten_var"] += 1
                else:
                    sayac["eklenecek"] += 1
                goruldu.add(imza)
                continue
            try:
                _, yeni = ek_dosya.ekle(conn, tid, icerik, ad)
            except ek_dosya.EkHatasi:
                sayac["alinmayan"] += 1
                continue
            sayac["eklenen" if yeni else "zaten_var"] += 1
    if not kuru:
        conn.commit()
    return {"klasor": True, **dict(sayac)}


# ============================================================================= bütün aktarım
def _kapali_ticketlar(conn: sqlite3.Connection) -> set[int]:
    return {r[0] for r in conn.execute("SELECT id FROM ticket WHERE durum IN ('ÇÖZÜLDÜ','KAPATILDI','İPTAL')")}


def _is_denetimi(conn: sqlite3.Connection, ticket_idler, k: dict | None) -> int:
    """Kapanan ticket'a bağlı "altyapı" işleri bekliyor'a döner (spec §3.7; ``POST /api/ticket/aktar`` ile aynı)."""
    try:
        from operasyon.v2 import akis
    except ImportError:
        return 0
    fonksiyon = getattr(akis, "ticket_degisti", None)
    if not callable(fonksiyon):
        return 0
    donen = 0
    for tid in ticket_idler:
        donen += len(list(fonksiyon(conn, int(tid), k) or []))
    conn.commit()
    return donen


def aktar(conn: sqlite3.Connection, okuma: Okuma, kuru: bool = False, k: dict | None = None) -> dict:
    """Bütün sayfalar + fotoğraflar. ``kuru`` → hiçbir şey yazılmaz (önizleme). Yedek ÇAĞIRANIN işidir.

    Her sayfa kendi işleminde yazılır; yarıda kesilen aktarım yeniden çalıştırılınca kaldığı yerden tamamlanır
    (anahtarlar etkisiz). Dönen sözlükte kişisel veri yoktur, yalnız sayılar.
    """
    t0 = time.monotonic()
    once = _kapali_ticketlar(conn)
    sonuc: dict = {"dosya": okuma.dosya_adi, "kuru": kuru}
    s = ticket_aktar.aktar(conn, okuma.ticket, kuru=kuru) if okuma.ticket else {"okunan": 0}
    sonuc["ticket"] = {a: s.get(a, 0) for a in ("okunan", "eklenen", "eklenecek", "zaten_var", "guncellenen",
                                                 "uygulamada_degismis", "binaya_baglanan", "binasiz")}
    sonuc["ticket"]["bilinmeyen_konu_durum"] = len(s.get("bilinmeyen_konu_durum", []))
    sonuc["guzergah"] = guzergah_aktar(conn, okuma.guzergah, kuru=kuru)
    sonuc["altyapi"] = altyapi_aktar(conn, okuma.altyapi, kuru=kuru)
    sonuc["foto"] = foto_aktar(conn, okuma, kuru=kuru)
    sonuc["bekliyora_donen_is"] = 0
    if not kuru:
        ozet = {"ticket": sonuc["ticket"].get("eklenen", 0), "guzergah": sonuc["guzergah"].get("eklenen", 0),
                "altyapi": sonuc["altyapi"].get("eklenen", 0), "foto": sonuc["foto"].get("eklenen", 0)}
        db.yonetim_kaydi(conn, k, "ps26_aktar", okuma.dosya_adi, ozet)
        db.ayar_yaz(conn, "ps26_son_aktarim", ayarlar.zaman_metni())
        conn.commit()
        sonuc["bekliyora_donen_is"] = _is_denetimi(conn, sorted(_kapali_ticketlar(conn) - once), k)
    sonuc["eklenecek_toplam"] = sum(sonuc[b].get("eklenecek", 0) + sonuc[b].get("guncellenen", 0)
                                    for b in ("ticket", "guzergah", "altyapi", "foto"))
    sonuc["sure_sn"] = round(time.monotonic() - t0, 2)
    return sonuc


# ============================================================================= komut satırı
def _yaz(sonuc: dict) -> None:
    kuru = sonuc["kuru"]
    ek = "eklenecek" if kuru else "eklenen"
    t, g, a, f = sonuc["ticket"], sonuc["guzergah"], sonuc["altyapi"], sonuc["foto"]
    print(f"{sonuc['dosya']}: {'ÖNİZLEME — hiçbir şey yazılmadı' if kuru else 'aktarıldı'} ({sonuc['sure_sn']} sn)")
    print(f"  TICKET    : {t.get('okunan', 0)} satır · {ek} {t.get(ek, 0)} · zaten vardı {t.get('zaten_var', 0)} · "
          f"Excel'den güncellenen {t.get('guncellenen', 0)} · uygulamada değiştiği için dokunulmayan "
          f"{t.get('uygulamada_degismis', 0)} · binaya bağlı {t.get('binaya_baglanan', 0)}")
    print(f"  GUZERGAH  : {g.get('okunan', 0)} dolu satır · {ek} {g.get(ek, 0)} · zaten vardı {g.get('zaten_var', 0)} · "
          f"güncellenen {g.get('guncellenen', 0)} · binaya bağlı {g.get('binaya_baglanan', 0)}")
    print(f"  ALTYAPI   : {a.get('okunan', 0)} satır · {ek} {a.get(ek, 0)} · zaten vardı {a.get('zaten_var', 0)} · "
          f"güncellenen {a.get('guncellenen', 0)} · Excel'de artık olmayan bekleyen {a.get('excelde_yok', 0)} · "
          f"{a.get('satici_sayisi', 0)} satıcı")
    if not f.get("klasor"):
        print("  Fotoğraf  : imgs klasörü bulunamadı (atlandı)")
    else:
        print(f"  Fotoğraf  : {f.get('zip', 0)} arşiv · eşleşen {f.get('eslesen_zip', 0)} · eşleşmeyen "
              f"{f.get('eslesmeyen_zip', 0)} · {ek} {f.get(ek, 0)} · zaten vardı {f.get('zaten_var', 0)} · "
              f"alınmayan (resim değil/bozuk/10 MB üstü) {f.get('alinmayan', 0)}")


def main(argv: list[str] | None = None) -> int:
    ayrac = argparse.ArgumentParser(description="PS26 data.xlsx'i Saha Sistemine aktarır (TICKET, GUZERGAH, "
                                                "ALTYAPI + imgs fotoğrafları). İki kez çalıştırmak güvenlidir.")
    ayrac.add_argument("dosya", nargs="?", default=None, help="data.xlsx (varsayılan: bulunan PS26 klasörü)")
    ayrac.add_argument("--imgs", default=None, help="Fotoğraf arşivleri klasörü (varsayılan: data.xlsx'in yanındaki imgs)")
    ayrac.add_argument("--db", default=None, help="Veritabanı (varsayılan: saha/saha.db ya da SAHA_DB)")
    ayrac.add_argument("--kuru", action="store_true", help="Önizleme: ne olacağını yaz, hiçbir şey değiştirme")
    a = ayrac.parse_args(argv)
    ayarlar.konsolu_hazirla()
    from . import goc, yedekle

    hedef = Path(a.db) if a.db else db.db_yolu()
    if not hedef.exists():
        print(f"Veritabanı bulunamadı: {hedef}")
        return 1
    # Fotoğraflar verilen veritabanının yanına (kopya üzerinde çalışırken canlının ek/ klasörüne değil).
    ek_yonlendir = bool(a.db) and not os.environ.get("SAHA_EK")
    if ek_yonlendir:
        os.environ["SAHA_EK"] = str(hedef.parent / "ek")
    conn = db.baglan(hedef)
    try:
        # Göç burada YAPILMAZ (yalnız sunucu açılışında); sürüm uymuyorsa araç durur.
        try:
            db.semayi_kur(conn)
        except goc.SurumUyumsuz as exc:
            print(exc.mesaj)
            return 3
        if a.dosya:
            dosya = Path(a.dosya)
        else:
            klasor = klasor_bul(conn)
            if klasor is None:
                print("PS26 klasörü bulunamadı. data.xlsx'in yolunu verin: python -m saha.ps26_aktar <data.xlsx>")
                return 1
            dosya = klasor / "data.xlsx"
        if not dosya.is_file():
            print(f"Dosya bulunamadı: {dosya}")
            return 1
        imgs = Path(a.imgs) if a.imgs else dosya.parent / "imgs"
        try:
            okuma = oku(dosya, imgs)
        except Ps26Hatasi as exc:
            print(exc.mesaj)
            return 2
        if not a.kuru:
            print(f"Yedek alındı: {yedekle.aktarim_yedegi(hedef)}")
        sonuc = aktar(conn, okuma, kuru=a.kuru)
    finally:
        conn.close()
        if ek_yonlendir:
            os.environ.pop("SAHA_EK", None)
    _yaz(sonuc)
    return 0


if __name__ == "__main__":
    sys.exit(main())
