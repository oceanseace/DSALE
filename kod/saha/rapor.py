"""Günlük Excel raporu (openpyxl).

Üç sayfa, tek bakışta okunur:
    GUN         satışçı bazında ziyaret / satış / ret / randevu / kalan
    ZIYARETLER  o günün bütün ziyaretleri, bina ve adresle birlikte
    KAPSAMA     bölge bazında "neye dokunuldu, ne kaldı"

Görsel dil ``dsale/excel_report.py`` ile aynı (lacivert bant, sarı vurgu) ama
bilerek sade: bu rapor her gün açılıp kapanacak.
"""
from __future__ import annotations

import datetime as dt
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import ayarlar, db


def rapor_dizini() -> Path:
    """Raporlar HANGİ veritabanı kullanılıyorsa onun yanına yazılır.

    Sabit bir klasör kullanılınca test ya da gösterim koşusu üretim
    klasörünü kirletiyordu.
    """
    return db.db_yolu().parent / "raporlar"

LACIVERT = "1B2A41"
VURGU = "FFC400"
ZEBRA = "F7F9FC"
CIZGI = "D9DEE7"
KOYU = "1F2937"
GRI = "6B7280"
YESIL_Z = "DCFCE7"

FMT_SAYI = "#,##0"
FMT_YUZDE = "0.0%"

_INCE = Side(style="thin", color=CIZGI)
_KENAR = Border(left=_INCE, right=_INCE, top=_INCE, bottom=_INCE)


def _baslik_bandi(ws, metin: str, alt: str, sutun_sayisi: int) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=sutun_sayisi)
    h = ws.cell(row=1, column=1, value=metin)
    h.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    h.fill = PatternFill("solid", fgColor=LACIVERT)
    h.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=sutun_sayisi)
    a = ws.cell(row=2, column=1, value=alt)
    a.font = Font(name="Calibri", size=10, color=GRI)
    a.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 18
    return 4


def _tablo_basligi(ws, satir: int, basliklar: list[str]) -> None:
    for i, b in enumerate(basliklar, start=1):
        h = ws.cell(row=satir, column=i, value=b)
        h.font = Font(name="Calibri", size=10, bold=True, color=LACIVERT)
        h.fill = PatternFill("solid", fgColor=VURGU)
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = _KENAR
    ws.row_dimensions[satir].height = 26
    ws.freeze_panes = ws.cell(row=satir + 1, column=1)


def _veri_satiri(ws, satir: int, degerler: list, bicimler: dict[int, str] | None = None,
                 vurgula: bool = False) -> None:
    bicimler = bicimler or {}
    zemin = PatternFill("solid", fgColor=ZEBRA) if satir % 2 == 0 else None
    if vurgula:
        zemin = PatternFill("solid", fgColor=YESIL_Z)
    for i, d in enumerate(degerler, start=1):
        h = ws.cell(row=satir, column=i, value=d)
        h.font = Font(name="Calibri", size=10, color=KOYU, bold=vurgula)
        h.border = _KENAR
        if zemin:
            h.fill = zemin
        if i in bicimler:
            h.number_format = bicimler[i]
            h.alignment = Alignment(horizontal="right")
        else:
            h.alignment = Alignment(horizontal="left", vertical="center")


def _genislikler(ws, genislikler: list[int]) -> None:
    for i, g in enumerate(genislikler, start=1):
        ws.column_dimensions[get_column_letter(i)].width = g


# ----------------------------------------------------------------------------- sayfalar
def _sayfa_gun(wb: Workbook, conn: sqlite3.Connection, gun: str) -> None:
    ws = wb.active
    ws.title = "GUN"
    basliklar = ["Satışçı", "Bölge", "Ziyaret", "Satış", "Satış adedi", "İlgilenmedi",
                 "Evde yok", "Randevu", "Girilemedi", "Altyapı", "Kalan", "Dönüşüm"]
    satir = _baslik_bandi(ws, "SAHA GÜNLÜK RAPORU", f"{gun} · Dehanet EÇM · 8 satışçı", len(basliklar))
    _tablo_basligi(ws, satir, basliklar)
    _genislikler(ws, [34, 8, 10, 9, 12, 12, 11, 10, 12, 10, 9, 11])

    kullanicilar = conn.execute(
        "SELECT id, ad, bolge FROM kullanici WHERE rol='satisci' ORDER BY bolge, id"
    ).fetchall()
    toplam = [0] * 9
    for k in kullanicilar:
        z = conn.execute(
            "SELECT COUNT(*) AS ziyaret, "
            "COALESCE(SUM(CASE WHEN sonuc='satis' THEN 1 ELSE 0 END),0) AS satis, "
            "COALESCE(SUM(satis_adedi),0) AS adet, "
            "COALESCE(SUM(CASE WHEN sonuc='ilgilenmedi' THEN 1 ELSE 0 END),0) AS ret, "
            "COALESCE(SUM(CASE WHEN sonuc='evde_yok' THEN 1 ELSE 0 END),0) AS evde_yok, "
            "COALESCE(SUM(CASE WHEN sonuc='randevu' THEN 1 ELSE 0 END),0) AS randevu, "
            "COALESCE(SUM(CASE WHEN sonuc='girilemedi' THEN 1 ELSE 0 END),0) AS girilemedi, "
            "COALESCE(SUM(CASE WHEN sonuc='altyapi_sorunu' THEN 1 ELSE 0 END),0) AS altyapi "
            "FROM ziyaret WHERE kullanici_id=? AND iptal=0 AND substr(zaman,1,10)=?",
            (k["id"], gun),
        ).fetchone()
        kalan = conn.execute(
            "SELECT COUNT(*) FROM gorev_bina gb JOIN gorev g ON g.id=gb.gorev_id "
            "WHERE g.kullanici_id=? AND g.tarih=? AND gb.durum='bekliyor'",
            (k["id"], gun),
        ).fetchone()[0]
        degerler = [z["ziyaret"], z["satis"], int(z["adet"]), z["ret"], z["evde_yok"],
                    z["randevu"], z["girilemedi"], z["altyapi"], kalan]
        toplam = [a + b for a, b in zip(toplam, degerler)]
        satir += 1
        _veri_satiri(
            ws, satir,
            [k["ad"], k["bolge"], *degerler, (z["satis"] / z["ziyaret"]) if z["ziyaret"] else 0],
            {i: FMT_SAYI for i in range(3, 12)} | {12: FMT_YUZDE},
        )

    satir += 1
    _veri_satiri(
        ws, satir,
        ["TOPLAM", "", *toplam, (toplam[1] / toplam[0]) if toplam[0] else 0],
        {i: FMT_SAYI for i in range(3, 12)} | {12: FMT_YUZDE},
        vurgula=True,
    )


def _sayfa_ziyaretler(wb: Workbook, conn: sqlite3.Connection, gun: str) -> None:
    ws = wb.create_sheet("ZIYARETLER")
    basliklar = ["Saat", "Satışçı", "Bölge", "Bina", "Mahalle", "Adres", "Daire",
                 "Boş kapı", "Sonuç", "Satış", "Not"]
    satir = _baslik_bandi(ws, "GÜNÜN ZİYARETLERİ", f"{gun} · her satır bir binaya temas", len(basliklar))
    _tablo_basligi(ws, satir, basliklar)
    _genislikler(ws, [8, 26, 7, 34, 18, 34, 8, 10, 16, 8, 40])

    kayitlar = conn.execute(
        "SELECT z.zaman, z.sonuc, z.satis_adedi, z.notu, k.ad AS satisci, "
        "       b.ad, b.site_adi, b.mahalle, b.sokak, b.kapi_no, b.bolge, b.daire, b.firsat "
        "FROM ziyaret z "
        "LEFT JOIN kullanici k ON k.id=z.kullanici_id "
        "LEFT JOIN bina b ON b.bina_serial=z.bina_serial "
        "WHERE z.iptal=0 AND substr(z.zaman,1,10)=? ORDER BY z.zaman",
        (gun,),
    ).fetchall()
    for z in kayitlar:
        satir += 1
        ad = (z["ad"] or z["site_adi"] or "").strip()
        adres = " ".join(x for x in [(z["sokak"] or "").strip(), (z["kapi_no"] or "").strip()] if x)
        _veri_satiri(
            ws, satir,
            [(z["zaman"] or "")[11:16], z["satisci"] or "", z["bolge"], ad, z["mahalle"] or "", adres,
             z["daire"] or 0, z["firsat"] or 0,
             ayarlar.SONUC_ETIKET.get(z["sonuc"], z["sonuc"]), z["satis_adedi"] or 0, z["notu"] or ""],
            {7: FMT_SAYI, 8: FMT_SAYI, 10: FMT_SAYI},
            vurgula=(z["sonuc"] == "satis"),
        )
    if not kayitlar:
        ws.cell(row=satir + 1, column=1, value="Bu tarihte kayıtlı ziyaret yok.").font = Font(
            name="Calibri", size=10, italic=True, color=GRI
        )


def _sayfa_kapsama(wb: Workbook, conn: sqlite3.Connection, gun: str) -> None:
    ws = wb.create_sheet("KAPSAMA")
    basliklar = ["Bölge", "Bina", "Gidilen", "Görüşülen", "Kalan", "Kapsama",
                 "Boş kapı", "Satılmamış kapı", "Satış"]
    satir = _baslik_bandi(
        ws, "KAPSAMA — NEREYE DOKUNULDU", f"{gun} günü sonu itibarıyla · boşluklar doldukça oran yükselir",
        len(basliklar),
    )
    _tablo_basligi(ws, satir, basliklar)
    _genislikler(ws, [10, 12, 12, 13, 11, 12, 13, 17, 10])

    # TARİHE GÖRE: eskiden hangi tarih istenirse istensin O ANKİ kapsama
    # basılıyor ama alt başlıkta istenen tarih yazıyordu — rapor sessizce
    # yanlış bilgi veriyordu. Artık o günün sonundaki durum hesaplanır.
    temas = ",".join("'%s'" % x for x in ayarlar.TEMAS_SONUCLARI)
    sorgu = f"""
        WITH son AS (
            SELECT z.bina_serial,
                   MAX(z.zaman) AS son_zaman,
                   COALESCE(SUM(CASE WHEN z.sonuc='satis' THEN z.satis_adedi ELSE 0 END),0) AS satis
            FROM ziyaret z
            WHERE z.iptal=0 AND substr(z.zaman,1,10) <= ?
            GROUP BY z.bina_serial
        ),
        gorusulen AS (
            SELECT DISTINCT z.bina_serial FROM ziyaret z
            WHERE z.iptal=0 AND substr(z.zaman,1,10) <= ? AND z.sonuc IN ({temas})
        )
        SELECT b.bolge AS bolge, COUNT(*) AS toplam,
               SUM(CASE WHEN son.bina_serial IS NOT NULL THEN 1 ELSE 0 END) AS dokunulan,
               SUM(CASE WHEN g.bina_serial IS NOT NULL THEN 1 ELSE 0 END) AS temas,
               COALESCE(SUM(b.firsat),0) AS firsat,
               COALESCE(SUM(MAX(b.firsat - COALESCE(son.satis,0), 0)),0) AS kalan_firsat,
               COALESCE(SUM(COALESCE(son.satis,0)),0) AS satis
        FROM bina b
        LEFT JOIN son ON son.bina_serial = b.bina_serial
        LEFT JOIN gorusulen g ON g.bina_serial = b.bina_serial
        GROUP BY b.bolge ORDER BY b.bolge
    """
    for s in conn.execute(sorgu, (gun, gun)).fetchall():
        satir += 1
        toplam, dokunulan = s["toplam"], s["dokunulan"] or 0
        _veri_satiri(
            ws, satir,
            [s["bolge"] if s["bolge"] is not None else "Atanmamış", toplam, dokunulan,
             s["temas"] or 0, toplam - dokunulan, (dokunulan / toplam) if toplam else 0,
             int(s["firsat"]), int(s["kalan_firsat"]), int(s["satis"])],
            {2: FMT_SAYI, 3: FMT_SAYI, 4: FMT_SAYI, 5: FMT_SAYI, 6: FMT_YUZDE,
             7: FMT_SAYI, 8: FMT_SAYI, 9: FMT_SAYI},
        )

    satir += 2
    ws.cell(row=satir, column=1, value=(
        "Gidilen: binaya gidildi (sonuç ne olursa olsun) · "
        "Görüşülen: kapı açıldı, biriyle konuşuldu · "
        "Satılmamış kapı: boş kapı eksi satılan abonelik"
    )).font = Font(name="Calibri", size=9, italic=True, color=GRI)


# ----------------------------------------------------------------------------- giriş
def gunluk_rapor_yaz(conn: sqlite3.Connection, gun: dt.date | str, hedef: Path | None = None) -> Path:
    """Günlük raporu üretir ve dosya yolunu döndürür."""
    tarih = gun.isoformat() if isinstance(gun, dt.date) else str(gun)[:10]
    wb = Workbook()
    _sayfa_gun(wb, conn, tarih)
    _sayfa_ziyaretler(wb, conn, tarih)
    _sayfa_kapsama(wb, conn, tarih)

    yol = Path(hedef) if hedef else (rapor_dizini() / f"Saha_Gun_Raporu_{tarih}.xlsx")
    yol.parent.mkdir(parents=True, exist_ok=True)
    wb.save(yol)
    return yol
