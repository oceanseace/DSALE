"""Bölgeleme çıktılarının bütünlük denetimi (CCPD-2.5).

    python scripts/dogrula.py                    # outputs/N08_res_hp + app veri paketi
    python scripts/dogrula.py --klasor outputs/N10_res_hp
    python scripts/dogrula.py --ozet             # yalnız atama.csv özeti (determinizm karşılaştırması)

Denetlenenler
  1. Her bina tam olarak bir bölgeye atanmış; bölge toplamları master toplamlarına eşit.
  2. N=8 sapması ≤ ±%0,5; paketteki her N ≤ ±%1 (N ≥ 40 için sert sınır ±%2).
  3. Bölge poligonları örtüşmüyor (ikili kesişim ≤ 0,01 km²) — N = 8/10/20/35/50.
  4. Ada kalmamış: bütün komşuları başka bölgede olan 100 HP üstü birim yok (N=8).
  5. Bölge adları ve kısa adları her N'de tekil.
  6. anlık etiket dizileri birim sayısı uzunluğunda; son kare gerçek atamayla ≥ %95 uyumlu.
  7. meta.mahalle_sayisi = (ilçe, mahalle) çifti sayısı.
  8. Bina adı sözlüğünde "Null" / "None" yok.
  9. Bütün planlarda parametreler.algoritma = CCPD-2.5.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
import shapely  # noqa: E402
from shapely.geometry import shape  # noqa: E402

from dsale import config  # noqa: E402
from dsale.enrich import M_PER_DEG_LAT, _m_per_deg_lon  # noqa: E402
from dsale.partition import ALGORITMA_SURUMU, _yerlesik_ada  # noqa: E402

ORTUSME_SINIRI_KM2 = 0.01
POLIGON_N = (8, 10, 20, 35, 50)

_sonuc: list[tuple[bool, str]] = []


def kontrol(ad: str, tamam: bool, ayrinti: str = "") -> bool:
    _sonuc.append((tamam, ad))
    print(f"  {'✓' if tamam else '✗'} {ad}{('  — ' + ayrinti) if ayrinti else ''}")
    return tamam


def _alan_km2(g) -> float:
    """Derece cinsinden geometrinin yaklaşık km² alanı (Bursa enlemi)."""
    return float(shapely.area(g)) * (M_PER_DEG_LAT / 1000) * (_m_per_deg_lon(40.2) / 1000)


def poligon_ortusmesi(bolgeler: list[dict]) -> tuple[float, str]:
    g = [(b["kod"], shapely.make_valid(shape(b["poligon"]))) for b in bolgeler if b.get("poligon")]
    en = (0.0, "")
    for i in range(len(g)):
        for j in range(i + 1, len(g)):
            a = _alan_km2(shapely.intersection(g[i][1], g[j][1]))
            if a > en[0]:
                en = (a, f"{g[i][0]}×{g[j][0]}")
    return en


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--klasor", default=str(config.OUTPUTS / "N08_res_hp"))
    ap.add_argument("--ozet", action="store_true", help="yalnız atama.csv md5 + sapma (determinizm için)")
    a = ap.parse_args()
    klasor = Path(a.klasor)

    atama_yolu = klasor / "atama.csv"
    if a.ozet:
        h = hashlib.md5(atama_yolu.read_bytes()).hexdigest()
        oz = json.load(open(klasor / "bolgeler.json", encoding="utf-8"))
        print(h, f"{min(b['sapma'] for b in oz['bolgeler']):+.6%}/{max(b['sapma'] for b in oz['bolgeler']):+.6%}")
        return 0

    df = pd.read_csv(config.MASTER_CSV, encoding="utf-8-sig", dtype={"tellcordia_id": str, "location_id": str})
    oz = json.load(open(klasor / "bolgeler.json", encoding="utf-8"))
    at = pd.read_csv(atama_yolu, encoding="utf-8-sig", dtype={"bina_serial": str})
    n = int(oz["n"])
    print(f"\n{klasor.name}  ·  N={n}  ·  {len(df):,} bina")

    # 1 — atama bütünlüğü
    print("\n1) Atama bütünlüğü")
    kontrol("her bina tam olarak bir kez atanmış",
            len(at) == len(df) and at["bina_serial"].nunique() == len(df) and set(at["bina_serial"]) == set(df["bina_serial"]),
            f"{len(at)} satır / {at['bina_serial'].nunique()} tekil")
    kontrol("bölge numaraları 1..N", bool(at["bolge"].between(1, n).all()))
    d = df.merge(at[["bina_serial", "bolge"]], on="bina_serial")
    top = d.groupby("bolge")[["bina", "res_hp", "toplam_hp", "aktif_res", "firsat"]].sum()
    kontrol("bölge toplamları = master toplamları",
            all(int(top[c].sum()) == int(df[c].sum()) for c in ["bina", "res_hp", "toplam_hp", "aktif_res", "firsat"]))
    json_top = {c: sum(b[c] for b in oz["bolgeler"]) for c in ["bina", "res_hp", "aktif_res", "firsat"]}
    kontrol("bolgeler.json toplamları = master",
            all(json_top[c] == int(df["bina" if c == "bina" else c].sum()) for c in json_top))

    # 2 — denge
    print("\n2) Denge")
    sap = [b["sapma"] for b in oz["bolgeler"]]
    kontrol(f"N={n} sapma ≤ ±%0,5", max(abs(x) for x in sap) <= 0.005, f"{min(sap):+.3%} / {max(sap):+.3%}")

    planlar_yolu = config.ROOT / "app" / "src" / "data" / "generated" / "planlar.json"
    planlar = json.load(open(planlar_yolu, encoding="utf-8"))["planlar"] if planlar_yolu.exists() else {}
    if planlar:
        kotu = [(int(k), p["sapma_min"], p["sapma_maks"]) for k, p in planlar.items()
                if max(abs(p["sapma_min"]), abs(p["sapma_maks"])) > (0.02 if int(k) >= 40 else 0.01)]
        en_kotu = max(((max(abs(p["sapma_min"]), abs(p["sapma_maks"])), int(k)) for k, p in planlar.items()), default=(0, 0))
        kontrol(f"paketteki {len(planlar)} planın tamamı sınır içinde", not kotu,
                f"en kötü N={en_kotu[1]} ±{en_kotu[0]:.3%}" + (f" · sınır dışı: {kotu}" if kotu else ""))
        kontrol("bütün planlar ≤ ±%1",
                all(max(abs(p["sapma_min"]), abs(p["sapma_maks"])) <= 0.01 for p in planlar.values()))

    # 3 — poligonlar
    print("\n3) Bölge poligonları")
    en_ort, cift = poligon_ortusmesi(oz["bolgeler"])
    kontrol(f"N={n} poligon örtüşmesi ≤ {ORTUSME_SINIRI_KM2} km²", en_ort <= ORTUSME_SINIRI_KM2,
            f"en büyük {en_ort:.2e} km² ({cift})")
    for k in POLIGON_N:
        p = planlar.get(str(k))
        if not p:
            continue
        e, c = poligon_ortusmesi(p["bolgeler"])
        kontrol(f"N={k} poligon örtüşmesi ≤ {ORTUSME_SINIRI_KM2} km²", e <= ORTUSME_SINIRI_KM2, f"en büyük {e:.2e} km² ({c})")
    eksik = [b["kod"] for b in oz["bolgeler"] if not b.get("etiket")]
    kontrol("her bölgenin etiket noktası var", not eksik, f"eksik: {eksik}" if eksik else "")
    dis = [b["kod"] for b in oz["bolgeler"]
           if b.get("poligon") and not shapely.contains(shapely.make_valid(shape(b["poligon"])), shapely.points(*b["etiket"]))]
    kontrol("etiket noktaları poligonun içinde", not dis, f"dışarıda: {dis}" if dis else "")

    # 4 — ada
    print("\n4) Ada (enclave)")
    g = d.groupby("site_grup", sort=True)
    lat0, lon0 = float(df["lat"].mean()), float(df["lon"].mean())
    X = np.c_[(g["lon"].mean() - lon0) * _m_per_deg_lon(lat0) / 1000, (g["lat"].mean() - lat0) * M_PER_DEG_LAT / 1000]
    bw = g["res_hp"].sum().to_numpy().astype(float)
    et = g["bolge"].first().to_numpy() - 1
    ada = _yerlesik_ada(X, et)
    buyuk = {i: k for i, k in ada.items() if bw[i] > 100}
    kontrol("100 HP üstü ada kalmadı", not buyuk,
            f"toplam {len(ada)} küçük ada" + (f" · büyük: {[(int(bw[i]), k + 1) for i, k in buyuk.items()]}" if buyuk else ""))

    # 5 — adlar
    print("\n5) Bölge adları")
    tekrar = []
    for etiket, veri in ([(f"N={n}", oz["bolgeler"])]
                         + [(f"N={k}", p["bolgeler"]) for k, p in sorted(planlar.items(), key=lambda x: int(x[0]))]):
        adlar = [b["ad"] for b in veri if not b.get("bos")]
        kisalar = [b["kisa_ad"] for b in veri if not b.get("bos")]
        if len(set(adlar)) != len(adlar):
            tekrar.append(f"{etiket} ad {sorted(x for x in set(adlar) if adlar.count(x) > 1)}")
        if len(set(kisalar)) != len(kisalar):
            tekrar.append(f"{etiket} kısa {sorted(x for x in set(kisalar) if kisalar.count(x) > 1)}")
    kontrol("bütün planlarda ad ve kısa ad tekil", not tekrar, " | ".join(tekrar[:6]))
    yalova = [b["kod"] for b in oz["bolgeler"] if b.get("yalova_bina")]
    yal_ad = [b["ad"] for b in oz["bolgeler"] if b.get("yalova_bina")]
    kontrol("Yalova bölgesinin adında Yalova geçiyor", all("Yalova" in x for x in yal_ad), f"{yalova}: {yal_ad}")

    # 6–9 — uygulama veri paketi
    print("\n6-9) Uygulama veri paketi")
    gen = config.ROOT / "app" / "src" / "data" / "generated"
    if planlar:
        binalar = json.load(open(gen / "binalar.json", encoding="utf-8"))
        meta = json.load(open(gen / "meta.json", encoding="utf-8"))
        birim_sayisi = len(binalar["sozluk"]["birim"])
        ana = planlar[str(n)]
        ae = ana.get("anlik_etiket") or []
        kontrol("anlık etiket dizisi var", len(ae) > 0, f"{len(ae)} kare")
        kontrol("anlık etiket uzunlukları = birim sayısı", all(len(s) == birim_sayisi for s in ae), f"birim {birim_sayisi}")
        if ae:
            son = np.array([ord(c) - 48 for c in ae[-1]])
            uyum = float((son == np.array(ana["birim_bolge"])).mean())
            kontrol("son kare gerçek atamayla ≥ %95 uyumlu", uyum >= 0.95, f"%{uyum * 100:.1f}")
        kontrol("meta.mahalle_sayisi = (ilçe, mahalle) çifti",
                meta["mahalle_sayisi"] == df.groupby(["ilce", "mahalle"]).ngroups,
                f"{meta['mahalle_sayisi']} / {df.groupby(['ilce', 'mahalle']).ngroups}")
        kotu_ad = [x for x in binalar["sozluk"]["ad"] if str(x).strip().lower() in ("null", "none", "nan")]
        kontrol("bina adı sözlüğünde 'Null' yok", not kotu_ad, str(kotu_ad[:5]))
        yanlis = [k for k, p in planlar.items() if (p.get("parametreler") or {}).get("algoritma") != ALGORITMA_SURUMU]
        kontrol(f"bütün planlar {ALGORITMA_SURUMU}", not yanlis, f"farklı: {yanlis[:10]}")
        kontrol("her planın birim dizisi tam", all(len(p["birim_bolge"]) == birim_sayisi for p in planlar.values()))
        kontrol("hiçbir birim atanmamış kalmamış", all(min(p["birim_bolge"]) >= 1 for p in planlar.values()))
    else:
        print("  (planlar.json yok — paketi üretmek için: python bolge.py --n 8)")

    # 10. Excel BINA_ATAMA'da da metin "Null" kalmamalı (app paketi temiz olsa bile).
    xlsx = next(iter(sorted(klasor.glob("*.xlsx"))), None)
    if xlsx is not None:
        import openpyxl

        wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
        kotu = 0
        if "BINA_ATAMA" in wb.sheetnames:
            ws = wb["BINA_ATAMA"]
            for satir in ws.iter_rows(min_row=2, values_only=True):
                kotu += sum(1 for h in satir if isinstance(h, str) and h.strip().lower() in ("null", "none", "nan"))
        wb.close()
        kontrol("Excel BINA_ATAMA'da 'Null' metni yok", kotu == 0, f"{kotu} hücre")
    else:
        print("  (Excel yok — üretmek için: python bolge.py --n 8)")

    hatali = [ad for tamam, ad in _sonuc if not tamam]
    print(f"\n{'=' * 70}\n{len(_sonuc) - len(hatali)}/{len(_sonuc)} denetim geçti")
    if hatali:
        print("BAŞARISIZ: " + " · ".join(hatali))
        return 1
    print("✓ hepsi tamam")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
